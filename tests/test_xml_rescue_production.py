"""Actual production composition without a pair matching prerequisite."""
import io
from pathlib import Path
import sys
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'v1.2-xml-experiment/res/app/webapp'),
               str(ROOT / 'v1.2-xml-experiment/res/app')]
from app import app


def source(role):
    document = Document()
    document.add_paragraph('教学目标：' + ('教师专有目标' if role == 'teacher' else '学生专有目标'))
    document.add_paragraph('重点难点：' + ('教师专有重点' if role == 'teacher' else '学生专有重点'))
    document.add_paragraph('即时训练', style='Heading 1')
    document.add_paragraph('1. ' + ('教师原稿问题甲' if role == 'teacher' else '学生原稿问题乙'))
    if role == 'teacher':
        document.add_paragraph('【答案】甲的独有答案与解析。')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def configure(tmp_path, monkeypatch):
    for key, value in dict(TESTING=True, RESULT_ROOT=tmp_path / 'results',
        RUNTIME_ROOT=tmp_path / 'jobs', C0_RUN_JOBS_SYNCHRONOUSLY=True,
        C0_DISABLE_JOB_SUBMISSION=False, C0_FORCE_FALLBACK_REASON=None).items():
        monkeypatch.setitem(app.config, key, value)


def output_text(path):
    return ''.join(node.text or '' for node in Document(path).element.body.iter(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))


def test_unmatched_pair_generates_from_own_body_and_metadata(tmp_path, monkeypatch):
    configure(tmp_path, monkeypatch)
    client = app.test_client()
    response = client.post('/api/jobs', data={'files': [
        (io.BytesIO(source('teacher')), '测试专题教师版.docx'),
        (io.BytesIO(source('student')), '测试专题学生版.docx')]})
    job = client.get('/api/jobs/' + response.get_json()['job_id']).get_json()
    assert job['status'] == 'done' and job['renderer'] == 'XML'
    assert job['pair_alignment_status'] == 'NOT_REQUIRED_NOT_VERIFIED'
    teacher = output_text(job['teacher_output_path'])
    student = output_text(job['student_output_path'])
    assert '教师原稿问题甲' in teacher and '学生原稿问题乙' not in teacher
    assert '学生原稿问题乙' in student and '教师原稿问题甲' not in student
    assert '教师专有目标' in teacher and '学生专有目标' in student
    assert '教师专有目标' not in student and '独有答案' not in student
    assert job['product_integrity']['accepted']


def test_unproven_teacher_only_publishes_teacher_without_unsafe_student(tmp_path, monkeypatch):
    configure(tmp_path, monkeypatch)
    client = app.test_client()
    response = client.post('/api/jobs', data={'files': [(io.BytesIO(source('teacher')), '教师版.docx')]})
    job = client.get('/api/jobs/' + response.get_json()['job_id']).get_json()
    assert job['status'] == 'done' and job['renderer'] == 'XML'
    assert set(job['output_roles']) == {'teacher'}
    assert job['student_preparation_route'] == 'NOT_GENERATED_SAFETY_UNPROVEN'
    assert job['student_preparation']['make_student_called'] is False
    assert any('仅生成教师版' in warning for warning in job['warnings'])
    assert job['product_integrity']['accepted']
