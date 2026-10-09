"""Synthetic backend batch integration; these tests do not claim real COM UAT."""
import io
import hashlib
import json
from pathlib import Path
import sys
import threading
from types import SimpleNamespace
import zipfile

from docx import Document
import pytest
from product_fixture_utils import write_product_fixture as _write_product_fixture


def write_product_fixture(*args, **kwargs):
    # These renderer doubles exercise lifecycle, not semantic extraction.
    # Return a populated knowledge section under the new product contract.
    kwargs.setdefault('knowledge_paragraphs', ['测试替身的知识讲解内容'])
    return _write_product_fixture(*args, **kwargs)

ROOT = Path(__file__).resolve().parents[1]
WEBAPP = ROOT / "v1.2-xml-experiment" / "res" / "app" / "webapp"
sys.path.insert(0, str(WEBAPP))
sys.path.insert(0, str(WEBAPP.parent))
import app as app_module
from app import app
from job_service import JobService


def docx(*texts):
    buffer = io.BytesIO()
    document = Document()
    for text in texts or ("题目",):
        document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def archive(entries):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        for name, data in entries:
            package.writestr(name, data)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path):
    app.config.update(TESTING=True, RESULT_ROOT=tmp_path / "results",
                      RUNTIME_ROOT=tmp_path / "runtime",
                      C0_DISABLE_JOB_SUBMISSION=False, C0_RUN_JOBS_SYNCHRONOUSLY=True,
                      C0_FORCE_FALLBACK_REASON=None, STUDENTIZER_EVIDENCE_PROVIDER=None,
                      OPEN_FOLDER=lambda _folder: None)
    return app.test_client()


def service():
    return app_module._jobs()


def post(client, files, **options):
    data = {"subject": "数学", "grade": "高一", "handout_type": "复习讲义",
            "academic_year": "2026-2027学年", "template_type": "1v1",
            "docx_mode": "auto", "engine_mode": "stable_v09", **options,
            "files": [(io.BytesIO(content), name) for name, content in files]}
    response = client.post("/api/jobs", data=data, content_type="multipart/form-data")
    assert response.status_code == 202, response.get_json()
    return client.get("/api/jobs/" + response.get_json()["job_id"]).get_json()


@pytest.fixture
def renderer(monkeypatch):
    import renderer_orchestrator
    import slot_router
    import template_slot_composer
    calls = {"xml": [], "fallback": [], "selected": [], "make_student": []}
    plan = SimpleNamespace(template_sha256="fixture-template", slots={
        name: () for name in ("knowledge", "immediate", "final")}, units=(),
        slot_labels={"knowledge": "知识精讲", "immediate": "即时训练", "final": "六、巩固练习"},
        explicit_final_heading=True,
        # The test renderer is a route stub; paired training-only plans expose
        # the same deterministic fixture signature so the production pair gate
        # is exercised without pretending to reproduce A-Line classification.
        training_question_routes=((1, "immediate"),))

    def build(source, *args, **_kwargs):
        if args:
            plan.template_type = args[0]
        if "FORCED_UNSUPPORTED" in " ".join(p.text for p in Document(source).paragraphs):
            raise slot_router.SlotRoutingError("UNSUPPORTED_BOOKMARK_SCOPE", "safe forced fixture")
        return plan

    def render(source, _plan, output):
        calls["xml"].append(Path(source).read_bytes())
        write_product_fixture(output, getattr(_plan, "template_type", "1v1"))
        return SimpleNamespace(output_path=str(output), resource_report={"unsupported": []},
                               package_report={"valid": True, "errors": []})

    def fallback(job, reason):
        calls["fallback"].append((job.topic, reason))
        outputs = [job.output_doc]
        write_product_fixture(job.output_doc, job.template_type)
        if job.student_source_doc:
            write_product_fixture(job.student_output_doc, job.template_type)
            outputs.append(job.student_output_doc)
        elif not job.student_only:
            engine.make_student(job.source_doc, job.student_output_doc)
            outputs.append(job.student_output_doc)
        return {"output_paths": outputs, "whole_job": True,
                "baseline_sha": renderer_orchestrator.V09_BASELINE_SHA}

    def selected(job):
        calls["selected"].append(job.topic)
        running_batch = next((record for record in app_module._jobs().list()
                              if record.get("is_batch") and record.get("status") == "running"), None)
        if running_batch:
            parent = app_module._jobs().get(running_batch["job_id"])
            calls.setdefault("batch_progress", []).append(
                (parent.get("current_topic"), parent.get("completed"), parent.get("total")))
        outputs = [job.output_doc]
        write_product_fixture(job.output_doc, job.template_type)
        if not job.student_only:
            if not job.student_source_doc:
                raise AssertionError("stable paired render requires a prepared/supplied student source")
            write_product_fixture(job.student_output_doc, job.template_type)
            outputs.append(job.student_output_doc)
        return {"output_paths": outputs, "fallback_reason": None,
                "selected_engine": "V0.9", "baseline_sha": renderer_orchestrator.V09_BASELINE_SHA}

    def make_student(source, output):
        calls["make_student"].append(source)
        write_product_fixture(output, getattr(plan, "template_type", "1v1"),
                              paragraphs=["已由测试替身去答案的学生版"])

    monkeypatch.setattr(slot_router, "build_slot_routing_plan", build)
    monkeypatch.setattr(template_slot_composer, "render_slots", render)
    monkeypatch.setattr(renderer_orchestrator, "render_v09_whole_job", fallback)
    monkeypatch.setattr(renderer_orchestrator, "render_v09_selected", selected)
    engine = SimpleNamespace(make_student=make_student)
    monkeypatch.setattr(renderer_orchestrator, "_load_v09_engine", lambda: engine)
    return calls


def assert_clean(job):
    folder = Path(job["result_dir"])
    files = list(folder.rglob("*"))
    assert all("_" not in path.name for path in files)
    assert all(path.is_dir() or path.suffix == ".docx" for path in files)
    assert len([path for path in files if path.is_file()]) == len(job["output_paths"])


def test_three_supplied_pairs_are_independent_stable_v09_items(client, renderer):
    teacher, student = docx("教师独有内容"), docx("学生独有内容")
    files = [(f"专题{i} 教师版.docx", teacher) for i in (1, 2, 3)]
    files += [(f"专题{i} 学生版.docx", student) for i in (3, 1, 2)]
    final = post(client, files)
    assert final["status"] == "done" and final["batch_outcome"] == "ALL_SUCCESS"
    assert (final["total"], final["completed"], final["failed"], final["produced"]) == (3, 3, 0, 6)
    assert all(item["renderer"] == "V0.9" and item["renderer_route"] == "STABLE_V09"
               and item["input_version"] == "TEACHER_AND_STUDENT"
               for item in final["items"])
    assert not renderer["fallback"] and not renderer["make_student"]
    assert renderer["selected"] == ["专题1", "专题2", "专题3"]
    assert not renderer["xml"]
    assert all(not item["student_preparation"]["make_student_called"] for item in final["items"])
    assert len(client.get("/api/jobs").get_json()["jobs"]) == 1  # child jobs are internal
    assert_clean(final)


def test_product_integrity_gate_blocks_duplicate_staged_output_before_publication(
        client, renderer, monkeypatch):
    import renderer_orchestrator

    source = docx(*["block-%02d %s" % (index, "长题干内容" * 18) for index in range(14)])

    def duplicate_all_blocks(job):
        source_path, output_path = job.source_doc, job.output_doc
        original = Document(source_path)
        paragraphs = [paragraph.text for paragraph in original.paragraphs]
        write_product_fixture(output_path, "1v1", paragraphs=paragraphs * 2)
        if not job.student_only:
            write_product_fixture(job.student_output_doc, "1v1", paragraphs=paragraphs * 2)
            return {"output_paths": [str(output_path), job.student_output_doc],
                    "selected_engine": "V0.9", "fallback_reason": None}
        return {"output_paths": [str(output_path)], "selected_engine": "V0.9", "fallback_reason": None}

    monkeypatch.setattr(renderer_orchestrator, "render_v09_selected", duplicate_all_blocks)
    result = post(client, [("专题 教师版.docx", source)])
    assert result["status"] == "error"
    assert "PRODUCT_INTEGRITY_GATE" in result["error"]
    assert result["product_integrity"]["accepted"] is False
    assert any(error["reason_code"] == "PRODUCT_REPEATED_BLOCK_SEQUENCE"
               for error in result["product_integrity"]["errors"])
    assert result["output_paths"] == []
    assert list(Path(result["result_dir"]).glob("*.docx")) == []


def test_recursive_generated_input_is_rejected_before_any_renderer(client, renderer):
    sections = ("课堂启动", "知识回顾", "知识精讲", "即时训练", "归纳总结", "巩固练习")
    paragraphs = []
    for _ in range(2):
        paragraphs.append("学科教师辅导讲义")
        paragraphs.extend("数学" + heading for heading in sections)
    result = post(client, [("疑似生成成品 教师版.docx", docx(*paragraphs))])
    assert result["status"] == "error"
    assert "POSSIBLE_GENERATED_OUTPUT_REINGESTION" in result["error"]
    assert result["product_integrity"]["accepted"] is False
    assert not renderer["xml"] and not renderer["fallback"] and not renderer["selected"] and not renderer["make_student"]
    assert result["output_paths"] == []


def _student_job_options():
    return {"subject": "数学", "grade": "高一", "handout_type": "复习讲义",
            "academic_year": "2026-2027学年", "template_type": "1v1",
            "split_mode": "smart", "docx_mode": "auto"}


def test_runtime_input_hash_must_match_upload_provenance(client, renderer):
    source = docx("accepted upload bytes")
    created = service().create_inputs([("专题 学生版.docx", source)], _student_job_options())
    runtime_source = Path(created["student_source_path"])
    runtime_source.write_bytes(docx("mutated runtime bytes"))
    app_module._execute_job(created["job_id"])
    result = service().get(created["job_id"])
    assert result["status"] == "error"
    assert "INPUT_PROVENANCE_HASH_MISMATCH" in result["error"]
    assert result["product_integrity"]["accepted"] is False
    assert not renderer["xml"] and not renderer["fallback"]


def test_existing_publication_target_is_rechecked_by_product_gate(client, renderer):
    sections = ("课堂启动", "知识回顾", "知识精讲", "即时训练", "归纳总结", "巩固练习")
    bad_output = docx(*( ["学科教师辅导讲义"] + ["数学" + heading for heading in sections] ) * 2)
    created = service().create_inputs([("专题 学生版.docx", docx("single clean input"))],
                                      _student_job_options())
    job_id = created["job_id"]
    record = service().get(job_id)
    target = Path(record["result_dir"]) / app_module._visible_output_name(
        record["options"], record["items"][0]["topic"], "学生版")
    target.write_bytes(bad_output)
    bad_sha = hashlib.sha256(bad_output).hexdigest()
    record_path = service()._job_dir(job_id) / "job.json"
    record["publication"] = {
        "role_paths": {"student": str(target.resolve())},
        "expected_sha256": {"student": bad_sha},
        "published": {"student": {"path": str(target.resolve()), "sha256": bad_sha}},
    }
    JobService._write_json(record_path, record)

    app_module._execute_job(job_id)
    final = service().get(job_id)
    assert final["status"] == "error"
    assert "PRODUCT_INTEGRITY_GATE_FINAL" in final["error"]
    assert final["product_integrity"]["accepted"] is False
    assert final["product_integrity"]["final_outputs"][0]["accepted"] is False
    assert not target.exists()


@pytest.mark.parametrize("start_with_bad_target", [False, True])
def test_final_gate_rechecks_hash_before_acceptance_or_cleanup(
        client, renderer, monkeypatch, start_with_bad_target):
    import product_integrity

    clean_source = docx("single clean source")
    bad_template = docx(*(( ["学科教师辅导讲义"] +
                           ["数学" + heading for heading in
                            ("课堂启动", "知识回顾", "知识精讲", "即时训练", "归纳总结", "巩固练习")] ) * 2))
    created = service().create_inputs([("专题 学生版.docx", clean_source)], _student_job_options())
    job_id = created["job_id"]
    record = service().get(job_id)
    target = Path(record["result_dir"]) / app_module._visible_output_name(
        record["options"], record["items"][0]["topic"], "学生版")
    if start_with_bad_target:
        target.write_bytes(bad_template)
        bad_sha = hashlib.sha256(bad_template).hexdigest()
        record["publication"] = {
            "role_paths": {"student": str(target.resolve())},
            "expected_sha256": {"student": bad_sha},
            "published": {"student": {"path": str(target.resolve()), "sha256": bad_sha}},
        }
        JobService._write_json(service()._job_dir(job_id) / "job.json", record)

    replacement = bad_template if not start_with_bad_target else docx("replaced after inspection")
    original_validator = product_integrity.validate_product_integrity
    replaced = {"value": False}

    def replace_final_after_inspection(source_path, output_path, *, plan=None):
        report = original_validator(source_path, output_path, plan=plan)
        if Path(output_path).resolve() == target.resolve() and not replaced["value"]:
            target.write_bytes(replacement)
            replaced["value"] = True
        return report

    monkeypatch.setattr(product_integrity, "validate_product_integrity", replace_final_after_inspection)
    app_module._execute_job(job_id)
    final = service().get(job_id)
    assert replaced["value"] is True
    assert final["status"] == "error"
    assert "PRODUCT_OUTPUT_HASH_CHANGED" in final["error"]
    assert final["product_integrity"]["accepted"] is False
    assert target.is_file() and hashlib.sha256(target.read_bytes()).digest() == hashlib.sha256(replacement).digest()


def test_corrupt_item_does_not_remove_two_successful_outputs(client, renderer):
    final = post(client, [("正常A 学生版.docx", docx()), ("损坏 教师版.docx", b"broken"),
                          ("正常B 学生版.docx", docx())])
    assert final["status"] == "partial" and final["batch_outcome"] == "PARTIAL_SUCCESS"
    assert (final["completed"], final["failed"], len(final["output_paths"])) == (2, 1, 2)
    assert final["has_result"] and "download_available" not in final
    assert all(Path(path).is_file() for path in final["output_paths"])
    assert client.get("/api/open/" + final["job_id"]).status_code == 200
    assert client.get("/api/download/" + final["job_id"]).status_code == 404
    restarted = JobService(app.config["RESULT_ROOT"], runtime_root=app.config["RUNTIME_ROOT"])
    assert restarted.get(final["job_id"])["status"] == "partial"
    assert len(restarted.list()) == 1
    assert_clean(final)


def test_stable_engine_item_failure_keeps_successful_siblings(client, renderer, monkeypatch):
    import renderer_orchestrator

    original = renderer_orchestrator.render_v09_selected

    def fail_one(job):
        if job.topic == "专题4":
            raise RuntimeError("simulated selected V0.9 failure")
        return original(job)

    monkeypatch.setattr(renderer_orchestrator, "render_v09_selected", fail_one)
    final = post(client, [(f"专题{i} 学生版.docx", docx("source %d" % i)) for i in range(1, 5)])
    assert final["status"] == "partial"
    assert [item["status"] for item in final["items"]] == ["done", "done", "done", "error"]
    assert all(item["renderer"] == "V0.9" for item in final["items"][:3])
    failed = final["items"][-1]
    assert "simulated selected V0.9 failure" in failed["error"]
    assert failed["renderer_route"] == "STABLE_V09"
    assert failed["fallback_reason"] is None
    assert not renderer["fallback"]
    assert len(renderer["selected"]) == 3
    assert len(final["output_paths"]) == 3


def test_selected_xml_never_switches_to_v09_when_unsupported(client, renderer):
    final = post(client, [("Unsupported 学生版.docx", docx("FORCED_UNSUPPORTED"))],
                 engine_mode="xml_restricted")
    assert final["status"] == "error"
    assert final["renderer"] == "XML_UNSUPPORTED"
    assert final["fallback_reason"] is None
    assert "手动选择" in final["error"] and "稳定模式" in final["error"]
    assert not renderer["fallback"] and not renderer["selected"]


def test_zip_mixed_input_versions_preserves_user_student(client, renderer):
    student = docx("用户提供的学生内容")
    package = archive([("第一层/专题A 教师用.docx", docx()),
                       ("第一层/第二层/专题B 学生用.docx", student),
                       ("第一层/专题C 解析版.docx", docx()),
                       ("另一层/专题C 原卷版.docx", student),
                       ("第一层/~$专题C 解析版.docx", b"temporary")])
    final = post(client, [("中文混合输入.zip", package)])
    assert final["status"] == "done" and final["total"] == 3
    assert [item["input_version"] for item in final["items"]] == [
        "TEACHER_ONLY", "STUDENT_ONLY", "TEACHER_AND_STUDENT"]
    assert final["produced"] == 5 and len(renderer["make_student"]) == 1
    assert final["items"][0]["student_preparation"]["elapsed_seconds"] >= 0
    assert final["items"][0]["student_preparation"]["make_student_called"] is True
    assert final["items"][0]["student_preparation"]["wps_com_started"] is None
    assert all(not item["student_preparation"]["make_student_called"] for item in final["items"][1:])
    assert not renderer["xml"]
    assert len(renderer["selected"]) == 3
    child = service().get(final["items"][0]["child_job_id"])
    assert child["source_provenance"]["0"]["origin"].startswith(
        "中文混合输入.zip:第一层/专题A 教师用.docx")
    assert_clean(final)


def test_all_bad_items_are_all_failed_without_local_outputs(client, renderer):
    final = post(client, [(f"损坏{i} 教师版.docx", b"broken") for i in range(3)])
    assert final["status"] == "error" and final["batch_outcome"] == "ALL_FAILED"
    assert final["failed"] == 3 and final["completed"] == 0
    assert final["output_paths"] == [] and not final["has_result"]
    assert client.get("/api/download/" + final["job_id"]).status_code == 404
    assert not renderer["xml"] and not renderer["fallback"]


def test_recovery_loses_only_item_with_missing_final_output(client, renderer):
    final = post(client, [(f"专题{i} 学生版.docx", docx()) for i in range(3)])
    Path(final["items"][1]["output_paths"][0]).unlink()
    restarted = JobService(app.config["RESULT_ROOT"], runtime_root=app.config["RUNTIME_ROOT"])
    recovered = restarted.get(final["job_id"])
    assert recovered["status"] == "partial"
    assert recovered["completed"] == 2 and recovered["failed"] == 1
    assert all(Path(path).is_file() for path in recovered["output_paths"])


def test_batch_parent_restart_skips_already_completed_child(client, renderer):
    app.config["C0_DISABLE_JOB_SUBMISSION"] = True
    final = post(client, [(f"专题{i} 学生版.docx", docx()) for i in range(3)])
    parent = service().start_job(final["job_id"])
    app_module._execute_job(parent["items"][0]["child_job_id"])
    # Fresh service marks interrupted parent queued; completed child stays done.
    restarted = JobService(app.config["RESULT_ROOT"], runtime_root=app.config["RUNTIME_ROOT"])
    recovered = restarted.get(final["job_id"])
    assert recovered["status"] == "queued" and recovered["completed"] == 1
    app_module._execute_batch(final["job_id"])
    assert service().get(final["job_id"])["status"] == "done"
    assert len(renderer["selected"]) == 3  # the first child wasn't rendered twice


def test_batch_restart_reruns_incomplete_child_with_fresh_private_paths(client, renderer, monkeypatch):
    """A leftover internal renderer file must not block a resumed child."""
    app.config["C0_DISABLE_JOB_SUBMISSION"] = True
    final = post(client, [("恢复专题%d 教师版.docx" % i, docx("题目%d" % i)) for i in range(1, 4)])
    parent = service().start_job(final["job_id"])
    first_child = parent["items"][0]["child_job_id"]
    interrupted_child = parent["items"][1]["child_job_id"]

    app_module._execute_job(first_child)
    first_outputs = service().get(first_child)["output_paths"]
    first_bytes = [Path(path).read_bytes() for path in first_outputs]
    service().start_job(interrupted_child)  # persist the running state before interruption

    # Simulate an output written before the hard stop, with recovery unable to
    # remove it.  The retry must use another private path, never overwrite it.
    work = Path(service()._job_dir(interrupted_child)) / "work"
    stale = {
        work / ("teacher-output-%s.docx" % interrupted_child): docx("partial teacher"),
        work / ("student-output-%s.docx" % interrupted_child): docx("partial student"),
        work / "derived-student-source.docx": docx("partial student source"),
    }
    for path, content in stale.items():
        path.write_bytes(content)
    monkeypatch.setattr(JobService, "_clear_interrupted_work", staticmethod(lambda *_args: None))

    restarted = JobService(app.config["RESULT_ROOT"], runtime_root=app.config["RUNTIME_ROOT"])
    monkeypatch.setattr(app_module, "_jobs", lambda: restarted)
    recovered = restarted.get(final["job_id"])
    assert recovered["status"] == "queued" and recovered["completed"] == 1
    assert restarted.get(first_child)["status"] == "done"

    app_module._execute_batch(final["job_id"])

    finished = restarted.get(final["job_id"])
    assert finished["status"] == "done" and finished["completed"] == 3
    assert [Path(path).read_bytes() for path in first_outputs] == first_bytes
    assert all(path.is_file() and path.read_bytes() == content for path, content in stale.items())
    assert len(renderer["make_student"]) == 3  # completed child was not prepared again
    assert len(renderer["selected"]) == 3  # completed child wasn't rendered twice
    resumed_work = Path(restarted._job_dir(interrupted_child)) / "work"
    assert (resumed_work / ("teacher-output-attempt-2-%s.docx" % interrupted_child)).is_file()
    assert (resumed_work / ("derived-student-source-attempt-2-%s.docx" % interrupted_child)).is_file()


def test_local_outputs_have_no_name_collisions_for_same_topic_separate_mode(client, renderer):
    final = post(client, [("专题 学生版.docx", docx())] * 3, docx_mode="separate")
    assert final["status"] == "done" and final["total"] == 3
    paths = [Path(path) for path in final["output_paths"]]
    assert len(paths) == len(set(paths)) == 3
    assert len({path.parent for path in paths}) == 3
    assert all(path.is_file() and path.stat().st_size > 0 for path in paths)
    assert_clean(final)


def test_batch_metadata_exposes_current_topic_and_serial_order(client, renderer, monkeypatch):
    final = post(client, [(f"专题{i} 学生版.docx", docx()) for i in range(1, 4)])
    assert final["status"] == "done"
    assert renderer["batch_progress"] == [("专题1", 0, 3), ("专题2", 1, 3), ("专题3", 2, 3)]
    assert final["current_topic"] is None


def test_child_creation_failure_does_not_abort_batch(client, renderer, monkeypatch):
    jobs = service()
    original = jobs.create_inputs
    def fail_one(files, options, **kwargs):
        if kwargs.get("_topic") == "专题B":
            raise OSError("safe fixture write failure")
        return original(files, options, **kwargs)
    monkeypatch.setattr(jobs, "create_inputs", fail_one)
    final = post(client, [(f"专题{letter} 学生版.docx", docx()) for letter in "ABC"])
    assert final["status"] == "partial" and final["completed"] == 2 and final["failed"] == 1
    assert [item["status"] for item in final["items"]] == ["done", "error", "done"]
    assert all(Path(path).is_file() for path in final["output_paths"])


def test_concurrent_first_submissions_share_one_serial_executor(client, monkeypatch):
    """Real async queue, deliberately held during first construction and work."""
    from concurrent.futures import ThreadPoolExecutor

    app.config.update(C0_DISABLE_JOB_SUBMISSION=False, C0_RUN_JOBS_SYNCHRONOUSLY=False)
    # Isolate this test's executor/pending registry and restore the app afterward.
    monkeypatch.setattr(app, "extensions", {})
    barrier = threading.Barrier(9)
    constructor_started = threading.Event()
    second_constructor_started = threading.Event()
    release_constructor = threading.Event()
    first_callback_started = threading.Event()
    release_callbacks = threading.Event()
    guard = threading.Lock()
    executors, observed, failures = [], [], []
    active = {"current": 0, "maximum": 0}
    constructors = {"count": 0}

    def make_executor(**kwargs):
        with guard:
            constructors["count"] += 1
            if constructors["count"] > 1:
                second_constructor_started.set()
        constructor_started.set()
        if not release_constructor.wait(5):
            raise RuntimeError("test constructor timed out")
        executor = ThreadPoolExecutor(**kwargs)
        with guard:
            executors.append(executor)
        return executor

    def callback(job_id):
        with guard:
            active["current"] += 1
            active["maximum"] = max(active["maximum"], active["current"])
        first_callback_started.set()
        try:
            if not release_callbacks.wait(5):
                raise RuntimeError("test callback timed out")
            with guard:
                observed.append(job_id)
        finally:
            with guard:
                active["current"] -= 1

    def submit(index):
        try:
            barrier.wait(timeout=5)
            app_module._submit_job(str(index))
        except Exception as exc:
            with guard:
                failures.append(exc)

    monkeypatch.setattr(app_module, "ThreadPoolExecutor", make_executor)
    monkeypatch.setattr(app_module, "_execute_job", callback)
    threads = [threading.Thread(target=submit, args=(index,)) for index in range(8)]
    try:
        for thread in threads:
            thread.start()
        barrier.wait(timeout=5)
        assert constructor_started.wait(5)
        # Under the old code, another entrant can begin constructing a second
        # executor while the first is deliberately held. The fixed lock blocks
        # all entrants until registration completes.
        assert not second_constructor_started.wait(0.2)
        release_constructor.set()
        for thread in threads:
            thread.join(timeout=5)
            assert not thread.is_alive()
        assert first_callback_started.wait(5)
        with guard:
            assert constructors["count"] == 1 and active["maximum"] == 1
    finally:
        release_constructor.set()
        release_callbacks.set()
        for thread in threads:
            thread.join(timeout=5)
        for executor in executors:
            executor.shutdown(wait=True)
    assert not failures
    assert set(observed) == set(map(str, range(8))) and len(observed) == 8
    assert active["maximum"] == 1 and active["current"] == 0
    assert app.extensions["c0_job_pending"] == set()
