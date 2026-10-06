"""Build structurally valid V1.2 DOCX artifacts for integration test doubles."""
from __future__ import annotations

from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
APP_DIR = ROOT / "v1.2-xml-experiment" / "res" / "app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))


def write_product_fixture(output_path, template_type="1v1", *, source_path=None,
                          paragraphs=None):
    """Write the frozen V1.2 template shape plus optional source text blocks.

    Renderer doubles must return a DOCX that satisfies the production
    normalizer's structural contract. Keeping this helper on the real frozen
    template avoids weakening that contract for integration tests.
    """
    from template_block_plan import resolve_template
    from template_slot_composer import _resolve_content_carrier

    template_path, _digest = resolve_template(template_type)
    document = Document(str(template_path))
    carrier = _resolve_content_carrier(document)
    carrier_table = next((table for table in document.tables if table._tbl is carrier), None)
    if carrier_table is None or len(carrier_table.rows) <= 5 or not carrier_table.columns:
        raise AssertionError("frozen template content carrier is unresolved")

    content = list(paragraphs or ())
    if source_path is not None:
        source = Document(str(source_path))
        content.extend(paragraph.text for paragraph in source.paragraphs if paragraph.text)
    cell = carrier_table.cell(5, 0)
    for text in content:
        cell.add_paragraph(str(text))
    document.save(str(output_path))
    return Path(output_path)
