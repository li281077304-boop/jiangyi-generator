from __future__ import annotations

import base64
import sys
import zipfile
from pathlib import Path
import tempfile
import unittest
import hashlib

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from lxml import etree
from unittest import mock

APP = Path(__file__).resolve().parents[1] / "v1.2-xml-experiment" / "res" / "app"
sys.path.insert(0, str(APP))
from renderer_xml_minimal import (  # noqa: E402
    BlockSpan, ProjectionError, TemplateTarget, render_minimal,
)
import renderer_xml_minimal as renderer_module  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
O = "urn:schemas-microsoft-com:office:office"
DOC = "word/document.xml"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
    "SMlYVQAAAABJRU5ErkJggg==")


def _docx(path: Path, *, table: bool = False, nested: bool = False):
    doc = Document()
    p = doc.add_paragraph("duplicate")
    p.runs[0].bold = True
    if table:
        table_obj = doc.add_table(rows=1, cols=2)
        table_obj.cell(0, 0).text = "left"
        table_obj.cell(0, 1).text = "right"
        if nested:
            inner = table_obj.cell(0, 0).add_table(rows=1, cols=1)
            inner.cell(0, 0).text = "nested"
    p2 = doc.add_paragraph("duplicate")
    p2.runs[0].italic = True
    doc.add_paragraph()  # empty paragraphs retain their StructDoc identity
    doc.save(path)


def _rewrite_document(path, edit):
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    root = etree.fromstring(files[DOC])
    edit(root)
    files[DOC] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, content in files.items():
            zout.writestr(name, content)


def _counts(path):
    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read(DOC))
    return root, len(root.findall(".//{%s}body/{%s}p" % (W, W))), len(root.findall(".//{%s}tbl" % W))


def _add_empty_table_between_first_two_paragraphs(path):
    def edit(root):
        body = root.find("{%s}body" % W)
        table = etree.Element("{%s}tbl" % W)
        etree.SubElement(table, "{%s}tblPr" % W)
        etree.SubElement(table, "{%s}tblGrid" % W)
        body.insert(1, table)
    _rewrite_document(path, edit)


def _add_tracked_insertion(path):
    def edit(root):
        paragraph = root.find(".//{%s}body/{%s}p" % (W, W))
        insertion = etree.SubElement(paragraph, "{%s}ins" % W)
        run = etree.SubElement(insertion, "{%s}r" % W)
        text = etree.SubElement(run, "{%s}t" % W)
        text.text = "inserted"
    _rewrite_document(path, edit)


def _add_cell_revision(path, revision_tag):
    def edit(root):
        cell_properties = root.find(".//{%s}body/{%s}tbl/{%s}tr/{%s}tc/{%s}tcPr" %
                                    (W, W, W, W, W))
        if cell_properties is None:
            cell = root.find(".//{%s}body/{%s}tbl/{%s}tr/{%s}tc" % (W, W, W, W))
            cell_properties = etree.Element("{%s}tcPr" % W)
            cell.insert(0, cell_properties)
        etree.SubElement(cell_properties, "{%s}%s" % (W, revision_tag))
    _rewrite_document(path, edit)


class RendererMinimalTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _paths(self):
        return (self.tmp_path / name for name in ("source.docx", "template.docx", "out.docx"))

    def test_structdoc_node_mapping_uses_structure_for_duplicate_text(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        render_minimal(str(source), str(template), [BlockSpan("b1", "b1")], str(output),
                       TemplateTarget(0))

        root, body_paragraphs, _ = _counts(output)
        body = root.find("{%s}body" % W)
        inserted = body[0]
        self.assertEqual(inserted.tag, "{%s}p" % W)
        self.assertIn("duplicate", "".join(inserted.itertext()))
        self.assertEqual(body_paragraphs, 1)
        self.assertIsNotNone(inserted.find(".//{%s}i" % W))

    def test_whole_table_is_cloned_atomically_at_explicit_body_target(self):
        source, template, output = self._paths()
        _docx(source, table=True)
        template_doc = Document()
        template_doc.add_paragraph("anchor")
        template_doc.add_paragraph("tail")
        template_doc.save(template)

        render_minimal(str(source), str(template), [BlockSpan("b1", "b1")], str(output),
                       TemplateTarget(1))

        root, _, tables = _counts(output)
        body = root.find("{%s}body" % W)
        self.assertEqual(body[0].tag, "{%s}p" % W)
        self.assertEqual(body[1].tag, "{%s}tbl" % W)
        self.assertEqual(body[2].tag, "{%s}p" % W)
        self.assertEqual(tables, 1)
        self.assertIn("left", "".join(body[1].itertext()))
        self.assertIn("right", "".join(body[1].itertext()))

    def test_nested_table_and_empty_cell_paragraph_structure_survive_atomic_clone(self):
        source, template, output = self._paths()
        _docx(source, table=True, nested=True)
        Document().save(template)

        render_minimal(str(source), str(template), [BlockSpan("b1", "b1")], str(output),
                       TemplateTarget(0))

        root, _, tables = _counts(output)
        self.assertEqual(tables, 2)
        cloned_table = root.find(".//{%s}body/{%s}tbl" % (W, W))
        self.assertIn("nested", "".join(cloned_table.itertext()))
        outer_first_cell = cloned_table.find("{%s}tr/{%s}tc" % (W, W))
        empty_cell_paragraphs = [p for p in outer_first_cell.findall("{%s}p" % W)
                                 if not "".join(p.itertext()).strip()]
        self.assertGreaterEqual(len(empty_cell_paragraphs), 1)

    def test_raw_omml_inside_paragraph_is_preserved(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        def add_formula(root):
            paragraph = root.find(".//{%s}body/{%s}p" % (W, W))
            math = etree.SubElement(paragraph, "{%s}oMath" % M)
            run = etree.SubElement(math, "{%s}r" % M)
            text = etree.SubElement(run, "{%s}t" % M)
            text.text = "x"
        _rewrite_document(source, add_formula)

        render_minimal(str(source), str(template), [BlockSpan("b0", "b0")], str(output),
                       TemplateTarget(0))

        root, _, _ = _counts(output)
        self.assertEqual(len(root.findall(".//{%s}oMath" % M)), 1)

    def test_empty_body_paragraph_is_addressable_by_structdoc_id(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        render_minimal(str(source), str(template), [BlockSpan("b2", "b2")], str(output),
                       TemplateTarget(0))

        root, body_paragraphs, _ = _counts(output)
        self.assertEqual(body_paragraphs, 1)
        cloned = root.find(".//{%s}body/{%s}p" % (W, W))
        self.assertEqual("".join(cloned.itertext()).strip(), "")

    def test_external_hyperlink_relationship_is_remapped_and_validated(self):
        source, template, output = self._paths()
        source_doc = Document()
        paragraph = source_doc.add_paragraph("link")
        rid = source_doc.part.relate_to(
            "https://example.test/lesson", RT.HYPERLINK, is_external=True)
        hyperlink = etree.SubElement(paragraph._p, "{%s}hyperlink" % W)
        hyperlink.set(qn("r:id"), rid)
        source_doc.save(source)
        Document().save(template)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                                str(output), TemplateTarget(0))

        generated = Document(str(output))
        imported_link = next(generated.element.body.iter("{%s}hyperlink" % W))
        imported_rid = imported_link.get(qn("r:id"))
        self.assertTrue(generated.part.rels[imported_rid].is_external)
        self.assertEqual(generated.part.rels[imported_rid].target_ref,
                         "https://example.test/lesson")
        self.assertGreaterEqual(result.resource_report["stats"].get("hyperlinks_copied", 0), 1)
        self.assertTrue(result.package_report["valid"], result.package_report)

    def test_balanced_bookmark_and_internal_anchor_are_preserved_and_remapped(self):
        source, template, output = self._paths()
        _docx(source)
        template_doc = Document()
        template_paragraph = template_doc.add_paragraph("template")
        existing_start = etree.Element("{%s}bookmarkStart" % W)
        existing_start.set("{%s}id" % W, "7")
        existing_start.set("{%s}name" % W, "template_anchor")
        template_paragraph._p.insert(0, existing_start)
        existing_end = etree.Element("{%s}bookmarkEnd" % W)
        existing_end.set("{%s}id" % W, "7")
        template_paragraph._p.append(existing_end)
        template_doc.save(template)

        def add_bookmark_anchor(root):
            paragraphs = root.findall(".//{%s}body/{%s}p" % (W, W))
            paragraph = paragraphs[0]
            start = etree.Element("{%s}bookmarkStart" % W)
            start.set("{%s}id" % W, "7")
            start.set("{%s}name" % W, "inside")
            paragraph.insert(0, start)
            end = etree.Element("{%s}bookmarkEnd" % W)
            end.set("{%s}id" % W, "7")
            paragraphs[1].append(end)
            link = etree.SubElement(paragraph, "{%s}hyperlink" % W)
            link.set("{%s}anchor" % W, "inside")
        _rewrite_document(source, add_bookmark_anchor)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b1")], str(output),
                           TemplateTarget(0))
        body = Document(str(output)).element.body
        starts = list(body.iter("{%s}bookmarkStart" % W))
        ends = list(body.iter("{%s}bookmarkEnd" % W))
        imported_start = next(node for node in starts
                              if node.get("{%s}name" % W) == "inside")
        imported_end = next(node for node in ends
                            if node.get("{%s}id" % W) == imported_start.get("{%s}id" % W))
        self.assertNotEqual("7", imported_start.get("{%s}id" % W))
        self.assertEqual(imported_start.get("{%s}id" % W), imported_end.get("{%s}id" % W))
        self.assertEqual("inside", next(body.iter("{%s}hyperlink" % W)).get("{%s}anchor" % W))
        self.assertGreaterEqual(result.resource_report["stats"].get("bookmark_ids_remapped", 0), 1)

    def test_dangling_word_toc_anchor_keeps_text_and_drops_only_broken_link(self):
        source, template, output = self._paths()
        source_doc = Document()
        paragraph = source_doc.add_paragraph()
        hyperlink = etree.SubElement(paragraph._p, "{%s}hyperlink" % W)
        hyperlink.set("{%s}anchor" % W, "_Toc9")
        run = etree.SubElement(hyperlink, "{%s}r" % W)
        text = etree.SubElement(run, "{%s}t" % W)
        text.text = "\u76ee\u5f55\u6807\u9898"
        source_doc.save(source)
        Document().save(template)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                                str(output), TemplateTarget(0))

        generated = Document(str(output))
        imported = next(generated.element.body.iter("{%s}hyperlink" % W))
        self.assertIsNone(imported.get("{%s}anchor" % W))
        self.assertIn("\u76ee\u5f55\u6807\u9898", "".join(imported.itertext()))
        self.assertEqual(1, result.resource_report["stats"]["dangling_toc_anchors_dropped"])

    def test_missing_internal_anchor_preserves_visible_content(self):
        source, template, output = self._paths()
        source_doc = Document()
        paragraph = source_doc.add_paragraph()
        hyperlink = etree.SubElement(paragraph._p, "{%s}hyperlink" % W)
        hyperlink.set("{%s}anchor" % W, "custom_missing_target")
        run = etree.SubElement(hyperlink, "{%s}r" % W)
        etree.SubElement(run, "{%s}t" % W).text = "custom link"
        source_doc.save(source)
        Document().save(template)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                                str(output), TemplateTarget(0))
        link = next(Document(str(output)).element.body.iter("{%s}hyperlink" % W))
        self.assertIsNone(link.get("{%s}anchor" % W))
        self.assertIn("custom link", "".join(link.itertext()))
        self.assertEqual(1, result.resource_report["stats"]["dangling_internal_anchors_dropped"])

    def test_partial_bookmark_pair_and_unresolved_anchor_fail_closed(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        def add_pair_across_paragraphs(root):
            paragraphs = root.findall(".//{%s}body/{%s}p" % (W, W))
            start = etree.Element("{%s}bookmarkStart" % W)
            start.set("{%s}id" % W, "8")
            start.set("{%s}name" % W, "crossing")
            paragraphs[0].insert(0, start)
            end = etree.Element("{%s}bookmarkEnd" % W)
            end.set("{%s}id" % W, "8")
            paragraphs[1].append(end)
        _rewrite_document(source, add_pair_across_paragraphs)

        with self.assertRaisesRegex(ProjectionError, "cuts bookmark pair"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")], str(output),
                           TemplateTarget(0))
        self.assertFalse(output.exists())

        _docx(source)
        def add_unresolved_anchor(root):
            paragraph = root.find(".//{%s}body/{%s}p" % (W, W))
            link = etree.SubElement(paragraph, "{%s}hyperlink" % W)
            link.set("{%s}anchor" % W, "")
        _rewrite_document(source, add_unresolved_anchor)
        with self.assertRaisesRegex(ProjectionError, "anchor is missing or ambiguous"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")], str(output),
                           TemplateTarget(0))
        self.assertFalse(output.exists())

    def test_standalone_body_bookmark_marker_is_dropped_with_explicit_stat(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        def add_standalone_body_end(root):
            body = root.find("{%s}body" % W)
            paragraph = body.find("{%s}p" % W)
            start = etree.Element("{%s}bookmarkStart" % W)
            start.set("{%s}id" % W, "11")
            start.set("{%s}name" % W, "standalone_end")
            paragraph.insert(0, start)
            end = etree.Element("{%s}bookmarkEnd" % W)
            end.set("{%s}id" % W, "11")
            body.insert(1, end)
        _rewrite_document(source, add_standalone_body_end)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                                str(output), TemplateTarget(0))
        self.assertEqual(1, result.resource_report["stats"].get(
            "unmatched_bookmark_starts_dropped", 0))
        self.assertEqual(1, result.resource_report["stats"].get(
            "standalone_body_bookmark_markers_dropped", 0))
        self.assertEqual([], list(Document(str(output)).element.body.iter(
            "{%s}bookmarkStart" % W)))

        output.unlink()
        def add_dependent_link(root):
            paragraph = root.find(".//{%s}body/{%s}p" % (W, W))
            link = etree.SubElement(paragraph, "{%s}hyperlink" % W)
            link.set("{%s}anchor" % W, "standalone_end")
        _rewrite_document(source, add_dependent_link)
        with self.assertRaisesRegex(ProjectionError, "depends on omitted bookmark"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                           str(output), TemplateTarget(0))
        self.assertFalse(output.exists())

    def test_output_cannot_alias_source_or_template(self):
        source, template, _ = self._paths()
        _docx(source)
        Document().save(template)
        before = hashlib.sha256(source.read_bytes()).hexdigest()

        with self.assertRaisesRegex(ProjectionError, "must not alias"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")], str(source),
                           TemplateTarget(0))
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_source_is_read_once_for_both_structdoc_and_xml_parsing(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        with mock.patch.object(renderer_module, "_read_source_bytes",
                               wraps=renderer_module._read_source_bytes) as read_source:
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")], str(output),
                           TemplateTarget(0))
        read_source.assert_called_once()

    def test_zero_paragraph_table_between_endpoints_is_not_crossed(self):
        source, template, output = self._paths()
        _docx(source)
        _add_empty_table_between_first_two_paragraphs(source)
        Document().save(template)

        with self.assertRaisesRegex(ProjectionError, "crosses a top-level table"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b2")], str(output),
                           TemplateTarget(0))

    def test_nested_table_node_cannot_be_selected_as_atomic_payload(self):
        source, template, output = self._paths()
        _docx(source, table=True, nested=True)
        Document().save(template)

        with self.assertRaisesRegex(ProjectionError, "only a top-level bN table"):
            render_minimal(str(source), str(template),
                           [BlockSpan("b1.r0c0.n1", "b1.r0c0.n1")], str(output),
                           TemplateTarget(0))

    def test_tracked_change_fails_closed(self):
        source, template, output = self._paths()
        _docx(source)
        _add_tracked_insertion(source)
        Document().save(template)

        with self.assertRaisesRegex(ProjectionError, "tracked-change construct: ins"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")], str(output),
                           TemplateTarget(0))

    def test_cell_insertion_revision_fails_closed_in_table_fixture(self):
        source, template, output = self._paths()
        _docx(source, table=True)
        _add_cell_revision(source, "cellIns")
        Document().save(template)

        with self.assertRaisesRegex(ProjectionError, "tracked-change construct: cellIns"):
            render_minimal(str(source), str(template), [BlockSpan("b1", "b1")], str(output),
                           TemplateTarget(0))
        self.assertFalse(output.exists())

    def test_revision_tag_families_fail_closed_table_driven(self):
        tags = ("cellDel", "cellMerge", "numberingChange", "tblGridChange",
                "tblPrExChange", "conflictIns", "conflictDel")
        for tag in tags:
            with self.subTest(tag=tag):
                payload = etree.Element("{%s}tbl" % W)
                etree.SubElement(payload, "{%s}%s" % (W, tag))
                with self.assertRaisesRegex(ProjectionError,
                                            "tracked-change construct: %s" % tag):
                    renderer_module._validate_payload(payload)

    def test_image_relationship_is_remapped_and_saved_package_validates(self):
        source, template, output = self._paths()
        image_path = self.tmp_path / "pixel.png"
        image_path.write_bytes(PNG)
        source_doc = Document()
        source_doc.add_paragraph("source image").add_run().add_picture(str(image_path))
        source_doc.add_paragraph("source tail")
        source_doc.save(source)
        template_doc = Document()
        template_doc.add_paragraph("before")
        template_doc.add_paragraph("after")
        template_doc.save(template)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                                str(output), TemplateTarget(1))

        self.assertGreaterEqual(result.resource_report["stats"].get("image_relationships_copied", 0), 1)
        self.assertTrue(result.package_report["valid"], result.package_report)
        generated = Document(str(output))
        self.assertEqual([p.text for p in generated.paragraphs], ["before", "source image", "after"])
        with zipfile.ZipFile(output) as package:
            media_files = [name for name in package.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(media_files), 1)
            self.assertEqual(package.read(media_files[0]), PNG)

    def test_ole_embedding_is_copied_and_validated(self):
        source, template, output = self._paths()
        embedding_payload = b"dummy-ole-package-fixture"
        source_doc = Document()
        paragraph = source_doc.add_paragraph("OLE source")
        embedding = Part(PackURI("/word/embeddings/oleObject1.bin"),
                         "application/vnd.openxmlformats-officedocument.oleObject",
                         embedding_payload, source_doc.part.package)
        rel_id = source_doc.part.relate_to(
            embedding,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject")
        object_element = etree.SubElement(paragraph._p, "{%s}object" % W)
        ole_element = etree.SubElement(object_element, "{%s}OLEObject" % O)
        ole_element.set(qn("r:id"), rel_id)
        source_doc.save(source)
        Document().save(template)

        result = render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                                str(output), TemplateTarget(0))

        self.assertGreaterEqual(result.resource_report["stats"].get("ole_references_copied", 0), 1)
        self.assertTrue(result.package_report["valid"], result.package_report)
        with zipfile.ZipFile(output) as package:
            embedding_paths = [name for name in package.namelist()
                               if name.startswith("word/embeddings/") and not name.endswith(".rels")]
            self.assertEqual(len(embedding_paths), 1)
            self.assertEqual(package.read(embedding_paths[0]), embedding_payload)

    def test_unresolved_source_relationship_fails_closed_without_output(self):
        source, template, output = self._paths()
        _docx(source)
        Document().save(template)

        def add_missing_relation(root):
            paragraph = root.find(".//{%s}body/{%s}p" % (W, W))
            link = etree.SubElement(paragraph, "{%s}hyperlink" % W)
            link.set(qn("r:id"), "rId404")
        _rewrite_document(source, add_missing_relation)

        with self.assertRaisesRegex(ProjectionError, "missing_relationship"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b0")],
                           str(output), TemplateTarget(0))
        self.assertFalse(output.exists())

    def test_paragraph_span_crossing_table_fails_closed(self):
        source, template, output = self._paths()
        _docx(source, table=True)
        Document().save(template)

        with self.assertRaisesRegex(ProjectionError, "crosses .*table"):
            render_minimal(str(source), str(template), [BlockSpan("b0", "b2")], str(output),
                           TemplateTarget(0))
