import hashlib
import sys
from pathlib import Path
from collections import namedtuple
from copy import deepcopy
import pytest
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'v1.2-xml-experiment/res/app/webapp'), str(ROOT/'v1.2-xml-experiment/res/app')]
from xml_source_projection import project_equivalent_textboxes, MC, W, W14
from disk_preflight import check_disk_space, DiskSpaceError
import disk_preflight


def source(path, differing=False):
    d=Document();p=d.add_paragraph()._p
    ac=etree.SubElement(p,MC+'AlternateContent')
    c=etree.SubElement(ac,MC+'Choice');c.set('Requires','wps')
    box=etree.SubElement(c,W+'txbxContent');paragraph=etree.SubElement(box,W+'p')
    r=etree.SubElement(paragraph,W+'r');etree.SubElement(r,W+'t').text='提醒正文必须完整保留，不能把兼容分支当成另一份内容。'*3
    fb=etree.SubElement(ac,MC+'Fallback');other=deepcopy(box);other[0].set(W14+'paraId','12345678')
    if differing:other[0][0][0].text='另一份不可删除的正文'
    fb.append(other);d.save(path)
    return path


def test_equivalent_branches_preserve_original_and_package(tmp_path):
    p=source(tmp_path/'original.docx');before=p.read_bytes()
    out,e=project_equivalent_textboxes(p,tmp_path/'projected.docx')
    assert e['status']=='PROJECTED' and e['source_sha256']==hashlib.sha256(before).hexdigest()
    assert p.read_bytes()==before
    d=Document(out);assert len(list(d.element.iter(W+'txbxContent')))==1
    again,e2=project_equivalent_textboxes(out,tmp_path/'again.docx')
    assert e2['status']=='UNCHANGED' and again==out


def test_different_content_cannot_be_discarded(tmp_path):
    p=source(tmp_path/'original.docx',True)
    out,e=project_equivalent_textboxes(p,tmp_path/'projected.docx')
    assert e['status']=='UNCHANGED' and out==p


def test_disk_check_is_read_only_and_reports_all_locations(tmp_path,monkeypatch):
    usage=namedtuple('Usage','total used free')
    monkeypatch.setattr(disk_preflight.shutil,'disk_usage',lambda _:usage(100,99,1))
    with pytest.raises(DiskSpaceError) as e:
        check_disk_space(tmp_path/'runtime',tmp_path/'results',10,temporary_root=tmp_path/'temp')
    assert '磁盘空间不足' in str(e.value)
    assert not (tmp_path/'runtime').exists()
    assert {x for d in e.value.details for x in d['locations']}=={'任务目录','临时目录','成品目录'}


def test_http_space_refusal_happens_before_multipart_parse(monkeypatch):
    from app import app
    usage=namedtuple('Usage','total used free')
    monkeypatch.setattr(disk_preflight.shutil,'disk_usage',lambda _:usage(100,99,1))
    r=app.test_client().post('/api/jobs',data=b'bad multipart',content_type='multipart/form-data; boundary=x')
    assert r.status_code==507
    assert r.get_json()['error_code']=='INSUFFICIENT_DISK_SPACE'
