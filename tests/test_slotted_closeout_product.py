"""Final document slot evidence and empty-module contracts for the closeout."""
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'v1.2-xml-experiment/res/app'))
from exercise_partition import partition_exercises
from product_normalizer import inspect_delivered_slots, normalize_product_docx
from semantic_facade import analyze_source
from template_slot_composer import omit_empty_knowledge, render_slots, _paragraph_text


@pytest.mark.parametrize('template', ['1v1', 'class'])
def test_training_only_final_sections_and_continuous_headings(tmp_path, template):
    source = tmp_path / 'source.docx'
    doc = Document()
    for i in range(1, 21):
        doc.add_paragraph(f'{i}. Independent exercise number {i}?')
    doc.save(source)
    plan, _ = partition_exercises(source, template, analyze_source(source))
    rendered = tmp_path / 'rendered.docx'
    render_slots(str(source), plan, str(rendered), cover_metadata={'topic': '训练'})
    final = tmp_path / 'final.docx'
    normalize_product_docx(rendered, final, template_type=template, metadata={'topic': '训练'})
    output = Document(final)
    texts = [_paragraph_text(p) for p in output.element.body.iter(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')]
    assert not any(t in ('知识精讲', '知识精讲&例题讲解') for t in texts)
    assert '三、即时训练' in texts
    assert '四、归纳总结' in texts
    assert ('五、巩固练习' if template == '1v1' else '五、出门测试') in texts
    assert inspect_delivered_slots(final, template)['verified']
    before = output.element.xml
    omit_empty_knowledge(output, template)
    assert output.element.xml == before


def test_duplicate_slot_heading_is_not_verified(tmp_path):
    doc = Document()
    for text in ['即时训练', '1. exercise', '归纳总结', '六、巩固练习',
                 '2. exercise', '即时训练']:
        doc.add_paragraph(text)
    path = tmp_path / 'ambiguous.docx'
    doc.save(path)
    evidence = inspect_delivered_slots(path, '1v1')
    assert not evidence['verified']
    assert evidence['sections']['immediate']['reason'] == 'ANCHOR_NOT_UNIQUE'


@pytest.mark.parametrize('template', ['1v1', 'class'])
def test_blank_template_and_ornaments_are_not_slotted(tmp_path, template):
    from template_block_plan import resolve_template
    assert not inspect_delivered_slots(resolve_template(template)[0], template)['verified']


def test_empty_knowledge_with_media_is_never_removed():
    from docx.oxml import OxmlElement
    doc = Document()
    doc.add_paragraph('知识精讲')
    p = doc.add_paragraph()
    p._p.append(OxmlElement('w:drawing'))
    doc.add_paragraph('即时训练')
    assert omit_empty_knowledge(doc, '1v1')['status'] == 'RETAINED'


def test_image_only_slots_require_source_media_provenance(tmp_path):
    from PIL import Image
    image = tmp_path / 'question.png'
    Image.new('RGB', (20, 20), 'green').save(image)
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_picture(str(image))
    doc.save(source)
    final = tmp_path / 'final.docx'
    output = Document()
    output.add_paragraph('即时训练')
    output.add_picture(str(image))
    output.add_paragraph('五、归纳总结')
    output.add_paragraph('六、巩固练习')
    output.add_picture(str(image))
    output.save(final)
    assert not inspect_delivered_slots(final, '1v1')['verified']
    assert inspect_delivered_slots(final, '1v1', source_path=source)['verified']


def test_product_package_excludes_ocr_and_production_does_not_call_it():
    root = Path(__file__).resolve().parents[1]
    spec = (root / 'packaging/windows/v1.2_onedir.spec').read_text(encoding='utf-8')
    assert 'ocr_models' not in spec
    for package in ('rapidocr_onnxruntime', 'onnxruntime', 'cv2', 'numpy', 'shapely'):
        assert '"' + package + '"' in spec.split('excludes=')[1]
    app = (root / 'v1.2-xml-experiment/res/app/webapp/app.py').read_text(encoding='utf-8')
    assert 'from local_ocr' not in app
    assert 'analyze_source_with_ocr(' not in app
