"""Synthetic reviewed OOXML capability tests; no real sample/COM UAT claim."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import base64
import sys
import zipfile

from docx import Document
from docx.shared import RGBColor
from lxml import etree
import pytest

APP = Path(__file__).resolve().parents[1] / "v1.2-xml-experiment/res/app"
sys.path.insert(0, str(APP))
from studentizer import (AnswerBinding, StudentizerSemantics, StudentizerPolicy,
                         prepare_student, element_sha256, TAG, FALLBACK)

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")


def package(image=False):
    doc = Document()
    doc.add_paragraph("知识点：答案二字是这里的讲解，不要删除").runs[0].font.color.rgb = RGBColor(255, 0, 0)
    p = doc.add_paragraph("1. 求2+3，保留全部题干")
    if image:
        p.add_run().add_picture(BytesIO(PNG))
    doc.add_paragraph("【答案】5")
    doc.add_paragraph("下一段知识解释：不能顺着答案删除")
    doc.add_paragraph("2. 后续完整题目")
    doc.add_paragraph("【答案】不是批准删除的第二题答案")
    doc.add_paragraph("重复内容")
    doc.add_paragraph("重复内容")
    data = BytesIO(); doc.save(data)
    return data.getvalue()


def parts(data):
    with zipfile.ZipFile(BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def root(data):
    return etree.fromstring(parts(data)["word/document.xml"])


def body(data):
    return root(data).find(TAG("body"))


def rewrite(data, fn):
    entries = parts(data)
    tree = root(data); fn(tree.find(TAG("body")))
    entries["word/document.xml"] = etree.tostring(tree)
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n, value in entries.items():
            z.writestr(n, value)
    return out.getvalue()


def approved(data, **changes):
    binding = AnswerBinding("fixture-q1", "question_group", 1, 2, 2,
                            element_sha256(body(data)[2]), "reviewed_question_answer",
                            "synthetic-reviewed-fixture-q1")
    return StudentizerSemantics(sha256(data).hexdigest(), (replace(binding, **changes),),
                               protected_indices=(0, 1, 3, 4, 5, 6, 7))


def run(data, tmp_path, semantics=None, policy=None):
    source = tmp_path / "教师源.docx"
    output = tmp_path / "新学生源.docx"
    source.write_bytes(data)
    result = prepare_student(source, semantics, output, policy or StudentizerPolicy())
    assert source.read_bytes() == data
    assert result.com_invocations == 0
    return result, output


def assert_closed(result, output, code):
    assert result.status == FALLBACK and result.reason_code == code
    assert result.output_path is None and not output.exists()
    assert list(output.parent.glob('tmp*.docx')) == []


@pytest.mark.parametrize("image", [False, True])
def test_teacher_to_student_removes_only_reviewed_answer_preserves_objects_order_and_parts(tmp_path, image):
    data = package(image)
    result, output = run(data, tmp_path, approved(data))
    assert result.status == "XML_PREPARED" and result.validation["valid"]
    generated = output.read_bytes()
    assert result.source_sha256 == sha256(data).hexdigest()
    assert result.output_sha256 == sha256(generated).hexdigest()
    assert [element_sha256(n) for n in body(generated)] == [element_sha256(n) for i,n in enumerate(body(data)) if i != 2]
    assert "下一段知识解释" in ''.join(body(generated).itertext())
    assert "不是批准删除的第二题答案" in ''.join(body(generated).itertext())
    assert ''.join(body(generated).itertext()).count("重复内容") == 2
    assert result.mutations[0]["body_index"] == 2
    for name, value in parts(data).items():
        if name != "word/document.xml":
            assert parts(generated)[name] == value
    assert len(body(generated).xpath('.//w:drawing', namespaces=NS)) == int(image)


@pytest.mark.parametrize("version", ["STUDENT_ONLY", "TEACHER_AND_STUDENT"])
def test_supplied_student_bypasses_color_role_and_removal_byte_for_byte(tmp_path, version):
    data = package(image=True)
    result, output = run(data, tmp_path, None, StudentizerPolicy(input_version=version, source_role="student"))
    assert result.status == "STUDENT_SOURCE_BYPASS" and output.read_bytes() == data
    assert result.mutations == () and result.engine == "BYPASS"


@pytest.mark.parametrize("changes", [dict(owner_kind="section"), dict(relation="answer_role"), dict(review_id=""), dict(owner_start=2), dict(owner_end=4)])
def test_section_only_role_only_unreviewed_or_crossing_owner_refused(tmp_path, changes):
    data = package()
    result, output = run(data, tmp_path, approved(data, **changes))
    assert_closed(result, output, "STUDENTIZER_ANSWER_OWNERSHIP_UNPROVEN")


def test_red_and_answer_markers_without_reviewed_binding_do_not_authorize_deletion(tmp_path):
    data = package()
    result, output = run(data, tmp_path, StudentizerSemantics(sha256(data).hexdigest()))
    assert_closed(result, output, "STUDENTIZER_ANSWER_OWNERSHIP_UNPROVEN")


@pytest.mark.parametrize("which", ["source", "address"])
def test_identity_or_exact_fingerprint_mismatch_refused(tmp_path, which):
    data = package(); plan = approved(data)
    if which == "source": plan = replace(plan, source_sha256="0"*64)
    else: plan = replace(plan, bindings=(replace(plan.bindings[0], answer_sha256="0"*64),))
    result, output = run(data, tmp_path, plan)
    assert_closed(result, output, "STUDENTIZER_SOURCE_IDENTITY_MISMATCH")


def test_shared_material_protected_overlap_and_duplicate_removal_refused(tmp_path):
    data = package(); plan = approved(data)
    for name, revised in [("material", replace(plan, protected_indices=(2,))),
                          ("duplicate", replace(plan, bindings=plan.bindings*2))]:
        folder = tmp_path/name; folder.mkdir()
        result, output = run(data, folder, revised)
        assert_closed(result, output, "STUDENTIZER_PRESERVATION_CHECK_FAILED")


def test_table_cell_answer_cannot_be_removed_as_plain_body_paragraph(tmp_path):
    data = rewrite(package(), lambda b: b.replace(b[2], etree.fromstring('<w:tbl xmlns:w="'+NS['w']+'"><w:tr><w:tc><w:p><w:r><w:t>【答案】5</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')))
    result, output = run(data, tmp_path, approved(data))
    assert_closed(result, output, "STUDENTIZER_TABLE_SCOPE_UNSUPPORTED")


@pytest.mark.parametrize("xml,code", [
    ('<w:r><w:drawing/></w:r>', "STUDENTIZER_FORMULA_OBJECT_UNSUPPORTED"),
    ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"/>', "STUDENTIZER_FORMULA_OBJECT_UNSUPPORTED"),
    ('<w:r><w:pict><w:txbxContent><w:p/></w:txbxContent></w:pict></w:r>', "STUDENTIZER_TEXTBOX_SCOPE_UNSUPPORTED"),
    ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>', "STUDENTIZER_FIELD_SCOPE_UNSUPPORTED"),
    ('<w:bookmarkStart w:id="1" w:name="range"/>', "STUDENTIZER_BOOKMARK_SCOPE_UNSUPPORTED"),
    ('<w:ins w:id="1"><w:r><w:t>revision</w:t></w:r></w:ins>', "STUDENTIZER_REVISION_SCOPE_UNSUPPORTED"),
    ('<w:pPr><w:sectPr/></w:pPr>', "STUDENTIZER_MIXED_CONTENT_UNSUPPORTED"),
])
def test_unsupported_answer_boundary_constructs_fail_closed(tmp_path, xml, code):
    node = etree.fromstring('<x xmlns:w="'+NS['w']+'">'+xml+'</x>')[0]
    data = rewrite(package(), lambda b: b[2].append(node))
    result, output = run(data, tmp_path, approved(data))
    assert_closed(result, output, code)


def test_inline_prompt_answer_mixture_refused_even_if_reviewed_role_is_supplied(tmp_path):
    data = rewrite(package(), lambda b: setattr(b[2].find('.//'+TAG('t')), 'text', '题干前缀 【答案】5'))
    result, output = run(data, tmp_path, approved(data))
    assert_closed(result, output, "STUDENTIZER_MIXED_CONTENT_UNSUPPORTED")


def test_preserves_red_knowledge_and_unreviewed_answers_no_blank_insertion(tmp_path):
    data = package(); result, output = run(data, tmp_path, approved(data))
    assert result.status == 'XML_PREPARED'
    before = list(body(data)); after = list(body(output.read_bytes()))
    assert element_sha256(before[0]) == element_sha256(after[0])
    assert len(after) == len(before)-1


def test_source_alias_or_existing_output_is_never_overwritten(tmp_path):
    data = package(); source = tmp_path/'source.docx'; source.write_bytes(data)
    result = prepare_student(source, approved(data), source)
    assert result.status == FALLBACK and source.read_bytes() == data
    output=tmp_path/'existing.docx'; output.write_bytes(b'prior result')
    result=prepare_student(source,approved(data),output)
    assert result.status == FALLBACK and output.read_bytes() == b'prior result'


def test_validation_failure_does_not_publish_partial_student_file(tmp_path):
    data=package(image=True); entries=parts(data)
    names=[n for n in entries if n.startswith('word/media/')]
    for n in names: del entries[n]
    buffer=BytesIO()
    with zipfile.ZipFile(buffer,'w') as z:
        for n,v in entries.items():z.writestr(n,v)
    data=buffer.getvalue();result,output=run(data,tmp_path,approved(data))
    assert_closed(result,output,'STUDENTIZER_PACKAGE_VALIDATION_FAILED')


def test_missing_source_returns_explicit_failure_without_output(tmp_path):
    output=tmp_path/'output.docx'
    result=prepare_student(tmp_path/'missing.docx',None,output)
    assert_closed(result,output,'STUDENTIZER_XML_PREPARATION_FAILED')


@pytest.mark.parametrize('version', ['STUDENT_ONLY', 'TEACHER_AND_STUDENT'])
def test_student_bypass_requires_explicit_student_source_side(tmp_path, version):
    data=package();result,output=run(data,tmp_path,approved(data),StudentizerPolicy(input_version=version))
    assert_closed(result,output,'STUDENTIZER_ANSWER_OWNERSHIP_UNPROVEN')


def test_source_with_unknown_frozen_semantics_is_not_used_as_deletion_authority(tmp_path):
    from types import SimpleNamespace
    data=package();predictions=SimpleNamespace(source_sha256=sha256(data).hexdigest(),
                     units=[{'role':'answer','parent':'section1','spans':[['b2','b2']]}])
    result,output=run(data,tmp_path,predictions)
    assert_closed(result,output,'STUDENTIZER_ANSWER_OWNERSHIP_UNPROVEN')
