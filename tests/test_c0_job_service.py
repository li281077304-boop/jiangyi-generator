import importlib.util
import io
import json
from pathlib import Path
from types import SimpleNamespace
import sys
import zipfile

import pytest
from docx import Document
from product_fixture_utils import write_product_fixture


ROOT = Path(__file__).resolve().parents[1]
WEBAPP = ROOT / "v1.2-xml-experiment" / "res" / "app" / "webapp"
APP_DIR = WEBAPP.parent
INCIDENT_FIXTURE = (ROOT / "docs" / "v2" / "integration" / "fixtures" /
                    "incident-72cab43961d6488c91b7a9e049265263")
sys.path.insert(0, str(WEBAPP))
sys.path.insert(0, str(APP_DIR))

from app import app  # noqa: E402
import app as app_module  # noqa: E402
from job_service import JobService  # noqa: E402


def make_docx() -> bytes:
    buf = io.BytesIO()
    document = Document()
    document.add_paragraph("C0 validator fixture")
    document.add_paragraph("【答案】A")
    document.add_paragraph("【解析】示例解析")
    document.save(buf)
    return buf.getvalue()


def make_student_docx() -> bytes:
    buf = io.BytesIO()
    document = Document()
    document.add_paragraph("C1 student fixture")
    document.save(buf)
    return buf.getvalue()


def corrupt_document_xml(path: Path) -> None:
    replacement = path.with_suffix(".broken.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        for info in source.infolist():
            target.writestr(info, b"<w:document" if info.filename == "word/document.xml"
                            else source.read(info.filename))
    replacement.replace(path)


@pytest.fixture
def client(tmp_path):
    app.config.update(TESTING=True, RESULT_ROOT=tmp_path / "results",
                      RUNTIME_ROOT=tmp_path / "runtime-jobs",
                      C0_DISABLE_JOB_SUBMISSION=True,
                      C0_RUN_JOBS_SYNCHRONOUSLY=False,
                      C0_FORCE_FALLBACK_REASON=None, STUDENTIZER_EVIDENCE_PROVIDER=None,
                      V12_ALLOW_DIAGNOSTIC_V09=False)
    app.config.pop("OPEN_FOLDER", None)
    return app.test_client()


def post_one(client, name="课件.docx", content=None):
    form = {key: value for key, value in {
        "subject": "数学", "grade": "七年级", "handout_type": "课时讲义",
        "academic_year": "2026", "template_type": "class",
        "split_mode": "smart", "docx_mode": "auto",
    }.items()}
    form["files"] = (io.BytesIO(content or make_docx()), name)
    return client.post("/api/jobs", data=form, content_type="multipart/form-data")


def make_service(root=None):
    return JobService(root or Path(app.config["RESULT_ROOT"]),
                      runtime_root=Path(app.config["RUNTIME_ROOT"]))


def test_post_persists_single_docx_and_get_recovers_after_service_restart(client, tmp_path):
    source_bytes = make_docx()
    response = post_one(client, content=source_bytes)
    assert response.status_code == 202
    created = response.get_json()
    assert created["status"] == "queued"
    assert isinstance(created["created_at"], (int, float))
    assert isinstance(created["updated_at"], (int, float))
    assert created["created_at_iso"].endswith("+00:00")
    assert created["total"] == 2
    assert created["has_result"] is False
    assert "download_available" not in created
    root = tmp_path / "results"
    record_path = Path(app.config["RUNTIME_ROOT"]) / created["job_id"] / "job.json"
    record = json.loads(record_path.read_text(encoding="utf-8"))
    assert Path(record["source_path"]).read_bytes() == source_bytes
    provenance = record["source_provenance"]["0"]
    assert provenance["origin"] == "direct_upload:课件.docx"
    assert provenance["sha256"]
    assert Path(provenance["runtime_path"]).read_bytes() == source_bytes
    assert Path(record["result_dir"]).is_relative_to(root)
    assert record["options"]["template_type"] == "class"
    assert record["options"]["handoutType"] == "课时讲义"

    # A fresh service instance simulates process restart; it reads disk state.
    restarted = make_service(root)
    assert restarted.get(created["job_id"])["status"] == "queued"
    assert [item["job_id"] for item in restarted.list()] == [created["job_id"]]
    assert client.get("/api/jobs/" + created["job_id"]).get_json()["status"] == "queued"


def test_result_folder_uses_normalized_topic_and_preserves_original_topic_provenance(client):
    long_name = ("专题12.4 一次函数的实际应用（高效培优讲义）"
                 "数学新教材沪科版八年级上册 期末复习（教师版）.docx")
    created = post_one(client, long_name, make_docx()).get_json()
    assert created["items"][0]["topic"] == "专题12.4 一次函数的实际应用"
    assert created["items"][0]["original_topic"].startswith("专题12.4 一次函数的实际应用")
    assert created["topic_provenance"] == {
        "original": created["items"][0]["original_topic"],
        "display": "专题12.4 一次函数的实际应用",
        "normalizer": "V1.2_TOPIC_NORMALIZATION_V1",
    }
    assert Path(created["result_dir"]).name.endswith("专题12.4 一次函数的实际应用 课时讲义")


def test_upload_accepts_one_explicit_pair_and_rejects_unsupported_inputs(client):
    one = make_docx()
    multi = client.post("/api/jobs", data={"files": [(io.BytesIO(one), "a教师版.docx"),
                                                        (io.BytesIO(make_student_docx()), "a学生版.docx")]},
                        content_type="multipart/form-data")
    assert multi.status_code == 202
    assert multi.get_json()["input_version"] == "TEACHER_AND_STUDENT"
    assert post_one(client, "source.zip", one).status_code == 415
    assert post_one(client, "source.docx", b"not a package").status_code == 415
    three = client.post("/api/jobs", data={"files": [(io.BytesIO(one), "a教师版.docx"),
                                                        (io.BytesIO(make_student_docx()), "a学生版.docx"),
                                                        (io.BytesIO(one), "b.docx")]},
                        content_type="multipart/form-data")
    assert three.status_code == 202
    assert three.get_json()["is_batch"] is True
    assert three.get_json()["total"] == 2


def test_open_route_resolves_and_opens_result_folder(client, tmp_path):
    opened = []
    app.config["OPEN_FOLDER"] = lambda folder: opened.append(folder)
    created = post_one(client).get_json()
    response = client.get("/api/open/" + created["job_id"])
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["result_dir"] == created["result_dir"]
    assert opened == [Path(payload["result_dir"])]
    assert client.get("/api/open/" + "0" * 32).status_code == 404


def test_output_download_route_is_removed_and_does_not_change_generation_state(client):
    created = post_one(client).get_json()
    response = client.get("/api/download/" + created["job_id"])
    assert response.status_code == 404
    assert client.get("/api/jobs/" + created["job_id"]).get_json()["status"] == "queued"


def test_complete_job_uses_frozen_package_validator_and_local_role_outputs(client):
    created = post_one(client).get_json()
    job_dir = Path(created["result_dir"])
    outputs = [job_dir / "教师版.docx", job_dir / "学生版.docx"]
    for path in outputs:
        path.write_bytes(make_docx())
    completed = make_service().complete_job(created["job_id"], outputs[0], outputs[1])
    assert completed["status"] == "done"
    assert set(completed["package_validation"]) == {"teacher", "student"}
    assert all(report["valid"] for report in completed["package_validation"].values())
    assert completed["output_roles"] == {"teacher": str(outputs[0]), "student": str(outputs[1])}
    assert all(path.is_file() and path.stat().st_size > 0 for path in outputs)
    snapshot = client.get("/api/jobs/" + created["job_id"]).get_json()
    assert snapshot["status"] == "done"
    assert "download_available" not in snapshot
    restarted = make_service()
    recovered = restarted.get(created["job_id"])
    assert recovered["result_dir"] == str(job_dir)
    assert recovered["output_paths"] == [str(path) for path in outputs]
    assert "download_available" not in recovered
    assert recovered["created_at"] == created["created_at"]


@pytest.mark.parametrize("batch", [False, True])
def test_legacy_download_metadata_is_removed_without_losing_local_outputs(client, batch):
    if batch:
        created = client.post("/api/jobs", data={"files": [(io.BytesIO(make_student_docx()), f"专题{i} 学生版.docx") for i in range(3)]}, content_type="multipart/form-data").get_json()
        child_id = created["items"][0]["child_job_id"]
        child = make_service().get(child_id)
        output = Path(child["result_dir"]) / "学生版.docx"
        output.write_bytes(make_student_docx())
        make_service().complete_job(child_id, student_output_path=output)
    else:
        created = post_one(client).get_json()
        output = Path(created["result_dir"]) / "教师版.docx"
        output.write_bytes(make_docx())
        make_service().complete_job(created["job_id"], output)
    metadata = Path(app.config["RUNTIME_ROOT"]) / created["job_id"] / "job.json"
    record = json.loads(metadata.read_text(encoding="utf-8"))
    legacy = {"download_available": False, "delivery_status": "failed", "delivery_error_code": "DELIVERY_DOWNLOAD_FAILED", "delivery_error": "old failure"}
    record.update(legacy)
    metadata.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    recovered = make_service().get(created["job_id"])
    assert recovered["has_result"] and output.is_file()
    assert not set(legacy).intersection(recovered)
    assert not set(legacy).intersection(json.loads(metadata.read_text(encoding="utf-8")))
    assert client.get("/api/download/" + created["job_id"]).status_code == 404


def test_pair_alignment_failure_uses_v09_only_in_explicit_diagnostic_mode(
        client, tmp_path, monkeypatch):
    import renderer_orchestrator

    root = tmp_path / "results"
    app.config.update(C0_DISABLE_JOB_SUBMISSION=False,
                      C0_RUN_JOBS_SYNCHRONOUSLY=True,
                      V12_ALLOW_DIAGNOSTIC_V09=True)
    fallback_observed = {}

    def successful_fallback(job, reason):
        record = make_service(root).list()[0]
        fallback_observed.update({
            "renderer": record["renderer"],
            "fallback_reason": record["fallback_reason"],
            "baseline_sha": record["baseline_sha"],
            "reason_arg": reason,
            "plan_summary": record.get("plan_summary"),
        })
        outputs = [Path(job.output_doc), Path(job.student_output_doc)]
        for output in outputs:
            write_product_fixture(output, "class")
        return {"output_paths": [str(path) for path in outputs], "whole_job": True,
                "baseline_sha": renderer_orchestrator.V09_BASELINE_SHA}

    monkeypatch.setattr(renderer_orchestrator, "render_v09_whole_job", successful_fallback)
    response = _post_files(client, [("Unit 教师版.docx", make_docx()),
                                    ("Unit 学生版.docx", make_student_docx())])
    assert response.status_code == 202
    job_id = response.get_json()["job_id"]
    final = make_service(root).get(job_id)
    assert fallback_observed == {
        "renderer": "V0.9",
        "fallback_reason": "ALIGNMENT_UNRESOLVED",
        "baseline_sha": renderer_orchestrator.V09_BASELINE_SHA,
        "reason_arg": "ALIGNMENT_UNRESOLVED",
        "plan_summary": None,
    }
    assert final["status"] == "done", final.get("error")
    assert final["renderer"] == "V0.9"
    assert final["fallback_reason"] == "ALIGNMENT_UNRESOLVED"
    assert final["renderer_route"] == "V09_WHOLE_JOB"
    assert final["fallback_phase"] == "ROUTER"
    assert final["fallback_reason_code"] == "ALIGNMENT_UNRESOLVED"
    assert final["student_preparation_route"] == "BYPASS"
    assert final["baseline_sha"] == renderer_orchestrator.V09_BASELINE_SHA
    assert final.get("plan_summary") is None
    assert len(final["output_paths"]) == 2
    assert all(Path(path).is_file() for path in final["output_paths"])


def test_failed_fallback_attempt_is_persisted_before_runtime_failure(client, tmp_path, monkeypatch):
    import shutil
    import renderer_orchestrator

    root = tmp_path / "results"
    app.config.update(C0_DISABLE_JOB_SUBMISSION=False,
                      C0_RUN_JOBS_SYNCHRONOUSLY=True,
                      C0_FORCE_FALLBACK_REASON="UNSUPPORTED_REVISION_MARKUP",
                      V12_ALLOW_DIAGNOSTIC_V09=True)
    engine = SimpleNamespace(make_student=lambda source, target: write_product_fixture(
                                 target, "class"),
                             CLASS_TEMPLATE="unused-class-template.docx",
                             DEFAULT_TEMPLATE="unused-1v1-template.docx")
    monkeypatch.setattr(renderer_orchestrator, "_load_v09_engine", lambda: engine)
    observed = {}

    def fail_fallback(_job, reason):
        record = make_service(root).list()[0]
        observed.update({"renderer": record["renderer"],
                         "fallback_reason": record["fallback_reason"],
                         "baseline_sha": record["baseline_sha"],
                         "reason_arg": reason})
        raise RuntimeError("simulated whole-job fallback failure")

    monkeypatch.setattr(renderer_orchestrator, "render_v09_whole_job", fail_fallback)
    response = post_one(client)
    job_id = response.get_json()["job_id"]
    final = make_service(root).get(job_id)
    expected = renderer_orchestrator.V09_BASELINE_SHA
    assert observed == {"renderer": "V0.9",
                        "fallback_reason": "UNSUPPORTED_REVISION_MARKUP",
                        "baseline_sha": expected,
                        "reason_arg": "UNSUPPORTED_REVISION_MARKUP"}
    assert final["status"] == "error"
    assert final["renderer"] == "V0.9"
    assert final["fallback_reason"] == "UNSUPPORTED_REVISION_MARKUP"
    assert final["renderer_route"] == "V09_WHOLE_JOB"
    # Metadata now resolves before renderer selection so every output path
    # receives the same cover normalization and field-level warnings.
    assert final["fallback_phase"] == "METADATA"
    assert final["fallback_reason_code"] == "UNSUPPORTED_REVISION_MARKUP"
    assert final["baseline_sha"] == expected


def test_synthetic_pair_without_question_skeleton_fails_closed_without_summary(
        client, tmp_path):
    root = tmp_path / "results"
    app.config.update(C0_DISABLE_JOB_SUBMISSION=False,
                      C0_RUN_JOBS_SYNCHRONOUSLY=True)
    response = _post_files(client, [("Unit 教师版.docx", make_docx()),
                                    ("Unit 学生版.docx", make_student_docx())])
    job_id = response.get_json()["job_id"]
    final = make_service(root).get(job_id)
    assert final["status"] == "error"
    assert final["renderer"] == "XML_UNSUPPORTED"
    assert final["renderer_reason_code"] == "ALIGNMENT_UNRESOLVED"
    assert final.get("fallback_reason") is None
    assert final.get("plan_summary") is None


def test_whole_job_fallback_preserves_training_route_evidence(client):
    created = post_one(client, "training.docx").get_json()
    service = make_service()
    service.start_job(created["job_id"])
    service.update_route_evidence(created["job_id"], {
        "training_split_strategy": "GROUPED_VERTICAL",
        "training_question_routes": [
            {"question_number": 1, "slot": "knowledge"},
            {"question_number": 2, "slot": "immediate"},
            {"question_number": 3, "slot": "final"},
        ],
        "xml_renderer_attempted": True,
        "display_renumbering": {
            "status": "DISPLAY_RENUMBER_APPLIED", "applied": True,
            "slots": {"knowledge": [{"source_occurrence": 1, "new_number": 1}]},
        },
    })
    service.record_fallback_attempt(created["job_id"], "XML_RENDER_FAILED", "frozen-sha",
                                    detail="ProjectionError: unsupported test fixture", phase="RENDER")
    record = service.get(created["job_id"])
    assert record["renderer_route"] == "V09_WHOLE_JOB"
    assert record["fallback_phase"] == "RENDER"
    assert record["training_split_strategy"] == "GROUPED_VERTICAL"
    assert record["training_question_routes"] == [
        {"question_number": 1, "slot": "knowledge"},
        {"question_number": 2, "slot": "immediate"},
        {"question_number": 3, "slot": "final"},
    ]
    assert record["xml_renderer_attempted"] is True
    assert record["display_renumbering"] == {
        "planned_status": "DISPLAY_RENUMBER_APPLIED",
        "status": "DISPLAY_RENUMBER_NOT_APPLIED_FALLBACK", "applied": False,
        "slots": {"knowledge": [{"source_occurrence": 1, "new_number": 1}]},
    }


def test_xml_unsupported_route_is_persisted_without_fallback(client):
    created = post_one(client, "training.docx").get_json()
    service = make_service()
    service.start_job(created["job_id"])
    service.update_route_evidence(created["job_id"], {
        "renderer_route": "XML_UNSUPPORTED",
        "renderer_reason_code": "ALIGNMENT_UNRESOLVED",
        "renderer_reason_detail": "unmatched canonical question",
    })
    service.fail_job(created["job_id"], "XML_UNSUPPORTED: ALIGNMENT_UNRESOLVED")
    record = service.get(created["job_id"])
    assert record["status"] == "error"
    assert record["renderer"] == "XML_UNSUPPORTED"
    assert record["renderer_route"] == "XML_UNSUPPORTED"
    assert record["renderer_reason_code"] == "ALIGNMENT_UNRESOLVED"
    assert record["renderer_reason_detail"] == "unmatched canonical question"
    assert record.get("fallback_reason") is None


def test_restart_reconciles_done_job_when_a_final_docx_is_missing(client):
    created = post_one(client).get_json()
    job_dir = Path(created["result_dir"])
    teacher = job_dir / "教师版.docx"
    student = job_dir / "学生版.docx"
    teacher.write_bytes(make_docx())
    student.write_bytes(make_docx())
    service = make_service()
    service.complete_job(created["job_id"], teacher, student)
    student.unlink()
    recovered = service.get(created["job_id"])
    assert recovered["status"] == "error"
    assert "download_available" not in recovered
    assert recovered["has_result"] is False


def test_duplicate_role_path_never_completes_job(client):
    created = post_one(client).get_json()
    output = Path(created["result_dir"]) / "教师版.docx"
    output.write_bytes(make_docx())
    service = make_service()
    with pytest.raises(ValueError, match="different files"):
        service.complete_job(created["job_id"], output, output)
    assert service.get(created["job_id"])["status"] == "queued"
    assert service.get(created["job_id"])["has_result"] is False


def test_corrupt_document_xml_is_rejected_on_complete_and_recovery(client):
    created = post_one(client).get_json()
    job_dir = Path(created["result_dir"])
    teacher, student = job_dir / "教师版.docx", job_dir / "学生版.docx"
    teacher.write_bytes(make_docx())
    student.write_bytes(make_docx())
    service = make_service()
    corrupt_document_xml(student)
    with pytest.raises(ValueError, match="package validation failed"):
        service.complete_job(created["job_id"], teacher, student)
    assert service.get(created["job_id"])["status"] == "queued"
    student.write_bytes(make_docx())
    service.complete_job(created["job_id"], teacher, student)
    corrupt_document_xml(student)
    recovered = service.get(created["job_id"])
    assert recovered["status"] == "error"
    assert recovered["has_result"] is False
    assert "download_available" not in recovered
    response = client.get("/api/download/" + created["job_id"])
    assert response.status_code == 404


def test_workspace_javascript_handles_queued_and_running_states():
    js = (WEBAPP / "static" / "workspace.js").read_text(encoding="utf-8")
    assert 'function isActiveStatus(status) { return status === "queued" || status === "running"; }' in js
    assert "if (isActiveStatus(job.status)) state.pollTimer" in js
    assert 'job.status === "queued" ? "排队中"' in js
    assert 'localStorage.setItem("handout_current_job", response.job_id)' in js
    template = (WEBAPP / "templates" / "index.html").read_text(encoding="utf-8")
    assert '<option value="queued">排队中</option>' in template
    assert 'id="resultDelivery"' in template
    assert 'id="openResult"' in template and "打开成品文件夹" in template
    assert "本地输出路径：" in js and "下载 ZIP" not in template


def test_semantic_facade_hashes_one_source_snapshot_and_returns_same_units():
    from semantic_facade import analyze_source, PREDICTOR_PATH

    source = make_docx()
    snapshot = analyze_source(source)
    assert snapshot.source_sha256 == __import__("hashlib").sha256(source).hexdigest()
    # Identity and semantic-unit shape are stable against the authoritative
    # predictor directly; the facade does not implement new rules.
    spec = importlib.util.spec_from_file_location("test_authoritative_predictor", PREDICTOR_PATH)
    predictor = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(predictor)
    expected_index, expected_units = predictor.predict(snapshot.document)
    assert snapshot.node_index.order_ids == expected_index.order_ids
    assert snapshot.node_index.order_ids
    assert snapshot.units == expected_units
    assert isinstance(snapshot.units, list)


def _fake_slot_plan():
    return SimpleNamespace(
        template_sha256="fixture-template",
        slots={slot: () for slot in ("knowledge", "immediate", "final")},
        slot_labels={"knowledge": "知识精讲&例题讲解", "immediate": "即时训练",
                     "final": "六、出门测试"},
        units=(), explicit_final_heading=True,
    )


def _post_files(client, files, *, template_type="class"):
    data = {
        "subject": "物理", "grade": "九年级", "handout_type": "复习讲义",
        "academic_year": "2026-2027学年", "template_type": template_type,
        "split_mode": "smart", "docx_mode": "auto",
        "files": [(io.BytesIO(content), name) for name, content in files],
    }
    return client.post("/api/jobs", data=data, content_type="multipart/form-data")


def _fake_render_slots(source, _plan, output):
    template_type = "class" if _plan.slot_labels.get("final") == "六、出门测试" else "1v1"
    write_product_fixture(output, template_type)
    return SimpleNamespace(output_path=str(output), inserted_nodes=1,
                           resource_report={"unsupported": []},
                           package_report={"valid": True, "errors": []},
                           page_layout=None)


@pytest.mark.parametrize("template_type", ["1v1", "class"])
def test_paired_teacher_student_routes_each_input_without_make_student(
        client, monkeypatch, template_type):
    import renderer_orchestrator
    import template_slot_composer

    app.config.update(C0_DISABLE_JOB_SUBMISSION=False, C0_RUN_JOBS_SYNCHRONOUSLY=True)
    monkeypatch.setattr(renderer_orchestrator, "_load_v09_engine",
                        lambda: (_ for _ in ()).throw(AssertionError("make_student engine loaded")))
    seen_sources = []
    seen_cover_metadata = []

    def render(source, plan, output, *, cover_metadata=None):
        seen_sources.append(Path(source).read_bytes())
        seen_cover_metadata.append(cover_metadata)
        return _fake_render_slots(source, plan, output)

    monkeypatch.setattr(template_slot_composer, "render_slots", render)
    teacher_path = next(INCIDENT_FIXTURE.glob("*解析版*.docx"))
    student_path = next(INCIDENT_FIXTURE.glob("*原卷版*.docx"))
    teacher = teacher_path.read_bytes()
    student = student_path.read_bytes()
    response = _post_files(client, [("Unit 教师版.docx", teacher), ("Unit 学生版.docx", student)],
                           template_type=template_type)
    job = response.get_json()
    final = make_service().get(job["job_id"])
    assert final["status"] == "done", final.get("error")
    assert final["input_version"] == "TEACHER_AND_STUDENT"
    assert final["renderer"] == "XML"
    assert final["renderer_route"] == "XML"
    assert final["canonical_alignment"]["unmatched_real_question_count"] == 0
    assert final["canonical_alignment"]["ambiguous_real_question_count"] == 0
    assert len(final["canonical_occurrence_routes"]) == 60
    projection = final["canonical_projection"]
    assert projection["status"] == "PROJECTED"
    assert projection["source_block_counts"] == {"student": 356, "teacher": 860}
    assert len(projection["student_block_routes"]) == 356
    assert len(projection["teacher_block_routes"]) == 860
    for occurrence in final["canonical_occurrence_routes"]:
        assert projection["student_block_routes"][occurrence["student_node"][1:]] == \
            occurrence["destination_slot"]
        assert projection["teacher_block_routes"][occurrence["teacher_node"][1:]] == \
            occurrence["destination_slot"]
    assert final["student_preparation"]["make_student_called"] is False
    assert seen_sources == [teacher, student]
    assert len(seen_cover_metadata) == 2
    # The incident source has no reliable cover fields. Both outputs must use
    # the same teacher-owned metadata result, including an explicit empty state.
    assert seen_cover_metadata[0] == seen_cover_metadata[1]
    metadata = final["lesson_metadata"]
    assert metadata["status"] == "UNAVAILABLE"
    assert metadata["objectives"] == metadata["difficulties"] == ""
    assert metadata["warning"]
    assert metadata["full_objectives"] == metadata["objectives"]
    assert metadata["full_difficulties"] == metadata["difficulties"]
    assert metadata["cover_display"] == seen_cover_metadata[0]["cover_display"]
    renumbering = final["display_renumbering"]
    assert renumbering["status"] == "DISPLAY_RENUMBER_APPLIED"
    assert set(renumbering["slots"]) == {"knowledge", "immediate", "final"}
    assert sum(len(items) for items in renumbering["slots"].values()) == 60
    for items in renumbering["slots"].values():
        assert [item["new_number"] for item in items] == list(range(1, len(items) + 1))
        assert all(item["source_nodes"]["teacher"] and item["source_nodes"]["student"]
                   for item in items)
    assert final["items"][0]["topic"] == "Unit"
    result_dir = Path(final["result_dir"])
    assert {path.name for path in result_dir.iterdir()} == {
        Path(final["teacher_output_path"]).name, Path(final["student_output_path"]).name}
    assert all("_" not in path.name for path in result_dir.iterdir())
    assert not any(path.name in {"work", "job.json"} for path in result_dir.iterdir())


def test_student_only_creates_no_teacher_output_and_never_calls_make_student(client, monkeypatch):
    import renderer_orchestrator
    import template_slot_composer

    app.config.update(C0_DISABLE_JOB_SUBMISSION=False, C0_RUN_JOBS_SYNCHRONOUSLY=True)
    monkeypatch.setattr(renderer_orchestrator, "_load_v09_engine",
                        lambda: (_ for _ in ()).throw(AssertionError("student-only must not load COM")))
    seen = []

    def render(source, plan, output):
        seen.append(Path(source).read_bytes())
        return _fake_render_slots(source, plan, output)

    monkeypatch.setattr(template_slot_composer, "render_slots", render)
    student = make_student_docx()
    response = _post_files(client, [("Chapter 学生版.docx", student)])
    job = response.get_json()
    final = make_service().get(job["job_id"])
    assert final["status"] == "done", final.get("error")
    assert final["input_version"] == "STUDENT_ONLY"
    assert final["output_roles"].keys() == {"student"}
    assert final["teacher_output_path"] is None
    assert final["items"][0]["teacher"] == "未提供"
    assert final["student_preparation"]["make_student_called"] is False
    assert seen == [student]
    result_dir = Path(final["result_dir"])
    assert len(list(result_dir.iterdir())) == 1
    assert list(result_dir.iterdir())[0].name.endswith("学生版.docx")


def test_unknown_input_version_is_explicit_and_does_not_create_result_folder(client):
    response = _post_files(client, [("Chapter.docx", make_student_docx())])
    assert response.status_code == 422
    assert response.get_json()["error_code"] == "UNKNOWN_INPUT_VERSION"
    assert list(Path(app.config["RESULT_ROOT"]).iterdir()) == []


def test_unrelated_teacher_student_pair_is_rejected_without_result_folder(client):
    response = _post_files(client, [("math 教师版.docx", make_docx()),
                                    ("chem 学生版.docx", make_student_docx())])
    assert response.status_code == 415
    assert "error" in response.get_json()
    assert list(Path(app.config["RESULT_ROOT"]).iterdir()) == []


def test_failed_make_student_attempt_is_persisted_truthfully(client, monkeypatch):
    import renderer_orchestrator

    app.config.update(C0_DISABLE_JOB_SUBMISSION=False, C0_RUN_JOBS_SYNCHRONOUSLY=True)
    app.config["V12_ALLOW_DIAGNOSTIC_V09"] = True
    engine = SimpleNamespace(
        make_student=lambda *_args: (_ for _ in ()).throw(RuntimeError("simulated COM failure")),
        CLASS_TEMPLATE="unused-class-template.docx",
        DEFAULT_TEMPLATE="unused-1v1-template.docx",
    )
    monkeypatch.setattr(renderer_orchestrator, "_load_v09_engine", lambda: engine)
    def failed_whole_job(job, _reason):
        engine.make_student(job.source_doc, job.student_output_doc)
    monkeypatch.setattr(renderer_orchestrator, "render_v09_whole_job", failed_whole_job)
    response = post_one(client, "Chapter 教师版.docx", make_docx())
    job_id = response.get_json()["job_id"]
    final = make_service().get(job_id)
    assert final["status"] == "error"
    assert final["student_preparation"]["make_student_called"] is True
    assert final["student_preparation"]["wps_com_started"] is None
    assert final["student_preparation"]["elapsed_seconds"] is not None
