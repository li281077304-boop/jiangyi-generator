import io
from pathlib import Path
import sys

import pytest
from docx import Document

from product_fixture_utils import write_product_fixture

ROOT = Path(__file__).resolve().parents[1]
WEBAPP = ROOT / "v1.2-xml-experiment" / "res" / "app" / "webapp"
APP_DIR = WEBAPP.parent
sys.path.insert(0, str(WEBAPP))
sys.path.insert(0, str(APP_DIR))

from app import app  # noqa: E402
import app as app_module  # noqa: E402
from job_service import JobService  # noqa: E402


def source_docx(text: str) -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.add_paragraph("【答案】A")
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def stable_client(tmp_path):
    app.config.update(TESTING=True, RESULT_ROOT=tmp_path / "results",
                      RUNTIME_ROOT=tmp_path / "runtime", C0_DISABLE_JOB_SUBMISSION=False,
                      C0_RUN_JOBS_SYNCHRONOUSLY=True, C0_FORCE_FALLBACK_REASON=None,
                      STUDENTIZER_EVIDENCE_PROVIDER=None)
    app.config.pop("OPEN_FOLDER", None)
    return app.test_client()


def submit(client, files, *, engine_mode="stable_v09", template_type="1v1"):
    form = {"subject": "数学", "grade": "高一", "handout_type": "复习讲义",
            "academic_year": "2026-2027学年", "template_type": template_type,
            "engine_mode": engine_mode, "split_mode": "smart", "docx_mode": "auto",
            "files": [(io.BytesIO(data), name) for name, data in files]}
    return client.post("/api/jobs", data=form, content_type="multipart/form-data")


@pytest.mark.parametrize("kind", ["teacher_only", "pair", "student_only"])
def test_stable_mode_uses_only_selected_v09_and_preserves_input_roles(
        stable_client, monkeypatch, kind):
    import renderer_orchestrator
    import slot_router
    import studentizer_planner
    import template_slot_composer

    make_student_calls = []
    render_calls = []

    def make_student(source, target):
        make_student_calls.append((source, target))
        write_product_fixture(target, "1v1")

    engine = type("Engine", (), {"make_student": staticmethod(make_student)})()
    monkeypatch.setattr(renderer_orchestrator, "_load_v09_engine", lambda: engine)

    def selected(job):
        render_calls.append(job)
        outputs = [Path(job.output_doc)]
        if not job.student_only:
            outputs.append(Path(job.student_output_doc))
        for output in outputs:
            write_product_fixture(output, job.template_type)
        return {"output_paths": [str(path) for path in outputs],
                "fallback_reason": None, "selected_engine": "V0.9"}

    monkeypatch.setattr(renderer_orchestrator, "render_v09_selected", selected)

    def forbidden(*_args, **_kwargs):
        raise AssertionError("stable mode must bypass XML analysis and Studentizer")

    monkeypatch.setattr(slot_router, "analyze_source", forbidden)
    monkeypatch.setattr(studentizer_planner, "prepare_complete_student", forbidden)
    monkeypatch.setattr(template_slot_composer, "render_slots", forbidden)

    if kind == "teacher_only":
        files = [("专题 教师版.docx", source_docx("teacher source"))]
    elif kind == "pair":
        files = [("专题 教师版.docx", source_docx("teacher source")),
                 ("专题 学生版.docx", source_docx("student source"))]
    else:
        files = [("专题 学生版.docx", source_docx("student source"))]
    response = submit(stable_client, files, template_type="class")
    assert response.status_code == 202

    service = JobService(Path(app.config["RESULT_ROOT"]), runtime_root=Path(app.config["RUNTIME_ROOT"]))
    job = service.get(response.get_json()["job_id"])
    assert job["status"] == "done", job.get("error")
    assert job["renderer"] == "V0.9"
    assert job["renderer_route"] == "STABLE_V09"
    assert job["selected_engine"] == "V0.9"
    assert job["fallback_reason"] is None
    assert job["baseline_sha"] == renderer_orchestrator.V09_BASELINE_SHA
    assert len(render_calls) == 1
    assert len(job["output_paths"]) == (1 if kind == "student_only" else 2)
    assert all(Path(path).is_file() for path in job["output_paths"])
    assert make_student_calls and len(make_student_calls) == 1 if kind == "teacher_only" else not make_student_calls
    if kind == "teacher_only":
        assert job["student_preparation_route"] == "V09_MAKE_STUDENT"
        assert job["student_preparation"]["make_student_called"] is True
        assert any("请在使用前检查" in warning for warning in job["warnings"])
    else:
        assert job["student_preparation_route"] == "BYPASS"


def test_job_without_engine_option_defaults_to_stable_v09(tmp_path):
    from job_service import JobService

    service = JobService(tmp_path / "results", runtime_root=tmp_path / "runtime")
    job = service.create_inputs([("专题 学生版.docx", source_docx("student"))], {
        "template_type": "1v1", "docx_mode": "auto"})
    assert job["options"]["engine_mode"] == "stable_v09"
    assert job["selected_engine"] == "V0.9"
    assert job["renderer_route"] == "STABLE_V09"


def test_batch_children_keep_selected_engine_mode(tmp_path):
    service = JobService(tmp_path / "results", runtime_root=tmp_path / "runtime")
    batch = service.create_batch([
        ("One 学生版.docx", source_docx("one")),
        ("Two 学生版.docx", source_docx("two")),
        ("Three 学生版.docx", source_docx("three")),
    ], {"engine_mode": "xml_restricted", "docx_mode": "separate"})
    assert batch["options"]["engine_mode"] == "xml_restricted"
    for item in batch["items"]:
        child = service.get(item["child_job_id"])
        assert child["options"]["engine_mode"] == "xml_restricted"
        assert child["selected_engine"] == "XML"
