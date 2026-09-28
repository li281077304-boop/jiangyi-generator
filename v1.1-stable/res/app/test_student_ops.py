# -*- coding: utf-8 -*-
"""Regression tests for red-answer removal in student documents."""
import os
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from student_ops import strip_red


def _add_run(paragraph, text, color=None):
    run = paragraph.add_run(text)
    if color is not None:
        run.font.color.rgb = color
    return run


def _all_paragraph_texts(doc):
    texts = []
    seen_cells = set()

    def _walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                cell_key = cell._tc
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                texts.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested_table in cell.tables:
                    _walk_table(nested_table)

    texts.extend(paragraph.text for paragraph in doc.paragraphs)
    for table in doc.tables:
        _walk_table(table)
    return texts


class StripRedTests(unittest.TestCase):
    def _strip(self, build_document):
        with tempfile.TemporaryDirectory() as tempdir:
            source = os.path.join(tempdir, "teacher.docx")
            output = os.path.join(tempdir, "student.docx")
            document = Document()
            build_document(document)
            document.save(source)

            removed = strip_red(source, output)

            self.assertGreater(removed, 0)
            return _all_paragraph_texts(Document(output))

    def test_removes_red_answer_from_body_paragraph(self):
        def build(document):
            paragraph = document.add_paragraph()
            _add_run(paragraph, "黑色正文")
            _add_run(paragraph, "红色答案", RGBColor(0xFF, 0x00, 0x00))

        texts = self._strip(build)

        self.assertIn("黑色正文", texts)
        self.assertNotIn("红色答案", "".join(texts))

    def test_removes_red_answer_from_table_cell(self):
        def build(document):
            table = document.add_table(rows=1, cols=2)
            _add_run(table.cell(0, 0).paragraphs[0], "黑色题干")
            _add_run(table.cell(0, 1).paragraphs[0], "红色答案", RGBColor(0xFF, 0x00, 0x00))

        texts = self._strip(build)

        self.assertIn("黑色题干", texts)
        self.assertNotIn("红色答案", "".join(texts))

    def test_removes_red_answers_from_body_table_and_nested_table(self):
        def build(document):
            paragraph = document.add_paragraph()
            _add_run(paragraph, "正文保留")
            _add_run(paragraph, "正文答案", RGBColor(0xFF, 0x00, 0x00))

            table = document.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            _add_run(cell.paragraphs[0], "表格题干")
            _add_run(cell.paragraphs[0], "表格答案", RGBColor(0xFF, 0x00, 0x00))

            nested_table = cell.add_table(rows=1, cols=1)
            _add_run(nested_table.cell(0, 0).paragraphs[0], "嵌套题干")
            _add_run(nested_table.cell(0, 0).paragraphs[0], "嵌套答案", RGBColor(0xFF, 0x00, 0x00))

        texts = self._strip(build)

        remaining = "".join(texts)
        for expected in ("正文保留", "表格题干", "嵌套题干"):
            self.assertIn(expected, remaining)
        for removed in ("正文答案", "表格答案", "嵌套答案"):
            self.assertNotIn(removed, remaining)

    def test_preserves_non_text_xml_in_red_run(self):
        with tempfile.TemporaryDirectory() as tempdir:
            source = os.path.join(tempdir, "teacher.docx")
            output = os.path.join(tempdir, "student.docx")
            document = Document()
            red_run = _add_run(
                document.add_paragraph(), "红色答案", RGBColor(0xFF, 0x00, 0x00))
            red_run._element.append(OxmlElement("w:drawing"))
            document.save(source)

            self.assertEqual(strip_red(source, output), 1)
            with zipfile.ZipFile(output) as package:
                document_xml = package.read("word/document.xml").decode("utf-8")

            self.assertNotIn("红色答案", document_xml)
            self.assertIn("<w:drawing", document_xml)


if __name__ == "__main__":
    unittest.main()
