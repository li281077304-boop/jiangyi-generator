# -*- coding: utf-8 -*-
"""
student_ops.py —— 学生版加工

流程（make_student 统一入口）：
  1. strip_red：清空红色答案 run（保留段落结构与公式/OLE 对象）
  2. _strip_answer_markers：用 Word COM 删除「答案/解析」标记段（覆盖黑色、灰底等非红色答案）
  3. _add_blanks：为简答题补空白行

关键约束：
  - 删除段落必须走 COM，让 Word 自己处理，否则公式/OLE 会损坏。
  - PS1 脚本缺失时跳过对应步骤，不影响其他步骤结果。
"""
import os
import re
import subprocess
import locale

from split_engine import _is_red_hex
from docutils import POWERSHELL
import config
import logger

log = logger.get_logger("student_ops")


def strip_red(src_path: str, out_path: str) -> int:
    """
    清除红色答案文字（保留段落结构与公式/OLE对象）。
    用 python-docx 仅清空红色 run 的文本，不复制段落，故公式不受损。
    返回清除的 run 数。
    """
    from docx import Document
    NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    doc = Document(src_path)
    n = 0
    for para in doc.paragraphs:
        for run in para.runs:
            col = run.font.color
            if col is not None and col.rgb is not None and _is_red_hex(str(col.rgb).upper()):
                run.text = ""
                n += 1
    # XML 层兜底（python-docx runs 可能漏掉嵌套结构）
    for para in doc.paragraphs:
        for r_elem in para._element.findall(f".//{NS}r"):
            for ce in r_elem.findall(f".//{NS}color"):
                val = ce.get(f"{NS}val", "")
                if val and _is_red_hex(val.upper()):
                    for te in r_elem.findall(f".//{NS}t"):
                        if te.text:
                            te.text = ""
                            n += 1
    doc.save(out_path)
    return n


def _add_blanks(src_path: str, out_path: str):
    """在学生版文档末尾，为简答题补空白行。返回 (简答题数, 空白行数, 日志)。"""
    if not os.path.exists(config.PS_ADD_BLANKS):
        return 0, 0, "[make_student] 缺 _add_blanks_com.ps1，跳过留白"
    r = subprocess.run(
        [POWERSHELL, "-NonInteractive", "-ExecutionPolicy", "Bypass",
         "-File", config.PS_ADD_BLANKS,
         "-SrcPath", os.path.abspath(src_path), "-OutPath", os.path.abspath(out_path)],
        capture_output=True, text=True, encoding=locale.getpreferredencoding(),
        errors="replace", timeout=300)
    log_out = r.stdout.strip()
    if r.returncode not in (0, 1) and r.stderr.strip():
        log_out += "\n[PS stderr] " + r.stderr[:400]
    q_count = 0
    blank_count = 0
    for line in log_out.splitlines():
        m = re.search(r'SHORT_ANSWER=(\d+)', line)
        if m:
            q_count = int(m.group(1))
        m = re.search(r'BLANKS=(\d+)', line)
        if m:
            blank_count = int(m.group(1))
    return q_count, blank_count, log_out


def _strip_answer_markers(src_path: str, out_path: str):
    """
    用 Word COM 删除"答案/解析"标记段（覆盖黑色、灰底等非红色答案）。
    集中答案区（文档后 1/3 的短标题型标记）删到结尾；逐题答案只删标记段本身。
    返回 (删除段数, 日志)。
    """
    if not os.path.exists(config.PS_STRIP_ANSWER):
        # 脚本缺失则跳过（不影响 red 清除的结果）
        return 0, "[make_student] 缺 _strip_answer_com.ps1，跳过标记段删除"
    r = subprocess.run(
        [POWERSHELL, "-NonInteractive", "-ExecutionPolicy", "Bypass",
         "-File", config.PS_STRIP_ANSWER,
         "-SrcPath", os.path.abspath(src_path), "-OutPath", os.path.abspath(out_path)],
        capture_output=True, text=True, encoding=locale.getpreferredencoding(),
        errors="replace", timeout=300)
    log_out = r.stdout.strip()
    if r.returncode not in (0, 1) and r.stderr.strip():
        log_out += "\n[PS stderr] " + r.stderr[:400]
    deleted = 0
    for line in log_out.splitlines():
        if "DELETED=" in line:
            m = re.search(r'DELETED=(\d+)', line)
            if m:
                deleted = int(m.group(1))
                break
    return deleted, log_out


def make_student(src_path: str, out_path: str):
    """
    生成学生版统一入口：先 strip_red 清红色答案 run，再删答案/解析标记段（黑/灰底），
    最后为简答题加空白行。
    返回 (红色清除数, 标记段删除数, 简答题数, 空白行数, 日志)。
    """
    n_red = strip_red(src_path, out_path)
    n_del, log1 = _strip_answer_markers(out_path, out_path)
    n_q, n_blk, log2 = _add_blanks(out_path, out_path)
    return n_red, n_del, n_q, n_blk, log1 + "\n" + log2
