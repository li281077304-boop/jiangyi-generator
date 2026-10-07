# -*- coding: utf-8 -*-
"""Source-backed lesson objective/difficulty and training-only coverage."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from lesson_metadata import (LessonMetadataUnavailable, read_lesson_source_lines,
                             resolve_lesson_metadata)


class LessonMetadataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="lesson-metadata-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def _source(self, name, paragraphs=(), cover=None):
        doc = Document()
        for value in paragraphs:
            doc.add_paragraph(value)
        if cover:
            table = doc.add_table(rows=0, cols=2)
            for label, value in cover:
                cells = table.add_row().cells
                cells[0].text = label
                cells[1].text = value
        path = self.root / name
        doc.save(path)
        return path

    def test_explicit_source_objectives_and_difficulties_take_priority(self):
        path = self._source("explicit.docx", cover=[
            ("教学目标", "按源稿目标保留"),
            ("重点难点", "重点：按源稿重点。难点：按源稿难点。"),
        ], paragraphs=["题型01 分子热运动", "1. 练习题干"])
        result = resolve_lesson_metadata(
            path, subject="物理", topic="分子动理论",
            knowledge_point_status="NO_KNOWLEDGE_POINT",
        )
        self.assertEqual(result.objectives, "按源稿目标保留")
        self.assertEqual(result.difficulties, "重点：按源稿重点。难点：按源稿难点。")
        self.assertEqual(result.source, "SOURCE")

    def test_training_only_uses_explicit_type_titles_without_fake_knowledge(self):
        path = self._source("training.docx", paragraphs=[
            "第2节 分子动理论的初步知识", "题型分组练",
            "题型01 物质的构成", "1. 判断题", "2. 说明题",
            "题型02 分子热运动", "3. 选择题",
        ])
        result = resolve_lesson_metadata(
            path, subject="数学", topic="分子动理论专题训练",
            knowledge_point_status="NO_KNOWLEDGE_POINT",
        )
        self.assertEqual(result.source, "TRAINING_TYPE_HEADINGS")
        self.assertEqual(result.training_titles, ("物质的构成", "分子热运动"))
        self.assertIn("物质的构成", result.objectives)
        self.assertIn("分子热运动", result.difficulties)
        for prohibited in ("核心知识点", "核心概念", "知识体系"):
            self.assertNotIn(prohibited, result.objectives + result.difficulties)

    def test_no_reliable_source_is_empty_and_warnable_without_template_text(self):
        path = self._source("unclassified.docx", paragraphs=["1. 普通练习题"])
        result = resolve_lesson_metadata(
            path, subject="数学", topic="未命名讲义",
            knowledge_point_status="NO_KNOWLEDGE_POINT",
        )
        self.assertEqual(result.objectives, "")
        self.assertEqual(result.difficulties, "")
        self.assertEqual(result.objectives_reason, "NO_RELIABLE_OBJECTIVES_SOURCE")
        self.assertEqual(result.difficulties_reason, "NO_RELIABLE_DIFFICULTIES_SOURCE")

    def test_existing_offline_rule_remains_the_knowledge_lesson_fallback(self):
        path = self._source("knowledge.docx", paragraphs=[
            "知识点1 一元一次方程的解法", "题型01 解方程", "1. 练习题",
        ])
        result = resolve_lesson_metadata(
            path, subject="数学", topic="一元一次方程",
            knowledge_point_status="KNOWLEDGE_POINT_PRESENT",
        )
        self.assertEqual(result.source, "V09_OFFLINE_RULE")
        self.assertIn("一元一次方程", result.objectives)
        self.assertTrue(result.difficulties)

    def test_numbered_question_after_cover_table_is_not_metadata(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "教学目标"
        table.cell(0, 1).text = "掌握一次函数的性质"
        table.cell(1, 0).text = "教学重难点"
        table.cell(1, 1).text = "重点：函数图像与性质"
        document.add_paragraph("1. 已知一次函数 y=2x+1，求其图像与性质。")
        source = self.root / "metadata-then-question.docx"
        document.save(source)

        lines = read_lesson_source_lines(source)
        result = resolve_lesson_metadata(
            source, subject="数学", topic="一次函数",
            knowledge_point_status="UNKNOWN", source_lines=lines)

        self.assertIn("掌握一次函数的性质", result.objectives)
        self.assertNotIn("已知一次函数", result.objectives)
        self.assertNotIn("已知一次函数", result.difficulties)

    def test_mixed_three_row_table_is_not_accepted_as_metadata_container(self):
        document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "教学目标"
        table.cell(0, 1).text = "掌握一次函数"
        table.cell(1, 0).text = "教学重难点"
        table.cell(1, 1).text = "重点：函数图像"
        table.cell(2, 0).text = "1."
        table.cell(2, 1).text = "已知一次函数，求解析式。"
        source = self.root / "mixed-table.docx"
        document.save(source)

        result = resolve_lesson_metadata(
            source, subject="数学", topic="一次函数",
            knowledge_point_status="UNKNOWN",
            source_lines=read_lesson_source_lines(source))

        self.assertEqual(result.objectives, "")
        self.assertEqual(result.difficulties, "")

    def test_full_difficulty_label_alias_is_not_split_as_a_prefix(self):
        path = self._source("full-difficulty-label.docx", cover=[
            ("教学目标", "掌握函数建模。"),
            ("教学重点难点", "函数建模与应用。"),
        ])
        result = resolve_lesson_metadata(
            path, subject="数学", topic="函数建模",
            knowledge_point_status="UNKNOWN",
        )
        self.assertEqual(result.difficulties, "函数建模与应用。")
        self.assertEqual(result.difficulties_source, "SOURCE")


if __name__ == "__main__":
    unittest.main()
