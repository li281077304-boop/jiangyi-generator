# -*- coding: utf-8 -*-
"""Regression checks for the focused XML package validator."""
import base64
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).parent))
from package_validator import validate_package


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
    "SMlYVQAAAABJRU5ErkJggg==")


def _fixture(directory):
    """Create a normal DOCX with a media relationship and a drawing."""
    image = os.path.join(directory, "one.png")
    Path(image).write_bytes(_PNG)
    output = os.path.join(directory, "fixture.docx")
    doc = Document()
    doc.add_paragraph("plain text")
    doc.add_paragraph().add_run().add_picture(image, width=Inches(0.1))
    doc.save(output)
    return output


def _rewrite_member(path, member, transform):
    """Copy a zip while replacing one member, preserving package structure."""
    changed = path + ".changed"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(changed, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            target.writestr(info, transform(data) if info.filename == member else data)
    shutil.move(changed, path)


class PackageValidatorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = _fixture(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_normal_docx_passes_and_reports_parts(self):
        report = validate_package(self.path)
        self.assertTrue(report["valid"], report)
        self.assertGreater(report["stats"]["package_parts"], 0)
        self.assertGreater(report["stats"]["drawing_image_references"], 0)

    def test_dangling_relationship_reference_is_invalid(self):
        _rewrite_member(
            self.path, "word/document.xml",
            lambda data: data.replace(b'r:embed="rId9"', b'r:embed="rId404"'),
        )
        report = validate_package(self.path)
        self.assertFalse(report["valid"])
        self.assertTrue(any("dangling relationship reference rId404" in e for e in report["errors"]), report)

    def test_missing_media_target_is_invalid(self):
        media = "word/media/image1.png"
        changed = self.path + ".changed"
        with zipfile.ZipFile(self.path) as source, zipfile.ZipFile(changed, "w") as target:
            for info in source.infolist():
                if info.filename != media:
                    target.writestr(info, source.read(info.filename))
        shutil.move(changed, self.path)
        report = validate_package(self.path)
        self.assertFalse(report["valid"])
        self.assertTrue(any("dangling internal relationship" in e for e in report["errors"]), report)

    def test_duplicate_drawing_id_is_invalid(self):
        def duplicate_id(data):
            # The standard fixture has one wp:docPr. Add a second inline copy
            # by duplicating its XML element is unnecessary: a second docPr
            # anywhere in document.xml is enough to exercise collision logic.
            return data.replace(b'</w:body>',
                b'<wp:docPr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" id="1" name="duplicate"/></w:body>')
        _rewrite_member(self.path, "word/document.xml", duplicate_id)
        report = validate_package(self.path)
        self.assertFalse(report["valid"])
        self.assertTrue(any("duplicate drawing id 1" in e for e in report["errors"]), report)

    def test_external_hyperlink_relationship_is_valid(self):
        document = Document()
        paragraph = document.add_paragraph("reference")
        relationship_id = document.part.relate_to(
            "https://example.test/reference", RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        paragraph._p.append(hyperlink)
        document.save(self.path)

        report = validate_package(self.path)

        self.assertTrue(report["valid"], report)
        self.assertGreater(report["stats"]["relationship_references"], 0)

    def test_ole_embedding_and_vml_preview_relationships_are_valid(self):
        image = os.path.join(self.tmp.name, "ole-preview.png")
        Path(image).write_bytes(_PNG)
        document = Document()
        paragraph = document.add_paragraph("OLE fixture")
        image_run = paragraph.add_run()
        image_run.add_picture(image, width=Inches(0.1))
        image_rid = image_run._r.xpath(".//a:blip/@r:embed")[0]

        preview = etree.SubElement(paragraph._p, "{%s}pict" %
                                   "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
        shape = etree.SubElement(preview, "{%s}shape" %
                                 "urn:schemas-microsoft-com:vml")
        image_data = etree.SubElement(shape, "{%s}imagedata" %
                                      "urn:schemas-microsoft-com:vml")
        image_data.set(qn("r:id"), image_rid)

        embedding = Part(
            PackURI("/word/embeddings/oleObject1.bin"),
            "application/vnd.openxmlformats-officedocument.oleObject",
            b"dummy OLE binary fixture",
            document.part.package,
        )
        ole_rid = document.part.relate_to(
            embedding,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
        )
        object_element = etree.SubElement(paragraph._p, "{%s}object" %
                                          "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
        ole_element = etree.SubElement(object_element, "{%s}OLEObject" %
                                       "urn:schemas-microsoft-com:office:office")
        ole_element.set(qn("r:id"), ole_rid)
        document.save(self.path)

        report = validate_package(self.path)

        self.assertTrue(report["valid"], report)
        self.assertEqual(1, report["stats"]["ole_references"])
        self.assertEqual(1, report["stats"]["embedding_parts"])
        self.assertEqual(1, report["stats"]["vml_image_references"])


if __name__ == "__main__":
    unittest.main()
