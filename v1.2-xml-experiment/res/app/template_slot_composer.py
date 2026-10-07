# -*- coding: utf-8 -*-
"""Place already-imported XML blocks into the frozen template slot anchors.

The B-Line renderer remains unchanged: it performs the package-safe source
block import into a staging DOCX. This C-Line composer moves those imported
physical blocks from the temporary body tail into the template's three merged
content-cell anchors, then validates the final package again.
"""
from __future__ import annotations

from dataclasses import dataclass
from copy import deepcopy
import os
import re
from hashlib import sha256
from pathlib import Path
import tempfile
from typing import Callable

from docx import Document
from docx.oxml import OxmlElement
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Pt

from lesson_metadata import (LessonMetadata, build_cover_display, display_width_units,
                             COVER_DISPLAY_BUDGET)
from package_validator import validate_package
from renderer_xml_minimal import BlockSpan, ProjectionError, RenderResult
from slot_router import (SLOT_ORDER, SlotRoutingError, SlotRoutingPlan,
                         _visible_question_number)
from struct_doc import W_P, W_SECTPR, W_T, W_TBL, W_TC


@dataclass(frozen=True)
class SlotRenderResult:
    output_path: str
    inserted_nodes: int
    slot_nodes: dict[str, int]
    resource_report: dict
    package_report: dict
    display_renumbering: dict | None = None
    module2_end_divider_anchor: dict | None = None
    page_layout: dict | None = None


def _paragraph_text(paragraph) -> str:
    return "".join(node.text or "" for node in paragraph.iter(W_T)).strip()


_CROSS_REFERENCE = re.compile(r"第\s*(?:\d+|[一二三四五六七八九十百]+)\s*题")
_DISPLAY_PREFIX = re.compile(r"^\s*(\d{1,3})\s*[.．、)）]\s*\S")
_UNSAFE_NUMBER_TAGS = {
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + name
    for name in ("txbxContent", "fldChar", "instrText", "ins", "del", "hyperlink")
}
_COVER_ROW_HEIGHTS_PT = {"1v1": (23.9, 23.9), "class": (25.7, 27.8)}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W_PPR = W + "pPr"
W_SPACING = W + "spacing"
W_TR = W + "tr"
W_TCPR = W + "tcPr"
W_TC_NS = W + "tc"

# --- Module 2 end-divider first-page anchor ---------------------------------
# The physical node being positioned is the existing divider after the
# "二、知识回顾" content area, not the heading. Template-specific values are
# calibrated against WPS PDF output; only the existing spacer paragraphs after
# the heading are changed. No paragraph or page geometry is added or removed.
MODULE2_END_DIVIDER_TARGET_Y = {
    # WPS 12.0 PDF text-bottom (y1) of the module-2 ending divider.
    "1v1": 760.3,
    "class": 740.1,
}
# Fixed exact line height for each existing interval after the module-2 heading.
MODULE2_END_DIVIDER_SPACER_PT = {
    # WPS PDF measurement of the V0.9-normalized physics pair showed the
    # 1v1 divider bottom 17.38pt below its 760.3pt target at 21pt x 13
    # intervals. 19.65pt moves the existing divider to the measured target.
    "1v1": 19.65,
    "class": 22.0,
}
# The headings and divider occupy fixed positions inside the template's merged
# content cell. Only the existing spacer block after the module-2 heading moves.
_KNOWLEDGE_CELL_ROW_INDEX = 5
_LAUNCH_HEADING_INDEX = 0
_KNOWLEDGE_HEADING_INDEX = 5
_MODULE2_END_DIVIDER_INDEX = {"1v1": 19, "class": 16}
def _editable_number_nodes(paragraph) -> tuple[list, re.Match | None]:
    if any(node.tag in _UNSAFE_NUMBER_TAGS for node in paragraph.iter()):
        return [], None
    nodes = paragraph.xpath("./w:r/w:t")
    text = "".join(node.text or "" for node in nodes)
    match = _DISPLAY_PREFIX.match(text)
    if match is None or _visible_question_number(_paragraph_text(paragraph)) != int(match.group(1)):
        return [], None
    return nodes, match


def build_display_renumbering(plans: dict[str, SlotRoutingPlan], *,
                              canonical_occurrences: list[dict] | None = None) -> dict:
    """One canonical source-occurrence map, shared by all paired presentations.

    Source question numbers and routing remain untouched. Table-contained
    values are excluded; cross-references or unsupported question prefixes
    disable the whole paired display transformation instead of guessing.
    """
    if not plans:
        return {"status": "NOT_APPLICABLE", "slots": {}}
    if canonical_occurrences is not None:
        return _build_canonical_display_renumbering(plans, canonical_occurrences)
    role_signatures = {}
    nodes_by_role = {}
    cross_reference_role = None
    for role, plan in plans.items():
        if sha256(plan.source_path.read_bytes()).hexdigest() != plan.source_sha256:
            raise SlotRoutingError("DISPLAY_RENUMBER_SOURCE_CHANGED", "source changed after routing")
        document = Document(str(plan.source_path))
        if _CROSS_REFERENCE.search(_paragraph_text(document.element.body)):
            cross_reference_role = cross_reference_role or role
        block_by_id = {record["block_id"]: record for record in plan.block_records}
        unit_roles = {unit["unit_id"]: unit["role"] for unit in plan.units}
        body_blocks = [child for child in document.element.body if child.tag in (W_P, W_TBL)]
        starts = []
        signature = []
        for record in plan.block_records:
            node_id = record["block_id"]
            if record.get("kind") != "paragraph":
                # Top-level tables are atomic; measurements and serials in
                # their cells are never display question starts.
                continue
            contributors = record.get("unit_contributors", ())
            if not any(unit_roles.get(unit_id) == "question_group" for unit_id in contributors):
                continue
            seq = int(record["source_index"])
            paragraph = body_blocks[seq] if 0 <= seq < len(body_blocks) else None
            if paragraph is not None and paragraph.tag != W_P:
                paragraph = None
            if paragraph is None:
                return {"status": "DISPLAY_RENUMBER_PARTIAL_UNSAFE_STRUCTURE", "applied": False,
                        "reason": "question block ordinal unresolved at %s:%s" % (role, node_id), "slots": {}}
            number = _visible_question_number(_paragraph_text(paragraph))
            if number is None:
                continue
            nodes, match = _editable_number_nodes(paragraph)
            if not nodes or match is None:
                return {"status": "DISPLAY_RENUMBER_PARTIAL_UNSAFE_STRUCTURE", "applied": False,
                        "reason": "non-editable question prefix at %s:%s" % (role, node_id), "slots": {}}
            slot = record["destination_slot"]
            starts.append(node_id)
            signature.append((number, slot))
        role_signatures[role] = tuple(signature)
        nodes_by_role[role] = starts
    signatures = list(role_signatures.values())
    if not signatures or not signatures[0]:
        return {"status": "NOT_APPLICABLE", "applied": False,
                "reason": "no validated top-level question starts", "slots": {}}
    if any(signature != signatures[0] for signature in signatures[1:]):
        return {"status": "DISPLAY_RENUMBER_PARTIAL_UNSAFE_STRUCTURE", "applied": False,
                "reason": "teacher/student routed question occurrences differ", "slots": {}}
    signature = signatures[0]
    slots = {slot: [] for slot in SLOT_ORDER}
    needs_renumber = False
    for occurrence, (number, slot) in enumerate(signature, 1):
        new_number = len(slots[slot]) + 1
        needs_renumber = needs_renumber or number != new_number
        display_number = number if cross_reference_role else new_number
        slots[slot].append({"source_occurrence": occurrence,
                            "source_question_number": number,
                            "source_order": occurrence,
                            "destination_slot": slot,
                            "new_number": display_number,
                            "source_nodes": {role: ids[occurrence - 1] for role, ids in nodes_by_role.items()}})
    if cross_reference_role:
        return {"status": "DISPLAY_RENUMBER_SKIPPED_CROSS_REFERENCE", "applied": False,
                "reason": "main-story cross-question reference found in " + cross_reference_role,
                "warning": "题目含跨题引用，已跳过展示题号重排",
                "slots": slots}
    status = "DISPLAY_RENUMBER_APPLIED" if needs_renumber else "DISPLAY_RENUMBER_NOT_NEEDED"
    return {"status": status, "applied": needs_renumber, "slots": slots}


def _build_canonical_display_renumbering(plans: dict[str, SlotRoutingPlan],
                                         occurrences: list[dict]) -> dict:
    """Build one numbering projection from the aligned pair occurrence map."""
    if not occurrences:
        return {"status": "NOT_APPLICABLE", "applied": False,
                "reason": "canonical question sequence is empty", "slots": {}}
    roles = tuple(sorted(plans))
    slots = {slot: [] for slot in SLOT_ORDER}
    cross_reference_role = None
    parsed = []
    for role in roles:
        plan = plans[role]
        if sha256(plan.source_path.read_bytes()).hexdigest() != plan.source_sha256:
            raise SlotRoutingError("DISPLAY_RENUMBER_SOURCE_CHANGED", "source changed after routing")
        document = Document(str(plan.source_path))
        if _CROSS_REFERENCE.search(_paragraph_text(document.element.body)):
            cross_reference_role = cross_reference_role or role
        body_blocks = [child for child in document.element.body if child.tag in (W_P, W_TBL)]
        parsed.append((role, plan, body_blocks))

    for ordinal, occurrence in enumerate(occurrences, 1):
        number = int(occurrence["question_number"])
        slot = occurrence.get("destination_slot")
        if slot not in SLOT_ORDER:
            raise SlotRoutingError("CANONICAL_ROUTE_INVALID",
                                   "canonical occurrence has no valid destination")
        source_nodes = {}
        for role, plan, body_blocks in parsed:
            node_id = occurrence.get(role + "_node")
            match_node = re.fullmatch(r"b(\d+)", str(node_id or ""))
            if not match_node:
                raise SlotRoutingError("DISPLAY_RENUMBER_PROJECTION_MISMATCH",
                                       "canonical question node is not a top-level block")
            seq = int(match_node.group(1))
            paragraph = body_blocks[seq] if 0 <= seq < len(body_blocks) else None
            if paragraph is None or paragraph.tag != W_P:
                raise SlotRoutingError("DISPLAY_RENUMBER_PROJECTION_MISMATCH",
                                       "canonical question endpoint is not a paragraph: %s:%s" %
                                       (role, node_id))
            visible = _visible_question_number(_paragraph_text(paragraph))
            nodes, prefix = _editable_number_nodes(paragraph)
            if visible != number or not nodes or prefix is None:
                raise SlotRoutingError("DISPLAY_RENUMBER_PROJECTION_MISMATCH",
                                       "canonical question number does not match editable source prefix at %s:%s" %
                                       (role, node_id))
            source_nodes[role] = node_id
        new_number = len(slots[slot]) + 1
        slots[slot].append({"source_occurrence": ordinal,
                            "canonical_occurrence_id": occurrence["canonical_occurrence_id"],
                            "source_question_number": number,
                            "source_order": ordinal,
                            "destination_slot": slot,
                            "new_number": new_number,
                            "source_nodes": source_nodes})
    if cross_reference_role:
        return {"status": "DISPLAY_RENUMBER_SKIPPED_CROSS_REFERENCE", "applied": False,
                "reason": "main-story cross-question reference found in " + cross_reference_role,
                "warning": "题目含跨题引用，已跳过展示题号重排", "slots": slots}
    needs_renumber = any(item["source_question_number"] != item["new_number"]
                         for slot_items in slots.values() for item in slot_items)
    return {"status": "DISPLAY_RENUMBER_APPLIED" if needs_renumber else "DISPLAY_RENUMBER_NOT_NEEDED",
            "applied": needs_renumber, "slots": slots}


def _apply_display_renumbering(offsets: dict, blocks_by_slot: dict, evidence: dict, role: str) -> None:
    if evidence.get("status") != "DISPLAY_RENUMBER_APPLIED":
        return
    for slot in SLOT_ORDER:
        elements = dict(zip((span.start for span in blocks_by_slot[slot]), offsets[slot]))
        for item in evidence["slots"][slot]:
            node_id = item["source_nodes"][role]
            element = elements.get(node_id)
            nodes, match = _editable_number_nodes(element) if element is not None and element.tag == W_P else ([], None)
            if match is None or int(match.group(1)) != item["source_question_number"]:
                raise SlotRoutingError("DISPLAY_RENUMBER_PROJECTION_MISMATCH", "imported prefix changed at " + node_id)
            # Change only digit characters, preserving every other text node,
            # run property, picture, OMML/OLE and punctuation byte-for-byte.
            start, end = match.span(1)
            cursor = 0
            inserted = False
            for node in nodes:
                text = node.text or ""
                node_end = cursor + len(text)
                if cursor < end and node_end > start:
                    a, b = max(0, start - cursor), min(len(text), end - cursor)
                    node.text = text[:a] + (str(item["new_number"]) if not inserted else "") + text[b:]
                    inserted = True
                cursor = node_end


def _resolve_content_carrier(document):
    """Return the frozen template's own content table, ignoring imported blocks.

    render_slots runs on the rendered output, which already carries the imported
    source blocks at the body tail. Sources that contain tables therefore make a
    naive "exactly one table" test wrong. The template carrier is identified by
    its own frozen structure instead: row 5 is one merged cell whose first
    paragraph is 一、课堂启动 and whose sixth paragraph is 二、知识回顾.
    """
    candidates = []
    for table in document.element.body.findall(W_TBL):
        rows = table.findall(W_TR)
        if len(rows) <= _KNOWLEDGE_CELL_ROW_INDEX:
            continue
        cells = rows[_KNOWLEDGE_CELL_ROW_INDEX].findall(W_TC_NS)
        if len(cells) != 1:
            continue
        paragraphs = cells[0].findall(W_P)
        if len(paragraphs) <= _KNOWLEDGE_HEADING_INDEX:
            continue
        launch = _paragraph_text(paragraphs[_LAUNCH_HEADING_INDEX])
        heading = _paragraph_text(paragraphs[_KNOWLEDGE_HEADING_INDEX])
        if launch == "一、课堂启动" and heading == "二、知识回顾":
            candidates.append(table)
    return candidates[0] if len(candidates) == 1 else None


def _anchor_module2_end_divider(document, template_type: str) -> dict:
    """Anchor the existing module-2 end divider to page one's safe bottom.

    Deterministic and template-specific: the existing spacer block between the
    知识回顾 heading and its end divider is given one fixed exact line height per
    interval. No paragraph is added, no page geometry (page size, margins,
    footer, fonts, sizes, character spacing) is changed, and the frozen
    template file is never touched -- this runs on the output copy only.
    """
    target = MODULE2_END_DIVIDER_TARGET_Y[template_type]
    spacer_pt = MODULE2_END_DIVIDER_SPACER_PT[template_type]
    table = _resolve_content_carrier(document)
    if table is None:
        raise SlotRoutingError("TEMPLATE_MODULE2_END_UNRESOLVED",
                               "template content carrier is not a single table")
    rows = table.findall(W_TR)
    cells = rows[_KNOWLEDGE_CELL_ROW_INDEX].findall(W_TC_NS)
    paragraphs = cells[0].findall(W_P)
    divider_index = _MODULE2_END_DIVIDER_INDEX[template_type]
    if (divider_index >= len(paragraphs) or
            not _paragraph_text(paragraphs[divider_index]).startswith("~")):
        raise SlotRoutingError("TEMPLATE_MODULE2_END_UNRESOLVED",
                               "module 2 end divider is not at the validated template position")
    spacer_indices = tuple(range(_KNOWLEDGE_HEADING_INDEX + 1, divider_index))
    if not spacer_indices:
        raise SlotRoutingError("TEMPLATE_MODULE2_END_UNRESOLVED",
                               "template has no spacer block before the module 2 end divider")
    for index in spacer_indices:
        _set_exact_interval(paragraphs[index], spacer_pt)
    return {"target_y": target, "spacer_pt": spacer_pt,
            "spacer_intervals": len(spacer_indices),
            "anchor_node": "MODULE2_END_DIVIDER",
            "anchor_paragraph_index": divider_index}


def _set_exact_interval(paragraph, points: float) -> None:
    """Fix the distance from this paragraph's top to the next paragraph's top."""
    pPr = paragraph.find(W_PPR)
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        paragraph.insert(0, pPr)
    spacing = pPr.find(W_SPACING)
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(W + "line", str(int(round(points * 20))))
    spacing.set(W + "lineRule", "exact")
    spacing.set(W + "before", "0")
    spacing.set(W + "after", "0")


def _set_page_break_before(paragraph) -> None:
    """Keep the knowledge-section heading with its page-two content stream."""
    pPr = paragraph.find(W_PPR)
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        paragraph.insert(0, pPr)
    if pPr.find(W + "pageBreakBefore") is not None:
        return
    page_break = OxmlElement("w:pageBreakBefore")
    style = pPr.find(W + "pStyle")
    if style is None:
        pPr.insert(0, page_break)
    else:
        pPr.insert(pPr.index(style) + 1, page_break)


def _place_knowledge_and_later_slots_on_page_two(document, knowledge_anchor) -> dict:
    """Split the one-cell carrier table at module 3 and page-break between tables.

    WPS ignores paragraph-level pageBreakBefore inside this merged template
    carrier row. A body-level page break between two continuation tables is
    honored while retaining the original cover/divider table and all routed
    source nodes. The frozen template itself is never modified.
    """
    cell = knowledge_anchor.getparent()
    row = cell.getparent() if cell is not None else None
    table = row.getparent() if row is not None else None
    body = document.element.body
    if (cell is None or cell.tag != W_TC or row is None or row.tag != W_TR
            or table is None or table.tag != W_TBL or table.getparent() is not body):
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "module-three anchor is not in a top-level template carrier table")
    rows = table.findall(W_TR)
    cells = row.findall(W_TC)
    if row is not rows[-1] or len(cells) != 1:
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "template carrier row is not the unique single-cell final row")
    cell_children = list(cell)
    try:
        split_at = cell_children.index(knowledge_anchor)
    except ValueError as exc:
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "module-three anchor is not a direct cell child") from exc
    trailing = cell_children[split_at:]
    if not trailing or any(node.tag not in (W_P, W_TBL) for node in trailing):
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "module-three continuation contains unsupported carrier nodes")
    preceding_dividers = [index for index, node in enumerate(cell_children[:split_at])
                          if node.tag == W_P and _paragraph_text(node).lstrip().startswith("~")]
    if not preceding_dividers:
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "module-two divider before module three was not found")
    divider_index = preceding_dividers[-1]
    review_titles = [index for index, node in enumerate(cell_children[:divider_index])
                     if node.tag == W_P and _paragraph_text(node) == "二、知识回顾"]
    if not review_titles:
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "module-two heading before its divider was not found")
    review_index = review_titles[-1]
    for node in cell_children[review_index + 1:divider_index]:
        if node.tag != W_P or _paragraph_text(node):
            continue
        p_pr = node.find(W_PPR)
        if p_pr is not None:
            for flag_name in ("pageBreakBefore", "keepNext"):
                flag = p_pr.find(W + flag_name)
                if flag is not None:
                    p_pr.remove(flag)
    # The frozen template contains trailing empty layout paragraphs between
    # the module-two divider and module-three heading. Keeping the full chain
    # at the end of the first table can force that table onto page 2 in WPS;
    # the explicit body-level break then leaves page 2 blank. Remove only
    # plain empty paragraphs in this interval; preserve drawings, fields,
    # bookmarks, equations, or any other semantic/resource-bearing node.
    spacer_nodes = cell_children[divider_index + 1:split_at]
    protected = []
    removable = []
    for node in spacer_nodes:
        if node.tag == W_P and not _paragraph_text(node):
            sensitive = any(local in {"drawing", "pict", "object", "bookmarkStart",
                                      "bookmarkEnd", "fldChar", "instrText", "hyperlink",
                                      "oMath", "oMathPara", "footnoteReference", "endnoteReference"}
                            for local in (child.tag.split("}")[-1] for child in node.iter()))
            (protected if sensitive else removable).append(node)
        else:
            protected.append(node)
    if protected:
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "non-empty or resource-bearing nodes follow the module-two divider")

    continuation = deepcopy(table)
    continuation_rows = continuation.findall(W_TR)
    continuation_row = continuation_rows[-1]
    for old_row in continuation_rows[:-1]:
        continuation.remove(old_row)
    continuation_cells = continuation_row.findall(W_TC)
    if len(continuation_cells) != 1:
        raise SlotRoutingError("TEMPLATE_PAGE_SPLIT_UNRESOLVED",
                               "continuation template row has multiple physical cells")
    continuation_cell = continuation_cells[0]
    for child in list(continuation_cell):
        if child.tag != W_TCPR:
            continuation_cell.remove(child)
    for node in trailing:
        continuation_cell.append(deepcopy(node))
    # Some source lesson titles carry their own pageBreakBefore. The new
    # body-level boundary already starts the complete module-three stream on
    # page 2; keeping that source break makes WPS create a blank page after
    # SaveAs/Reopen. Remove only this now-redundant first-content break.
    continuation_paragraphs = [node for node in continuation_cell if node.tag == W_P]
    first_content = continuation_paragraphs[1] if len(continuation_paragraphs) > 1 else None
    if first_content is not None:
        content_ppr = first_content.find(W_PPR)
        if content_ppr is not None:
            redundant_break = content_ppr.find(W + "pageBreakBefore")
            if redundant_break is not None:
                content_ppr.remove(redundant_break)

    for node in trailing:
        cell.remove(node)
    for node in removable:
        cell.remove(node)

    page_break_paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(W + "type", "page")
    run.append(page_break)
    page_break_paragraph.append(run)
    table_index = list(body).index(table)
    body.insert(table_index + 1, page_break_paragraph)
    body.insert(table_index + 2, continuation)
    return {"status": "SPLIT_AFTER_MODULE2_DIVIDER", "page_break": "BODY_LEVEL",
            "first_table_rows": len(table.findall(W_TR)),
            "continuation_table_rows": len(continuation.findall(W_TR)),
            "moved_nodes": len(trailing),
            "removed_plain_spacers": len(removable)}


def _set_paragraph_text(paragraph, value: str) -> None:
    text_nodes = list(paragraph.iter(W_T))
    if not text_nodes:
        raise SlotRoutingError("TEMPLATE_SLOT_ANCHOR_UNRESOLVED",
                               "template slot anchor has no editable text run")
    text_nodes[0].text = value
    for node in text_nodes[1:]:
        node.text = ""


def _clear_paragraph_numbering(paragraph) -> None:
    word_namespace = W_P[1:].split("}", 1)[0]
    p_pr = paragraph.find("{%s}pPr" % word_namespace)
    if p_pr is None:
        return
    num_pr = p_pr.find("{%s}numPr" % word_namespace)
    if num_pr is not None:
        p_pr.remove(num_pr)


def _set_cover_cell_value(cell, value: str) -> None:
    """Replace the value portion of a known cover cell while keeping its label.

    Template cover labels and formatting are part of the frozen template. The
    metadata projection only changes the final non-empty text node in the
    selected value cell; an empty value cell receives a run using its existing
    paragraph properties.
    """
    value = str(value or "").strip()
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    text_nodes = list(paragraph._p.iter(W_T))
    value_node = next((node for node in reversed(text_nodes)
                       if (node.text or "").strip()), None)
    if value_node is None:
        if text_nodes:
            value_node = text_nodes[-1]
        else:
            paragraph.add_run(value)
            return
    old_text = value_node.text or ""
    # Some template revisions store the field label in the value cell, while
    # generated documents already contain a complete value such as
    # ``重点：...；难点：...``. Preserve only a label-only prefix. Carrying
    # every character before the last colon makes normalization append a
    # second complete objective/difficulty value on each pass.
    label_match = re.match(
        r"^\s*((?:教学目标|学习目标|重点难点|教学重难点|重难点|教学重点|教学难点)\s*[:：]\s*)$",
        old_text,
    )
    label = label_match.group(1) if label_match else ""
    pieces = (label + value).splitlines() or [""]
    run = value_node.getparent()
    if run.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r":
        value_node.text = "；".join(piece for piece in pieces if piece)
        return
    for node in list(run):
        if node.tag == W_T and node is not value_node:
            node.text = ""
    value_node.text = pieces[0]
    insertion_index = run.index(value_node) + 1
    for piece in pieces[1:]:
        run.insert(insertion_index, OxmlElement("w:br"))
        insertion_index += 1
        text_node = OxmlElement("w:t")
        text_node.text = piece
        run.insert(insertion_index, text_node)
        insertion_index += 1


def _fill_cover_metadata(document, template_type: str, metadata: dict) -> None:
    """Project selected job metadata into existing validated template cells."""
    if template_type not in ("1v1", "class"):
        raise SlotRoutingError("TEMPLATE_COVER_METADATA_UNRESOLVED",
                               "unsupported template type: %s" % template_type)
    table = _resolve_cover_table(document, template_type)
    if table is None:
        raise SlotRoutingError("TEMPLATE_COVER_METADATA_UNRESOLVED",
                               "template cover table is missing or ambiguous")
    display = metadata.get("cover_display") or build_cover_display(
        LessonMetadata(metadata.get("objectives", ""), metadata.get("difficulties", ""),
                       "SOURCE", "UNKNOWN"), template_type)
    limit = COVER_DISPLAY_BUDGET[template_type]["max_display_width_units"]
    for key in ("objectives", "difficulties"):
        if any(mark in display[key] for mark in ("\n", "\r", "\t")) or display_width_units(display[key]) > limit:
            raise SlotRoutingError("COVER_DISPLAY_BUDGET_EXCEEDED", "cover display is outside the fixed budget")
    metadata = {**metadata, "objectives": display["objectives"], "difficulties": display["difficulties"]}
    required_rows = 4
    required_columns = 4 if template_type == "1v1" else 3
    if len(table.rows) < required_rows or len(table.columns) < required_columns:
        raise SlotRoutingError("TEMPLATE_COVER_METADATA_UNRESOLVED",
                               "template cover table dimensions do not match %s" % template_type)
    # PDF row borders of the original frozen templates, measured through WPS.
    # Their atLeast minima are smaller than the actual rendered geometry.
    # Lock the original rendered heights, rather than shrinking to the minima.
    for row_index, height in zip((2, 3), _COVER_ROW_HEIGHTS_PT[template_type]):
        row = table.rows[row_index]
        if row.height is None:
            raise SlotRoutingError("TEMPLATE_COVER_METADATA_UNRESOLVED", "cover row lacks fixed height")
        row.height = Pt(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    if template_type == "1v1":
        # The one-to-one cover has explicit subject, grade, topic, and course
        # type cells. Labels remain in their original runs.
        coordinates = {
            "subject": (0, 0), "grade": (0, 2), "topic": (1, 1),
            "handout_type": (1, 3), "objectives": (2, 1),
            "difficulty": (3, 1),
        }
        values = {key: metadata.get(key, "") for key in coordinates}
        values["difficulty"] = metadata.get("difficulties", metadata.get("difficulty", ""))
    else:
        # The class cover has no separate course-type cell. Keep all selected
        # information visible in its existing class-topic value cell.
        topic = str(metadata.get("topic") or "").strip()
        handout_type = str(metadata.get("handout_type") or "").strip()
        topic_value = topic
        if handout_type:
            topic_value = (topic + "（" + handout_type + "）") if topic else handout_type
        coordinates = {
            "grade": (0, 0), "subject": (0, 2), "topic": (1, 1),
            # The frozen class template contains example text from a math
            # lesson in these fields. No selected-job metadata supplies them,
            # so blank the stale defaults instead of shipping false content.
            "objectives": (2, 1), "difficulty": (3, 1),
        }
        values = {"grade": metadata.get("grade", ""),
                  "subject": metadata.get("subject", ""), "topic": topic_value,
                  "objectives": metadata.get("objectives", ""),
                  "difficulty": metadata.get("difficulties", metadata.get("difficulty", ""))}

    for key, (row, column) in coordinates.items():
        try:
            cell = table.cell(row, column)
        except IndexError as exc:
            raise SlotRoutingError("TEMPLATE_COVER_METADATA_UNRESOLVED",
                                   "cover cell %s is missing for %s" % (key, template_type)) from exc
        _set_cover_cell_value(cell, values[key])


def _resolve_cover_table(document, template_type: str):
    """Identify the unique frozen cover table by stable, unfilled labels.

    Objective/difficulty and subject/grade cells may already have been
    populated by the XML composer before the shared post-render normalizer
    runs. The cover's course/topic labels remain fixed and therefore identify
    the same table across both pre- and post-population states.
    """
    if template_type not in ("1v1", "class"):
        return None

    def compact(value: str) -> str:
        return re.sub(r"\s+", "", value or "")

    candidates = []
    for table in document.tables:
        if len(table.rows) < 4:
            continue
        required_columns = 4 if template_type == "1v1" else 3
        if len(table.columns) < required_columns:
            continue
        row1 = [compact(cell.text) for cell in table.rows[1].cells]
        if template_type == "1v1":
            matches = "授课主题" in row1[0] and "课程类型" in row1[2]
        else:
            matches = "班课主题" in row1[0]
        if matches:
            candidates.append(table)
    return candidates[0] if len(candidates) == 1 else None


def _find_anchor_paragraphs(document, plan: SlotRoutingPlan) -> dict[str, object]:
    anchors = dict(zip(SLOT_ORDER, plan.template_anchors))
    matches: dict[str, list[object]] = {slot: [] for slot in SLOT_ORDER}
    body = document.element.body
    # Keep the element objects themselves alive in this set. Storing only
    # hash(cell) lets lxml release a proxy and reuse its address while walking
    # merged-cell aliases, which can make distinct cells collide and hide all
    # three anchors.
    seen_cells = set()
    for cell in body.iter(W_TC):
        # Merged cells may be surfaced repeatedly by high-level wrappers; only
        # one physical w:tc in the package is an insertion target.
        if cell in seen_cells:
            continue
        seen_cells.add(cell)
        for paragraph in cell.findall(W_P):
            text = _paragraph_text(paragraph)
            for slot, anchor in anchors.items():
                if text == anchor:
                    matches[slot].append(paragraph)
    missing_or_ambiguous = {slot: len(items) for slot, items in matches.items() if len(items) != 1}
    if missing_or_ambiguous:
        raise SlotRoutingError(
            "TEMPLATE_SLOT_ANCHOR_UNRESOLVED",
            "template %s slot anchor counts are not unique: %s" %
            (plan.template_type, missing_or_ambiguous),
        )
    return {slot: items[0] for slot, items in matches.items()}


def render_slots(
    source_doc: str,
    plan: SlotRoutingPlan,
    output_path: str,
    *,
    cover_metadata: dict | None = None,
    display_renumbering: dict | None = None,
    source_role: str = "source",
    render_minimal_fn: Callable = None,
) -> SlotRenderResult:
    """Use the frozen importer, then place each slot stream at its template anchor."""
    if render_minimal_fn is None:
        from renderer_xml_minimal import render_minimal
        render_minimal_fn = render_minimal
    output = Path(output_path).resolve()
    if output.exists():
        raise FileExistsError("slot composer requires a fresh staging output")
    blocks_by_slot = {slot: tuple(plan.slots[slot]) for slot in SLOT_ORDER}
    if display_renumbering is None:
        display_renumbering = build_display_renumbering({source_role: plan})
    flat_blocks: list[BlockSpan] = [span for slot in SLOT_ORDER for span in blocks_by_slot[slot]]
    total = sum(len(items) for items in blocks_by_slot.values())
    excluded_ids = set(getattr(plan, "cover_metadata_blocks", ()))
    record_ids = {str(record.get("block_id")) for record in plan.block_records}
    cover_image_blocks = [
        BlockSpan(str(record["block_id"]), str(record["block_id"]))
        for record in plan.block_records
        if str(record.get("block_id")) in excluded_ids
        and record.get("kind") == "paragraph"
    ]
    cover_image_ids = {span.start for span in cover_image_blocks}
    if (not total or total + len(excluded_ids) != len(plan.block_records)
            or not excluded_ids <= record_ids
            or len(cover_image_ids) != len(cover_image_blocks)):
        raise SlotRoutingError("SLOT_ROUTING_DUPLICATE_OR_MISSING_BLOCK",
                               "slot streams plus projected cover metadata do not partition the source")
    ids = [span.start for span in flat_blocks]
    if len(ids) != len(set(ids)) or set(ids) & excluded_ids:
        raise SlotRoutingError("SLOT_ROUTING_DUPLICATE_OR_MISSING_BLOCK",
                               "slot streams duplicate or include projected cover metadata blocks")
    if set(ids) | excluded_ids != record_ids:
        raise SlotRoutingError("SLOT_ROUTING_DUPLICATE_OR_MISSING_BLOCK",
                               "slot streams leave source blocks outside routed content and cover metadata")

    imported_blocks = flat_blocks + cover_image_blocks
    rendered = render_minimal_fn(source_doc, str(plan.template_path), imported_blocks,
                                 str(output), plan.target)
    try:
        document = Document(str(output))
        if cover_metadata is not None:
            _fill_cover_metadata(document, plan.template_type, cover_metadata)
        # Presentation-layer layout anchor. Runs before the slot payload is
        # inserted, so it can only touch the template's own spacer block.
        module2_end_divider_anchor = _anchor_module2_end_divider(document, plan.template_type)
        anchors = _find_anchor_paragraphs(document, plan)
        body = document.element.body
        insert_at = plan.target.body_child_index
        sectpr_index = next((i for i, child in enumerate(body) if child.tag == W_SECTPR), len(body))
        import_total = total + len(cover_image_blocks)
        if insert_at < 0 or insert_at + import_total > sectpr_index:
            raise SlotRoutingError("TEMPLATE_SLOT_PAYLOAD_UNRESOLVED",
                                   "rendered body tail does not match the slot plan")
        payloads = list(body)[insert_at:insert_at + import_total]
        if any(element.tag not in (W_P, W_TBL) for element in payloads):
            raise SlotRoutingError("TEMPLATE_SLOT_PAYLOAD_UNRESOLVED",
                                   "renderer tail contains a non-paragraph/table block")
        slot_payloads = payloads[:total]
        cover_payloads = payloads[total:]
        if len(cover_payloads) != len(cover_image_blocks):
            raise SlotRoutingError("TEMPLATE_SLOT_PAYLOAD_UNRESOLVED",
                                   "renderer did not import each cover image occurrence")
        blip_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
        imagedata_tag = "{urn:schemas-microsoft-com:vml}imagedata"
        for span, element in zip(cover_image_blocks, cover_payloads):
            if (element.tag != W_P or not any(
                    node.tag in (blip_tag, imagedata_tag) for node in element.iter())):
                raise SlotRoutingError(
                    "COVER_IMAGE_PROJECTION_UNRESOLVED",
                    "cover image %s did not import as a complete image paragraph" % span.start,
                )
        offsets: dict[str, list[object]] = {}
        cursor = 0
        for slot in SLOT_ORDER:
            count = len(blocks_by_slot[slot])
            offsets[slot] = slot_payloads[cursor:cursor + count]
            cursor += count
        if cursor != len(slot_payloads):
            raise SlotRoutingError("TEMPLATE_SLOT_PAYLOAD_UNRESOLVED",
                                   "renderer imported a different number of physical blocks")

        _apply_display_renumbering(offsets, blocks_by_slot, display_renumbering, source_role)

        for element in payloads:
            body.remove(element)
        if cover_payloads:
            cover_table = _resolve_cover_table(document, plan.template_type)
            if cover_table is None:
                raise SlotRoutingError(
                    "TEMPLATE_COVER_METADATA_UNRESOLVED",
                    "cannot locate the unique cover table for source metadata images",
                )
            cover_table_element = cover_table._tbl
            if cover_table_element.getparent() is not body:
                raise SlotRoutingError(
                    "TEMPLATE_COVER_METADATA_UNRESOLVED",
                    "cover metadata table is not a direct body child",
                )
            insertion_index = body.index(cover_table_element)
            for offset, element in enumerate(cover_payloads):
                body.insert(insertion_index + offset, element)
        for slot in SLOT_ORDER:
            anchor = anchors[slot]
            # The 1v1 template shares the class template's printed knowledge
            # heading. Keep the frozen template file intact, but project the
            # user-approved template-specific label into this output copy.
            _set_paragraph_text(anchor, plan.slot_labels[slot])
            for element in offsets[slot]:
                anchor.addnext(element)
                anchor = element

        # WPS does not honor paragraph pageBreakBefore inside this merged
        # carrier row. Split the rendered copy after the frozen module-two
        # divider so module three and later routed slots start in a new
        # body-level table on page two.
        if plan.template_type == "1v1":
            page_layout = _place_knowledge_and_later_slots_on_page_two(
                document, anchors["knowledge"])
        else:
            # The frozen class template already places module three on page 2.
            # Duplicating its carrier table would add an unwanted blank page.
            page_layout = {"status": "NATIVE_TEMPLATE_PAGE_TWO",
                           "page_break": "TEMPLATE_FLOW"}

        output.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(prefix=".slot-composer-", suffix=".docx",
                                         dir=str(output.parent))
        os.close(fd)
        try:
            document.save(temp_name)
            package_report = validate_package(temp_name)
            if not package_report.get("valid"):
                raise SlotRoutingError("PACKAGE_VALIDATION_FAILED",
                                       "; ".join(package_report.get("errors", [])[:5]))
            os.replace(temp_name, output)
        except Exception:
            try:
                os.unlink(temp_name)
            except OSError:
                pass
            raise
        return SlotRenderResult(
            output_path=str(output),
            inserted_nodes=import_total,
            slot_nodes={slot: len(offsets[slot]) for slot in SLOT_ORDER},
            resource_report=rendered.resource_report,
            package_report=package_report,
            display_renumbering=display_renumbering,
            module2_end_divider_anchor=module2_end_divider_anchor,
            page_layout=page_layout,
        )
    except Exception:
        try:
            output.unlink(missing_ok=True)
        except OSError:
            pass
        raise
