"""Fixed cover projection and canonical display-number safety regressions."""
import hashlib
from dataclasses import replace
from pathlib import Path
import sys
from unittest.mock import Mock

from docx import Document
from docx.oxml import OxmlElement
from docx.enum.table import WD_ROW_HEIGHT_RULE

from lesson_metadata import LessonMetadata, build_cover_display, display_width_units
from product_integrity import validate_product_integrity
from semantic_facade import SemanticSnapshot
from slot_router import build_slot_routing_plan, _visible_question_number
from struct_doc import read_struct_doc_bytes, W_P, W_T, W_TBL
from struct_nodes import NodeIndex
from template_block_plan import resolve_template
from template_slot_composer import (build_display_renumbering, render_slots,
                                    _find_anchor_paragraphs, _fill_cover_metadata,
                                    _editable_number_nodes, _apply_display_renumbering,
                                    _anchor_module2_end_divider,
                                    MODULE2_END_DIVIDER_TARGET_Y, MODULE2_END_DIVIDER_SPACER_PT)
from renderer_xml_minimal import BlockSpan
from slot_router import SlotRoutingError
import pytest

sys.path.insert(0, str(Path(__file__).parent / "v09_fallback_runtime"))
from handout import _run_hidden_process


def training_plan(tmp_path, role, *, cross_reference=False, field=False):
    doc = Document()
    units = []
    for group, numbers in enumerate(((1, 2, 3, 4), (1, 2, 3))):
        seq = len(doc.paragraphs)
        doc.add_paragraph('题型%02d 方法训练' % (group + 1))
        section = 's%d' % group
        units.append(dict(id=section, role='section', spans=[['b%d' % seq]*2], parent=None))
        for order, number in enumerate(numbers):
            seq = len(doc.paragraphs)
            p = doc.add_paragraph()
            p.add_run(str(number))
            p.add_run('. 题目%d-%d：年份2026；子问(1)(2)；①；公式2。' % (group, order))
            if cross_reference and group == 0 and order == 0:
                p.add_run('参见第5题。')
            if field and group == 0 and order == 0:
                p._p.append(OxmlElement('w:fldChar'))
            if role == 'teacher':
                doc.add_paragraph('【答案】解答%d-%d' % (group, order))
            end = len(doc.paragraphs)-1
            units.append(dict(id='q%d' % seq, role='question_group',
                              spans=[['b%d' % seq, 'b%d' % end]], parent=section))
    source = tmp_path/(role+'.docx')
    doc.save(source)
    payload = source.read_bytes()
    struct = read_struct_doc_bytes(payload, name=source.name)
    snapshot = SemanticSnapshot(hashlib.sha256(payload).hexdigest(), source.name,
                                struct, NodeIndex(struct), units)
    return build_slot_routing_plan(source, '1v1', snapshot=snapshot)


def test_full_source_metadata_and_deterministic_clause_selection():
    meta = LessonMetadata('掌握方程解法。规范书写；'+ '较长目标'*60,
                          '难点：综合应用；重点：方程解法；'+ '次要说明'*60,
                          'SOURCE', 'KNOWLEDGE_POINT_PRESENT')
    original = (meta.full_objectives, meta.full_difficulties)
    for template in ('1v1', 'class'):
        result = build_cover_display(meta, template)
        assert result == build_cover_display(meta, template)
        assert result['objectives'].startswith('掌握方程解法')
        assert result['difficulties'] == '重点：方程解法；难点：综合应用'
        assert result['cover_metadata_compacted']
        for key in ('objectives', 'difficulties'):
            assert '\n' not in result[key]
            assert display_width_units(result[key]) <= result['budget']['max_display_width_units']
    assert original == (meta.full_objectives, meta.full_difficulties)


def test_long_source_difficulty_keeps_both_labels_and_records_ellipsis():
    meta = LessonMetadata('目标'*150, '重点：'+ '方法'*150+'。难点：'+'应用'*150,
                          'SOURCE','NO_KNOWLEDGE_POINT')
    result = build_cover_display(meta,'class')
    assert '重点：' in result['difficulties'] and '难点：' in result['difficulties']
    assert '…' in result['difficulties'] and '…' in result['objectives']
    assert display_width_units(result['difficulties']) <= 68
    assert len(meta.full_difficulties) > 500


def test_generated_short_rule_retains_method_and_capability_targets():
    meta = LessonMetadata('完整目标'*30,'完整难点'*30,'TRAINING_TYPE_HEADINGS',
                          'NO_KNOWLEDGE_POINT',('物质的构成',))
    result=build_cover_display(meta,'class')
    assert '题型方法' in result['objectives'] and '规范分析与解答' in result['objectives']
    assert '重点：' in result['difficulties'] and '难点：' in result['difficulties']
    assert result['compaction_methods']['objectives']=='SHORT_OFFLINE_RULE'
    assert display_width_units('中文')==4
    assert display_width_units('123')==3


def test_cover_projection_locks_existing_rows_without_font_or_margin_changes():
    for template in ('1v1','class'):
        doc=Document(str(resolve_template(template)[0]))
        margins=[(s.top_margin,s.bottom_margin,s.left_margin,s.right_margin) for s in doc.sections]
        cells=[doc.tables[0].cell(index,1) for index in (2,3)]
        fonts=[[node.xml for node in cell._tc.xpath('.//w:rFonts | .//w:sz')] for cell in cells]
        _fill_cover_metadata(doc,template,dict(objectives='目标'*100,difficulties='重点：方法。难点：应用。'))
        expected=(23.9,23.9) if template=='1v1' else (25.7,27.8)
        assert expected==tuple(doc.tables[0].rows[index].height.pt for index in (2,3))
        assert all(doc.tables[0].rows[index].height_rule==WD_ROW_HEIGHT_RULE.EXACTLY for index in (2,3))
        assert fonts==[[node.xml for node in cell._tc.xpath('.//w:rFonts | .//w:sz')] for cell in cells]
        assert margins==[(s.top_margin,s.bottom_margin,s.left_margin,s.right_margin) for s in doc.sections]


def test_canonical_pair_numbering_restarts_without_changing_route_or_content(tmp_path):
    plans={role:training_plan(tmp_path,role) for role in ('teacher','student')}
    signatures={role:plan.training_question_routes for role,plan in plans.items()}
    evidence=build_display_renumbering(plans)
    assert evidence['status']=='DISPLAY_RENUMBER_APPLIED'
    assert [row['source_question_number'] for row in evidence['slots']['knowledge']]==[1,1]
    expected={'knowledge':[1,2], 'immediate':[1,2], 'final':[1,2,3]}
    for role,plan in plans.items():
        output=tmp_path/(role+'-out.docx')
        render_slots(str(plan.source_path),plan,str(output),display_renumbering=evidence,source_role=role)
        doc=Document(str(output))
        anchors=_find_anchor_paragraphs(doc,replace(plan,template_anchors=tuple(plan.slot_labels[slot] for slot in ('knowledge','immediate','final'))))
        body_children=list(doc.element.body)
        tables=[node for node in body_children if node.tag==W_TBL]
        assert len(tables)==2
        first_table_index=body_children.index(tables[0])
        assert body_children[first_table_index+1].tag==W_P
        assert body_children[first_table_index+1].xpath('.//w:br[@w:type="page"]')
        first_text=''.join(node.text or '' for node in tables[0].iter(W_T))
        second_text=''.join(node.text or '' for node in tables[1].iter(W_T))
        assert '二、知识回顾' in first_text and '知识精讲' not in first_text
        assert '知识精讲' in second_text
        last_row=tables[0].findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')[-1]
        first_cell=last_row.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')[0]
        first_cell_paragraphs=[node for node in first_cell if node.tag==W_P]
        assert ''.join(t.text or '' for t in first_cell_paragraphs[-1].iter(W_T)).strip().startswith('~')
        knowledge_ppr=anchors['knowledge'].find('{%s}pPr' % W_P[1:].split('}',1)[0])
        assert knowledge_ppr is None or knowledge_ppr.find('{%s}pageBreakBefore' % W_P[1:].split('}',1)[0]) is None
        for slot,anchor in anchors.items():
            paragraphs=[]
            for following in anchor.itersiblings():
                if following in anchors.values():
                    break
                if following.tag==W_P:
                    paragraphs.append(following)
            numbers=[_visible_question_number(''.join(t.text or '' for t in p.iter(W_T)))
                     for p in paragraphs]
            assert [number for number in numbers if number is not None]==expected[slot]
        text='\n'.join(''.join(t.text or '' for t in p.iter(W_T)) for p in doc.element.body.iter(W_P))
        for group,count in ((0,4),(1,3)):
            for order in range(count):
                assert text.count('题目%d-%d' % (group,order))==1
        assert text.count('子问(1)(2)；①；公式2。')==7
        assert text.count('【答案】')==(7 if role=='teacher' else 0)
        assert '题型01' in text and '题型02' in text
        for heading in ('一、课堂启动','二、知识回顾','知识精讲','即时训练','五、归纳总结','六、巩固练习'):
            assert heading in text
        assert validate_product_integrity(plan.source_path,output,plan=plan)['accepted']
    assert signatures=={role:plan.training_question_routes for role,plan in plans.items()}


def test_cross_reference_disables_both_presentations(tmp_path):
    plans={role:training_plan(tmp_path,role,cross_reference=(role=='teacher')) for role in ('teacher','student')}
    evidence=build_display_renumbering(plans)
    assert evidence['status']=='DISPLAY_RENUMBER_SKIPPED_CROSS_REFERENCE'
    assert evidence['warning']
    assert all(item['source_question_number']==item['new_number']
               for rows in evidence['slots'].values() for item in rows)
    for role,plan in plans.items():
        output=tmp_path/(role+'-out.docx')
        render_slots(str(plan.source_path),plan,str(output),display_renumbering=evidence,source_role=role)
        text=''.join(t.text or '' for t in Document(output).element.body.iter(W_T))
        assert '4. 题目0-3' in text


def test_field_prefix_fails_safe_and_ordinary_routed_lessons_can_renumber(tmp_path):
    plan=training_plan(tmp_path,'teacher',field=True)
    assert build_display_renumbering({'teacher':plan})['status']=='DISPLAY_RENUMBER_PARTIAL_UNSAFE_STRUCTURE'
    ordinary=replace(training_plan(tmp_path,'teacher'),knowledge_point_status='KNOWLEDGE_POINT_PRESENT')
    assert build_display_renumbering({'teacher':ordinary})['status']=='DISPLAY_RENUMBER_APPLIED'
    ordinary_student=replace(training_plan(tmp_path,'student'),knowledge_point_status='KNOWLEDGE_POINT_PRESENT')
    output=tmp_path/'ordinary.docx';render_slots(str(ordinary_student.source_path),ordinary_student,str(output))
    assert '2. 题目0-3' in ''.join(t.text or '' for t in Document(output).element.body.iter(W_T))


def test_all_safe_prefixes_and_fragmented_digits_leave_other_numbers_untouched():
    for prefix in ('1.', '2．', '3、', '4)'):
        p=Document().add_paragraph(prefix+' question')._p
        assert _editable_number_nodes(p)[1] is not None
    for prefix in ('(1)', '①', '题型01', '例1', '变式1-1', '2026年'):
        p=Document().add_paragraph(prefix+' question')._p
        assert _editable_number_nodes(p)[1] is None
    doc=Document();p=doc.add_paragraph();p.add_run('1').bold=True;p.add_run('2');p.add_run('．内容12、(1)与①')
    evidence=dict(status='DISPLAY_RENUMBER_APPLIED',slots=dict(knowledge=[dict(source_nodes={'teacher':'b0'},source_question_number=12,new_number=3)],immediate=[],final=[]))
    _apply_display_renumbering(dict(knowledge=[p._p],immediate=[],final=[]),dict(knowledge=[BlockSpan('b0','b0')],immediate=[],final=[]),evidence,'teacher')
    assert p.text=='3．内容12、(1)与①'
    assert p.runs[0].bold is True and len(p.runs)==3


def test_caller_cannot_expand_the_validated_display_budget():
    doc=Document(str(resolve_template('class')[0]))
    with pytest.raises(SlotRoutingError) as raised:
        _fill_cover_metadata(doc,'class',{'cover_display':dict(objectives='目标'*100,difficulties='重点',budget={'max_display_width_units':9999})})
    assert raised.value.reason_code=='COVER_DISPLAY_BUDGET_EXCEEDED'


def test_atomic_table_number_is_not_edited(tmp_path):
    source=tmp_path/'table.docx';doc=Document();doc.add_paragraph('题型01 方法训练')
    doc.add_table(rows=1,cols=1).cell(0,0).text='1. 表内问题';doc.save(source)
    payload=source.read_bytes();struct=read_struct_doc_bytes(payload,name=source.name)
    units=[dict(id='s0',role='section',spans=[['b0','b0']],parent=None),
           dict(id='q0',role='question_group',spans=[['b1.r0c0.n0','b1.r0c0.n0']],parent='s0')]
    snapshot=SemanticSnapshot(hashlib.sha256(payload).hexdigest(),source.name,struct,NodeIndex(struct),units)
    plan=build_slot_routing_plan(source,'1v1',snapshot=snapshot)
    assert build_display_renumbering({'source':plan})['status']=='NOT_APPLICABLE'
    assert Document(source).tables[0].cell(0,0).text=='1. 表内问题'


# --- Module 2 end-divider first-page anchor ---------------------------------

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _content_cell_paragraphs(doc):
    return doc.tables[0]._tbl.findall(W+'tr')[5].findall(W+'tc')[0].findall(W+'p')


def test_module2_end_divider_targets_are_template_specific_and_distinct():
    assert MODULE2_END_DIVIDER_TARGET_Y['1v1']!=MODULE2_END_DIVIDER_TARGET_Y['class']
    assert MODULE2_END_DIVIDER_SPACER_PT['1v1']!=MODULE2_END_DIVIDER_SPACER_PT['class']
    assert MODULE2_END_DIVIDER_SPACER_PT['1v1']==19.65
    assert MODULE2_END_DIVIDER_SPACER_PT['class']==22.0
    # The anchor must sit above the footer band of its own template, with a
    # measured clearance rather than a physical page-edge placement.
    for template,footer_top in (('1v1',792.25),('class',769.90)):
        assert MODULE2_END_DIVIDER_TARGET_Y[template]<footer_top
        assert footer_top-MODULE2_END_DIVIDER_TARGET_Y[template]>=20.0


def test_module2_end_anchor_changes_only_the_existing_post_heading_spacers():
    for template in ('1v1','class'):
        doc=Document(str(resolve_template(template)[0]))
        cell=doc.tables[0]._tbl.findall(W+'tr')[5].findall(W+'tc')[0]
        before=[p.xml for p in cell.findall(W+'p')]
        before_body=[child.tag for child in doc.element.body]
        before_margins=[(s.top_margin,s.bottom_margin,s.left_margin,s.right_margin,
                         s.page_width,s.page_height) for s in doc.sections]
        before_rows=[row.height.pt if row.height else None for row in doc.tables[0].rows]
        evidence=_anchor_module2_end_divider(doc,template)
        after=[p.xml for p in cell.findall(W+'p')]
        assert len(before)==len(after), 'no paragraph may be added or removed'
        # Only paragraphs between the module-2 heading and end divider change.
        changed=[index for index,(old,new) in enumerate(zip(before,after)) if old!=new]
        divider = 19 if template == '1v1' else 16
        assert changed==list(range(6,divider)), changed
        assert evidence['spacer_intervals']==divider-6
        assert evidence['target_y']==MODULE2_END_DIVIDER_TARGET_Y[template]
        assert evidence['spacer_pt']==MODULE2_END_DIVIDER_SPACER_PT[template]
        assert evidence['anchor_node']=='MODULE2_END_DIVIDER'
        # Page geometry, cover rows and the heading text stay untouched.
        assert before_body==[child.tag for child in doc.element.body]
        assert before_margins==[(s.top_margin,s.bottom_margin,s.left_margin,s.right_margin,
                                 s.page_width,s.page_height) for s in doc.sections]
        assert before_rows==[row.height.pt if row.height else None for row in doc.tables[0].rows]
        paragraphs=_content_cell_paragraphs(doc)
        assert ''.join(t.text or '' for t in paragraphs[5].iter(W+'t'))=='二、知识回顾'
        assert ''.join(t.text or '' for t in paragraphs[0].iter(W+'t'))=='一、课堂启动'


def test_module2_end_anchor_uses_exact_line_rule_and_is_idempotent():
    for template in ('1v1','class'):
        doc=Document(str(resolve_template(template)[0]))
        _anchor_module2_end_divider(doc,template)
        first=_content_cell_paragraphs(doc)
        snapshot=[p.xml for p in first]
        divider = 19 if template == '1v1' else 16
        for index in range(6,divider):
            spacing=first[index].find(W+'pPr').find(W+'spacing')
            assert spacing.get(W+'lineRule')=='exact'
            assert spacing.get(W+'before')=='0' and spacing.get(W+'after')=='0'
            assert int(spacing.get(W+'line'))==int(round(MODULE2_END_DIVIDER_SPACER_PT[template]*20))
        _anchor_module2_end_divider(doc,template)
        assert snapshot==[p.xml for p in _content_cell_paragraphs(doc)]


def test_module2_end_anchor_refuses_a_moved_or_missing_template_heading():
    doc=Document(str(resolve_template('1v1')[0]))
    paragraphs=_content_cell_paragraphs(doc)
    heading=paragraphs[5]
    for node in list(heading.iter(W+'t')):
        heading.remove(node) if node.getparent() is heading else node.getparent().remove(node)
    with pytest.raises(SlotRoutingError) as raised:
        _anchor_module2_end_divider(doc,'1v1')
    assert raised.value.reason_code=='TEMPLATE_MODULE2_END_UNRESOLVED'


def test_module2_end_anchor_ignores_imported_blocks_that_carry_tables():
    """render_slots runs after the source blocks are imported at the body tail.

    A source containing a table therefore makes the rendered document hold more
    than one table; the anchor must still resolve the frozen template carrier
    instead of failing closed on a table count.
    """
    doc=Document(str(resolve_template('1v1')[0]))
    imported=doc.add_table(rows=1,cols=1)
    imported.cell(0,0).text='1. 导入题目'
    evidence=_anchor_module2_end_divider(doc,'1v1')
    assert evidence['target_y']==MODULE2_END_DIVIDER_TARGET_Y['1v1']
    assert evidence['spacer_pt']==MODULE2_END_DIVIDER_SPACER_PT['1v1']
    assert imported.cell(0,0).text=='1. 导入题目'
    paragraphs=_content_cell_paragraphs(doc)
    for index in range(6,19):
        spacing=paragraphs[index].find(W+'pPr').find(W+'spacing')
        assert spacing.get(W+'lineRule')=='exact'
        assert int(spacing.get(W+'line'))==int(round(MODULE2_END_DIVIDER_SPACER_PT['1v1']*20))


def test_render_slots_reports_module2_end_anchor_evidence(tmp_path):
    plan=training_plan(tmp_path,'teacher')
    for template in ('1v1','class'):
        output=tmp_path/('anchored-%s.docx'%template)
        candidate=build_slot_routing_plan(plan.source_path, template)
        result=render_slots(str(candidate.source_path),candidate,str(output))
        assert result.module2_end_divider_anchor['target_y']==MODULE2_END_DIVIDER_TARGET_Y[template]
        assert result.module2_end_divider_anchor['spacer_pt']==MODULE2_END_DIVIDER_SPACER_PT[template]


def test_v09_powershell_child_is_hidden_without_changing_capture_contract(monkeypatch):
    import handout

    runner = Mock(return_value=object())
    monkeypatch.setattr(handout.subprocess, "run", runner)
    monkeypatch.setattr(handout.os, "name", "nt")
    result = _run_hidden_process(["powershell.exe", "-File", "test.ps1"],
                                 capture_output=True, timeout=37, check=False)
    assert result is runner.return_value
    args, kwargs = runner.call_args
    assert args[0] == ["powershell.exe", "-File", "test.ps1"]
    assert kwargs["capture_output"] is True
    assert kwargs["timeout"] == 37
    assert kwargs["check"] is False
    assert kwargs["creationflags"] & handout.subprocess.CREATE_NO_WINDOW
    assert kwargs["startupinfo"].wShowWindow == handout.subprocess.SW_HIDE
