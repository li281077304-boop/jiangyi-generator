# -*- coding: utf-8 -*-
"""Narrow cover-metadata integration checks for both XML templates."""
from __future__ import annotations

import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path

from docx import Document

from package_validator import validate_package
from template_block_plan import resolve_template
from template_slot_composer import _fill_cover_metadata


class CoverMetadataProjectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="cover-metadata-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def _project(self, template_type):
        template, _ = resolve_template(template_type)
        output = self.root / (template_type + ".docx")
        document = Document(str(template))
        _fill_cover_metadata(document, template_type, {
            "subject": "化学", "grade": "高一", "topic": "复习讲义",
            "handout_type": "专题复习",
        })
        document.save(str(output))
        self.assertTrue(validate_package(str(output))["valid"])
        with zipfile.ZipFile(template) as before, zipfile.ZipFile(output) as after:
            before_parts = {name for name in before.namelist() if not name.endswith("/")}
            after_parts = {name for name in after.namelist() if not name.endswith("/")}
            self.assertEqual(before_parts, after_parts)
            rel_parts = [name for name in before_parts if name.endswith(".rels")]
            def relationships(package, part):
                root = ET.fromstring(package.read(part))
                return sorted(tuple(sorted(element.attrib.items())) for element in root)
            self.assertEqual({name: relationships(before, name) for name in rel_parts},
                             {name: relationships(after, name) for name in rel_parts})
        return Document(str(output))

    def test_one_to_one_selected_metadata_replaces_cover_values(self):
        document = self._project("1v1")
        table = document.tables[0]
        self.assertIn("化学", table.cell(0, 0).text)
        self.assertNotIn("数学", table.cell(0, 0).text)
        self.assertIn("高一", table.cell(0, 2).text)
        self.assertNotIn("高三", table.cell(0, 2).text)
        self.assertEqual(table.cell(1, 1).text.strip(), "复习讲义")
        self.assertEqual(table.cell(1, 3).text.strip(), "专题复习")

    def test_class_selected_metadata_uses_existing_cover_cells(self):
        document = self._project("class")
        table = document.tables[0]
        self.assertIn("高一", table.cell(0, 0).text)
        self.assertNotIn("五年级", table.cell(0, 0).text)
        self.assertIn("化学", table.cell(0, 2).text)
        self.assertNotIn("数学", table.cell(0, 2).text)
        self.assertEqual(table.cell(1, 1).text.strip(), "复习讲义（专题复习）")
        self.assertNotIn("因数与倍数", table.cell(1, 1).text)
        self.assertEqual(table.cell(2, 1).text.strip(), "")
        self.assertEqual(table.cell(3, 1).text.strip(), "")


if __name__ == "__main__":
    unittest.main()
