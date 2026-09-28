# -*- coding: utf-8 -*-
"""
xml_engine.py —— Flash 1.2 不启动 Word 的纯 XML 讲义生成引擎
（基于 Flash 1.1 的 XML 实验脚本抽出模块化版本）

能力：
  - scan(): body 直接子级扫描（段落+表格混排）
  - split_ideal(): 题型清单式文档分区（抽例题后按原讲义分区分块）
  - fill_cover(): 封面表格字段填充（科目/年级/授课主题/课程类型/教学目标/重难点）
  - set_page_break_before(): 强制某锚点从新页开始
  - image_attach_fix(): 边界前纯图片装饰段并入下一块（必考题型图跟下面走）
  - fill_blocks(): 复制 children 区间到模板锚点（含表格+段落），通用关系迁移（图片/OLE/超链接）

不依赖 Word COM，速度预期 1~3 秒/份（vs 1.1 COM 的 26~45 秒，约 15~30 倍提速）。
"""
import os
import re
import copy
import time
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 命名空间常量
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P_TAG = "{%s}p" % W_NS
TBL_TAG = "{%s}tbl" % W_NS
T_TAG = "{%s}t" % W_NS


# ===================== 扫描 =====================

def scan(docx_path):
    """
    body 直接子元素扫描：返回 (paras, children)
      paras:   [(text, element)] 仅段落（body 直接子级 w:p，按文档流顺序）
      children: [element] body 直接子级（w:p 与 w:tbl 混排）
      段落编号 = children 中 w:p 的顺序（1-based）
    """
    doc = Document(docx_path)
    children = [el for el in doc.element.body.iterchildren()
                if el.tag in (P_TAG, TBL_TAG)]
    paras = []
    for el in children:
        if el.tag == P_TAG:
            txt = "".join(t.text or "" for t in el.iter(T_TAG)).strip()
            paras.append((txt, el))
    return paras, children


# ===================== 拆分（ideal：抽例题+重编号后的题型清单文档） =====================

_QNUM = re.compile(r'^\s*(\d{1,3})\s*[.．、]\s*[（(]\s*(20\d\d|\d{2}-)')
_TYPE = re.compile(r'^【题型\s*\d+')
_SECT = re.compile(r'提升专练|真题感知')


def split_ideal(paras, block_size=10):
    """
    ideal 分区：
      知识精讲 = 1 .. 第一个【题型】前
      即时训练 = 题型1 .. 提升专练/真题感知 标题前（无分区则到第 N 题）
      巩固练习 = 提升专练/真题感知 .. 文末
    """
    first_type = next((i for i, (t, _) in enumerate(paras, 1) if _TYPE.match(t)), None)
    if first_type is None:
        return None
    qnums = [i for i, (t, _) in enumerate(paras, 1) if _QNUM.match(t)]
    sect_pos = next(
        (i for i, t in enumerate((t for t, _ in paras[first_type:]), first_type + 1)
         if _SECT.search(t)),
        None)
    last = len(paras)
    if sect_pos:
        return [("知识精讲", 1, first_type - 1),
                ("即时训练", first_type, sect_pos - 1),
                ("六、巩固练习", sect_pos, last)]
    cut = qnums[min(block_size, len(qnums)) - 1]
    return [("知识精讲", 1, first_type - 1),
            ("即时训练", first_type, cut),
            ("六、巩固练习", cut + 1, last)]


# ===================== 图片归属修正 =====================

def _has_picture(el):
    return (el.find('.//{%s}drawing' % W_NS) is not None or
            el.find('.//{%s}pict' % W_NS) is not None)

def _is_decor(t):
    if not t:
        return True
    return all(c in '/\\~-_=*#·.。、，| ' for c in t)

def image_attach_fix(blocks, paras):
    """边界前紧邻的纯图片装饰段并入下一块（返回新块列表，不修改入参）"""
    blocks = [list(b) for b in blocks]
    for k in range(len(blocks) - 1):
        mk, s, e = blocks[k]
        mk2, s2, e2 = blocks[k + 1]
        cut = None
        lo = max(s2 - 3, s)
        for j in range(s2 - 1, lo - 1, -1):
            t, el = paras[j - 1]
            if _has_picture(el) and _is_decor(t):
                cut = j
            else:
                break
        if cut is not None and cut - 1 >= s - 1:
            blocks[k] = [mk, s, cut - 1]
            blocks[k + 1] = [mk2, cut, e2]
    return [tuple(b) for b in blocks]


# ===================== 通用关系迁移 =====================

def migrate_rels(new_el, src_part, tpl_part):
    """
    遍历元素中的 r:embed / r:id / r:link（a:blip、v:imagedata、o:OLEObject、hyperlink），
    把对应关系（图片/oleObject/超链接）注册到模板并改写 rId。
    数学文档有 1116 个 OLE 对象，不迁移 OLE 关系会导致 Word 显示乱码/挤一坨。
    """
    migrated = 0
    for el in new_el.iter():
        for attr in ("{%s}embed" % R_NS, "{%s}id" % R_NS, "{%s}link" % R_NS):
            rid = el.get(attr)
            if not rid or rid not in src_part.rels:
                continue
            rel = src_part.rels[rid]
            target = rel.target_part
            if target is None:
                continue
            new_rid = tpl_part.relate_to(target, rel.reltype)
            el.set(attr, new_rid)
            migrated += 1
    return migrated


# ===================== 锚点定位 =====================

def find_anchors(tpl, markers):
    """在模板（XML 全量，含文本框/表格内段落）中找锚点段落元素"""
    anchors = {}
    for p_el in tpl.element.body.iter(P_TAG):
        txt = "".join(t.text or "" for t in p_el.iter(T_TAG)).strip()
        for m in markers:
            if txt.startswith(m) and m not in anchors:
                anchors[m] = p_el
                break
    return anchors


# ===================== 封面表格填充 =====================

def set_cell(tbl, row, col, text):
    """修改模板表格某单元格文本（保留第一个 run 格式）"""
    try:
        cell = tbl.rows[row].cells[col]
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(text)
    except Exception:
        pass

def fill_cover(tpl, meta):
    """
    meta 字典应含：subject / grade / topic / handout_type / objectives / difficulties
    模板第一个表格（Tables.Item(1)）：6x4 布局
      [1,1]科目 / [1,3]年级 / [2,2]授课主题 / [2,4]课程类型
      [3,2]教学目标 / [4,2]重点难点
    """
    if not tpl.tables:
        return False
    t = tpl.tables[0]
    set_cell(t, 0, 0, f"科 目：{meta.get('subject', '')}")
    set_cell(t, 0, 2, meta.get("grade", ""))
    set_cell(t, 1, 1, meta.get("topic", ""))
    set_cell(t, 1, 3, meta.get("handout_type", ""))
    set_cell(t, 2, 1, meta.get("objectives", ""))
    set_cell(t, 3, 1, meta.get("difficulties", ""))
    return True


# ===================== 分页符 =====================

def set_page_break_before(p_el):
    """给段落加 pageBreakBefore 属性（强制新页开始）"""
    pPr = p_el.find('{%s}pPr' % W_NS)
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p_el.insert(0, pPr)
    if pPr.find('{%s}pageBreakBefore' % W_NS) is None:
        pb = OxmlElement('w:pageBreakBefore')
        # 按 OOXML schema：插在 pStyle 之后
        idx = 0
        for i, c in enumerate(list(pPr)):
            if c.tag == '{%s}pStyle' % W_NS:
                idx = i + 1
        pPr.insert(idx, pb)


# ===================== 主流程 =====================

def build(src_path, tpl_path, out_path, meta=None, block_size=10, markers=None):
    """
    一站式生成讲义 XML 成品：
      meta 字典（封面字段）：subject/grade/topic/handout_type/objectives/difficulties
      markers: 锚点标记列表（默认 ["知识精讲", "即时训练", "六、巩固练习"]）
    返回 (blocks_used, stats_dict)
    """
    if markers is None:
        markers = ["知识精讲", "即时训练", "六、巩固练习"]
    if meta is None:
        meta = {}

    t_all = time.perf_counter()
    paras, children = scan(src_path)
    blocks = split_ideal(paras, block_size=block_size)
    blocks = image_attach_fix(blocks, paras)

    p_ci = [i for i, el in enumerate(children) if el.tag == P_TAG]

    src_doc = Document(src_path)
    tpl = Document(tpl_path)

    anchors = find_anchors(tpl, markers)
    if "知识精讲" in anchors:
        set_page_break_before(anchors["知识精讲"])
    fill_cover(tpl, meta)

    copied = migrated = 0
    for marker, s, e in blocks:
        anchor = anchors.get(marker)
        if anchor is None:
            continue
        ci_lo, ci_hi = p_ci[s - 1], p_ci[e - 1]
        cur = anchor
        for el in children[ci_lo:ci_hi + 1]:
            new_el = copy.deepcopy(el)
            migrated += migrate_rels(new_el, src_doc.part, tpl.part)
            cur.addnext(new_el)
            cur = new_el
            copied += 1

    tpl.save(out_path)
    elapsed = time.perf_counter() - t_all
    return blocks, {"copied": copied, "migrated": migrated, "elapsed_s": elapsed,
                    "out_size_kb": os.path.getsize(out_path) / 1024}


# ===================== CLI =====================

if __name__ == "__main__":
    import sys, json
    if len(sys.argv) < 4:
        print("用法: python xml_engine.py <源.docx> <模板.docx> <输出.docx> [topic]")
        sys.exit(1)
    meta = {"subject": "数学", "grade": "八年级", "topic": sys.argv[4] if len(sys.argv) > 4 else "实数",
            "handout_type": "专题复习", "objectives": "", "difficulties": ""}
    blocks, stats = build(sys.argv[1], sys.argv[2], sys.argv[3], meta=meta)
    print(json.dumps({"blocks": blocks, **stats}, ensure_ascii=False, indent=2))