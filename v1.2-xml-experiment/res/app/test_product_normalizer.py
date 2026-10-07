# -*- coding: utf-8 -*-
"""Shared post-render product normalization contract tests."""
from __future__ import annotations

import tempfile
import unittest
import hashlib
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from package_validator import validate_package
from product_normalizer import normalize_product_docx
from slot_router import SlotRoutingError
from template_block_plan import resolve_template
from template_slot_composer import (MODULE2_END_DIVIDER_TARGET_Y,
                                    _anchor_module2_end_divider,
                                    _fill_cover_metadata)


class ProductNormalizerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="product-normalizer-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def _normalize(self, template_type, name, metadata):
        template, _ = resolve_template(template_type)
        source = self.root / (name + "-renderer.docx")
        output = self.root / (name + "-normalized.docx")
        source.write_bytes(Path(template).read_bytes())
        evidence = normalize_product_docx(
            source, output, template_type=template_type, metadata=metadata)
        self.assertEqual(evidence["status"], "NORMALIZED")
        self.assertEqual(evidence["source_sha256"], hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertNotEqual(evidence["source_sha256"], evidence["output_sha256"])
        self.assertTrue(validate_package(str(output))["valid"])
        return source, output, evidence

    def test_both_templates_populate_metadata_and_clear_old_examples(self):
        for template_type in ("1v1", "class"):
            with self.subTest(template=template_type):
                _source, output, evidence = self._normalize(
                    template_type, template_type,
                    {"subject": "物理", "grade": "九年级", "topic": "电学复习",
                     "handout_type": "专题训练", "objectives": "",
                     "difficulties": "", "cover_display": {"objectives": "",
                                                                     "difficulties": ""}})
                document = Document(str(output))
                cover = document.tables[0]
                self.assertEqual(cover.cell(2, 1).text.strip(), "")
                self.assertEqual(cover.cell(3, 1).text.strip(), "")
                flattened = "".join(node.text or "" for node in
                                     document.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
                for stale in ("因数与倍数", "质数和合数", "掌握重难点题型"):
                    self.assertNotIn(stale, flattened)
                self.assertEqual(evidence["cover_fields"]["objectives"], "EMPTY_WITH_WARNING")
                self.assertEqual(evidence["module2_end_divider_anchor"]["target_y"],
                                 MODULE2_END_DIVIDER_TARGET_Y[template_type])

    def test_normalization_is_semantically_idempotent(self):
        template_type = "1v1"
        _source, once, first = self._normalize(
            template_type, "once",
            {"subject": "物理", "grade": "九年级", "topic": "电学复习",
             "handout_type": "专题训练", "objectives": "目标", "difficulties": "重点",
             "cover_display": {"objectives": "目标", "difficulties": "重点"}})
        twice = self.root / "twice.docx"
        second = normalize_product_docx(
            once, twice, template_type=template_type,
            metadata={"subject": "物理", "grade": "九年级", "topic": "电学复习",
                      "handout_type": "专题训练", "objectives": "目标", "difficulties": "重点",
                      "cover_display": {"objectives": "目标", "difficulties": "重点"}})
        first_doc, second_doc = Document(str(once)), Document(str(twice))
        self.assertEqual(len(first_doc.paragraphs), len(second_doc.paragraphs))
        self.assertEqual([p.text for p in first_doc.paragraphs],
                         [p.text for p in second_doc.paragraphs])
        self.assertEqual(first_doc.tables[0].cell(2, 1).text,
                         second_doc.tables[0].cell(2, 1).text)
        self.assertEqual(first_doc.tables[0].cell(3, 1).text,
                         second_doc.tables[0].cell(3, 1).text)
        self.assertEqual(first["module2_end_divider_anchor"],
                         second["module2_end_divider_anchor"])

    def test_output_hash_is_stable_when_renderer_zip_timestamps_change(self):
        template, _ = resolve_template("1v1")
        source_a = self.root / "renderer-a.docx"
        source_b = self.root / "renderer-b.docx"
        source_a.write_bytes(Path(template).read_bytes())
        with zipfile.ZipFile(template, "r") as original, zipfile.ZipFile(source_b, "w") as changed:
            for entry in original.infolist():
                stable = zipfile.ZipInfo(entry.filename, date_time=(2020, 1, 1, 0, 0, 0))
                stable.compress_type = entry.compress_type
                changed.writestr(stable, original.read(entry.filename))
        self.assertNotEqual(hashlib.sha256(source_a.read_bytes()).hexdigest(),
                            hashlib.sha256(source_b.read_bytes()).hexdigest())
        metadata = {"subject": "物理", "grade": "九年级", "topic": "专题",
                    "handout_type": "复习讲义", "objectives": "目标",
                    "difficulties": "重点", "cover_display": {
                        "objectives": "目标", "difficulties": "重点"}}
        output_a = self.root / "normalized-a.docx"
        output_b = self.root / "normalized-b.docx"
        first = normalize_product_docx(source_a, output_a, template_type="1v1",
                                       metadata=metadata)
        second = normalize_product_docx(source_b, output_b, template_type="1v1",
                                        metadata=metadata)
        self.assertEqual(first["output_sha256"], second["output_sha256"])

    def test_existing_renderer_difficulty_labels_are_not_copied_as_cell_prefix(self):
        for template_type in ("1v1", "class"):
            with self.subTest(template=template_type):
                template, _ = resolve_template(template_type)
                source = self.root / (template_type + "-renderer-populated.docx")
                output = self.root / (template_type + "-normalized-populated.docx")
                metadata = {
                    "subject": "物理", "grade": "九年级", "topic": "电学专题",
                    "handout_type": "复习讲义", "objectives": "掌握专题方法",
                    "difficulties": "重点：实验方法；难点：综合分析",
                    "cover_display": {"objectives": "掌握专题方法",
                                      "difficulties": "重点：实验方法；难点：综合分析"},
                }
                document = Document(str(template))
                _fill_cover_metadata(document, template_type, metadata)
                document.save(str(source))
                normalize_product_docx(source, output, template_type=template_type,
                                       metadata=metadata)
                cover = Document(str(output)).tables[0]
                self.assertEqual(cover.cell(3, 1).text.strip(),
                                 "重点：实验方法；难点：综合分析")

    def test_normalizer_resolves_cover_after_composer_clears_optional_identity_fields(self):
        for template_type in ("1v1", "class"):
            with self.subTest(template=template_type):
                template, _ = resolve_template(template_type)
                source = self.root / (template_type + "-composer-output.docx")
                output = self.root / (template_type + "-composer-normalized.docx")
                metadata = {"subject": "", "grade": "", "topic": "专题",
                            "handout_type": "复习讲义", "objectives": "",
                            "difficulties": "", "cover_display": {
                                "objectives": "", "difficulties": ""}}
                document = Document(str(template))
                _fill_cover_metadata(document, template_type, metadata)
                document.save(str(source))
                evidence = normalize_product_docx(
                    source, output, template_type=template_type, metadata=metadata)
                self.assertEqual(evidence["status"], "NORMALIZED")
                self.assertTrue(validate_package(str(output))["valid"])

    def test_only_structurally_bound_source_title_is_shortened(self):
        source_title = "九年级上学期物理期末复习（易错精选60题27大考点）"
        metadata = {"subject": "物理", "grade": "九年级", "topic": "易错题精选",
                    "handout_type": "期末复习", "objectives": "目标", "difficulties": "重点",
                    "cover_display": {"objectives": "目标", "difficulties": "重点"}}
        for template_type in ("1v1", "class"):
            with self.subTest(template=template_type):
                template, _ = resolve_template(template_type)
                source = self.root / (template_type + "-long-source-title.docx")
                output = self.root / (template_type + "-short-source-title.docx")
                document = Document(str(template))
                anchors = [paragraph for cell in document.element.body.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
                    for paragraph in cell.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
                    if "".join(node.text or "" for node in paragraph.iter(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
                    .strip() == "知识精讲&例题讲解"]
                self.assertEqual(len(anchors), 1)
                if template_type == "1v1":
                    from template_slot_composer import _set_paragraph_text
                    _set_paragraph_text(anchors[0], "知识精讲")
                    knowledge_label = "知识精讲"
                else:
                    knowledge_label = "知识精讲&例题讲解"
                title = OxmlElement("w:p")
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.text = source_title
                run.append(text)
                title.append(run)
                anchors[0].addnext(title)
                immediate_anchors = [paragraph for cell in document.element.body.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
                    for paragraph in cell.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
                    if "".join(node.text or "" for node in paragraph.iter(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
                    .strip() == "即时训练"]
                self.assertEqual(len(immediate_anchors), 1)
                duplicate = OxmlElement("w:p")
                duplicate_run = OxmlElement("w:r")
                duplicate_text = OxmlElement("w:t")
                duplicate_text.text = source_title
                duplicate_run.append(duplicate_text)
                duplicate.append(duplicate_run)
                immediate_anchors[0].addnext(duplicate)
                document.save(str(source))

                evidence = normalize_product_docx(
                    source, output, template_type=template_type, metadata=metadata)
                rendered = Document(str(output))
                knowledge = next(paragraph for cell in rendered.element.body.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
                    for paragraph in cell.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
                    if "".join(node.text or "" for node in paragraph.iter(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
                    .strip() == knowledge_label)
                actual_title = knowledge.getnext()
                actual_text = "".join(node.text or "" for node in actual_title.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")).strip()
                self.assertEqual(actual_text, "易错题精选")
                self.assertEqual(evidence["source_title_projection"], {
                    "status": "APPLIED", "source_title": source_title,
                    "display_title": "易错题精选"})
                duplicate_after = [paragraph for cell in rendered.element.body.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
                    for paragraph in cell.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
                    if "".join(node.text or "" for node in paragraph.iter(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
                    .strip() == source_title]
                self.assertEqual(len(duplicate_after), 1)
                second = self.root / (template_type + "-short-source-title-twice.docx")
                second_evidence = normalize_product_docx(
                    output, second, template_type=template_type, metadata=metadata)
                self.assertEqual(second_evidence["source_title_projection"]["status"], "UNCHANGED")
                second_doc = Document(str(second))
                self.assertIn("易错题精选", [paragraph.text.strip() for paragraph in second_doc.paragraphs]
                              + ["".join(node.text or "" for node in p.iter(
                                  "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")).strip()
                                 for p in second_doc.element.body.iter(
                                     "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")])

    def test_malformed_renderer_output_fails_closed_and_cleans_temp(self):
        source = self.root / "invalid.docx"
        output = self.root / "invalid-normalized.docx"
        source.write_bytes(b"not a docx")
        with self.assertRaises(Exception):
            normalize_product_docx(source, output, template_type="1v1", metadata={})
        self.assertFalse(output.exists())
        self.assertEqual(list(self.root.glob("*.normalize-*.docx")), [])

    def test_ambiguous_cover_and_content_carriers_fail_closed(self):
        template, _ = resolve_template("1v1")
        source = self.root / "ambiguous-cover.docx"
        output = self.root / "ambiguous-cover-normalized.docx"
        document = Document(str(template))
        document.element.body.append(deepcopy(document.tables[0]._tbl))
        document.save(str(source))
        with self.assertRaises(SlotRoutingError) as raised:
            normalize_product_docx(source, output, template_type="1v1", metadata={})
        self.assertEqual(raised.exception.reason_code, "TEMPLATE_COVER_METADATA_UNRESOLVED")
        self.assertFalse(output.exists())

        carrier_doc = Document(str(template))
        carrier_doc.element.body.append(deepcopy(carrier_doc.tables[0]._tbl))
        with self.assertRaises(SlotRoutingError) as raised:
            _anchor_module2_end_divider(carrier_doc, "1v1")
        self.assertEqual(raised.exception.reason_code, "TEMPLATE_MODULE2_END_UNRESOLVED")


if __name__ == "__main__":
    unittest.main()
