# -*- coding: utf-8 -*-
"""User-facing metadata warning and normalization evidence persistence."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from job_service import JobService


class JobServiceProductMetadataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="job-product-meta-")
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        source = root / "sample 教师版.docx"
        doc = Document()
        doc.add_paragraph("1. 样例题")
        doc.save(source)
        self.service = JobService(root / "results", runtime_root=root / "runtime")
        self.job = self.service.create_inputs(
            [(source.name, source.read_bytes())],
            {"docx_mode": "auto", "template_type": "1v1"},
        )
        self.service.start_job(self.job["job_id"])

    def test_missing_metadata_warning_and_normalization_survive_job_reads(self):
        warning = "原文未提供教学目标，已留空，请使用前补充。"
        details = {"status": "PARTIAL", "objectives": "", "difficulties": "重点",
                   "objectives_reason": "NO_RELIABLE_OBJECTIVES_SOURCE"}
        self.service.update_lesson_metadata(self.job["job_id"], details, warning=warning)
        normalization = {"status": "NORMALIZED", "version": "V1.2_PRODUCT_NORMALIZATION_V1",
                         "outputs": {"teacher": {"output_sha256": "abc123"}}}
        self.service.update_product_normalization(self.job["job_id"], normalization)

        persisted = self.service.get(self.job["job_id"])
        self.assertEqual(persisted["lesson_metadata"], details)
        self.assertEqual(persisted["warnings"], ["封面字段：" + warning])
        self.assertEqual(persisted["product_normalization"], normalization)

        # Re-evaluation replaces the prior cover warning instead of duplicating it.
        self.service.update_lesson_metadata(self.job["job_id"], details, warning=warning)
        self.assertEqual(self.service.get(self.job["job_id"])["warnings"],
                         ["封面字段：" + warning])

    def test_direct_two_docx_pair_resolves_to_one_ordinary_job(self):
        teacher = Document()
        teacher.add_paragraph("1. 测试题")
        student = Document()
        student.add_paragraph("1. 测试题")
        teacher_path = Path(self.temp.name) / "专题 教师版.docx"
        student_path = Path(self.temp.name) / "专题 学生版.docx"
        teacher.save(teacher_path)
        student.save(student_path)

        job = self.service.create_inputs(
            [(student_path.name, student_path.read_bytes()),
             (teacher_path.name, teacher_path.read_bytes())],
            {"docx_mode": "auto", "template_type": "1v1"},
        )

        self.assertEqual(job["status"], "queued")
        self.assertEqual(job["input_version"], "TEACHER_AND_STUDENT")
        self.assertEqual(job["items"][0]["topic"], "专题")
        self.assertEqual(Path(job["teacher_source_path"]).name, "input-2.docx")
        self.assertEqual(Path(job["student_source_path"]).name, "input-1.docx")
        self.assertIsNone(job["parent_job_id"])


if __name__ == "__main__":
    unittest.main()
