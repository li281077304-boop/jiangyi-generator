# -*- coding: utf-8 -*-
"""Package-level regression tests for the v1.2 XML block importer.

The OLE case intentionally uses a dummy embedded binary.  It proves OPC
relationship and binary preservation only; it is not a MathType rendering UAT.
"""
import hashlib
import os
import sys
import tempfile
import unittest
import zipfile

from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(__file__))
from block_importer import BlockImporter, PIC_NS, R_NS, V_NS, W_NS, WP_NS
from package_validator import validate_package


W = "{%s}" % W_NS
R = "{%s}" % R_NS
WP = "{%s}" % WP_NS
PIC = "{%s}" % PIC_NS

OLE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"

PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDAT\x08\xd7c\xf8\xcf"
    b"\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00IEND\xaeB`\x82"
)


def body_blocks(doc):
    return [el for el in doc.element.body.iterchildren() if el.tag in (W + "p", W + "tbl")]


def append_blocks(doc, elements):
    sect_pr = doc.element.body.sectPr
    for element in elements:
        sect_pr.addprevious(element)


def paragraph_text(element):
    return "".join(t.text or "" for t in element.iter(W + "t"))


def numbered(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = etree.SubElement(p_pr, W + "numPr")
    etree.SubElement(num_pr, W + "ilvl").set(W + "val", "0")
    etree.SubElement(num_pr, W + "numId").set(W + "val", str(num_id))


def add_hyperlink(paragraph, rid, text):
    hyperlink = etree.SubElement(paragraph._p, W + "hyperlink")
    hyperlink.set(R + "id", rid)
    run = etree.SubElement(hyperlink, W + "r")
    etree.SubElement(run, W + "t").text = text


class BlockImporterTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def save_output(self, doc, name="out.docx"):
        path = os.path.join(self.tmp.name, name)
        doc.save(path)
        return path

    def test_text_table_style_numbering_and_external_hyperlink(self):
        source = Document()
        source.styles.add_style("ImportedParagraph", WD_STYLE_TYPE.PARAGRAPH)
        source.styles.add_style("ImportedTable", WD_STYLE_TYPE.TABLE)
        source.styles.add_style("SourceBase", WD_STYLE_TYPE.PARAGRAPH)
        source.styles.add_style("ImportedCharacter", WD_STYLE_TYPE.CHARACTER)
        source.styles["ImportedParagraph"].base_style = source.styles["SourceBase"]
        etree.SubElement(source.styles["ImportedParagraph"]._element,
                         W + "link").set(W + "val", "ImportedCharacter")
        etree.SubElement(source.styles["ImportedCharacter"]._element,
                         W + "link").set(W + "val", "ImportedParagraph")
        paragraph = source.add_paragraph("Styled source text", style="ImportedParagraph")
        num_id = source.part.numbering_part.element.findall(W + "num")[0].get(W + "numId")
        source_abstract_id = next(
            n for n in source.part.numbering_part.element.findall(W + "num")
            if n.get(W + "numId") == num_id).find(W + "abstractNumId").get(W + "val")
        source_abstract = next(
            n for n in source.part.numbering_part.element.findall(W + "abstractNum")
            if n.get(W + "abstractNumId") == source_abstract_id)
        source_num_format = source_abstract.find(".//" + W + "numFmt").get(W + "val")
        style_ppr = source.styles["ImportedParagraph"]._element.get_or_add_pPr()
        style_num_pr = etree.SubElement(style_ppr, W + "numPr")
        etree.SubElement(style_num_pr, W + "ilvl").set(W + "val", "0")
        etree.SubElement(style_num_pr, W + "numId").set(W + "val", num_id)
        numbered(paragraph, num_id)
        table = source.add_table(rows=1, cols=1)
        table.style = "ImportedTable"
        table.cell(0, 0).text = "Table payload"
        link_p = source.add_paragraph("Link: ")
        hyperlink_rid = source.part.relate_to("https://example.test/lesson", RT.HYPERLINK, is_external=True)
        add_hyperlink(link_p, hyperlink_rid, "external")

        destination = Document()
        # force a meaningful numbering-ID collision with a valid target definition
        target_num = next(n for n in destination.part.numbering_part.element.findall(W + "num")
                          if n.get(W + "numId") == num_id)
        target_abstract_id = target_num.find(W + "abstractNumId").get(W + "val")
        target_abstract = next(n for n in destination.part.numbering_part.element.findall(W + "abstractNum")
                               if n.get(W + "abstractNumId") == target_abstract_id)
        target_num_fmt = target_abstract.find(".//" + W + "numFmt")
        target_num_fmt.set(W + "val", "lowerLetter")
        importer = BlockImporter(source, destination)
        copied, report = importer.import_blocks(body_blocks(source))
        append_blocks(destination, copied)
        output = self.save_output(destination)

        self.assertEqual([], report["unsupported"])
        self.assertTrue(any(s.get(W + "styleId") == "ImportedParagraph"
                            for s in destination.part._styles_part.element.findall(W + "style")))
        self.assertTrue(any(s.get(W + "styleId") == "ImportedTable"
                            for s in destination.part._styles_part.element.findall(W + "style")))
        self.assertTrue(any(s.get(W + "styleId") == "SourceBase"
                            for s in destination.part._styles_part.element.findall(W + "style")))
        self.assertTrue(any(s.get(W + "styleId") == "ImportedCharacter"
                            for s in destination.part._styles_part.element.findall(W + "style")))
        copied_num = copied[0].find(".//" + W + "numId").get(W + "val")
        self.assertNotEqual(num_id, copied_num)
        self.assertTrue(any(n.get(W + "numId") == copied_num
                            for n in destination.part.numbering_part.element.findall(W + "num")))
        imported_style = next(s for s in destination.part._styles_part.element.findall(W + "style")
                              if s.get(W + "styleId") == "ImportedParagraph")
        self.assertEqual(copied_num,
                         imported_style.find(".//" + W + "numId").get(W + "val"))
        target_abstract_after = next(n for n in destination.part.numbering_part.element.findall(W + "abstractNum")
                                     if n.get(W + "abstractNumId") == target_abstract_id)
        self.assertEqual("lowerLetter",
                         target_abstract_after.find(".//" + W + "numFmt").get(W + "val"))
        imported_abstract_id = next(n for n in destination.part.numbering_part.element.findall(W + "num")
                                    if n.get(W + "numId") == copied_num).find(W + "abstractNumId").get(W + "val")
        imported_abstract = next(n for n in destination.part.numbering_part.element.findall(W + "abstractNum")
                                 if n.get(W + "abstractNumId") == imported_abstract_id)
        self.assertEqual(source_num_format,
                         imported_abstract.find(".//" + W + "numFmt").get(W + "val"))
        self.assertIn("Table payload", "".join(paragraph_text(e) for e in copied))
        hyperlink = copied[-1].find(".//" + W + "hyperlink")
        self.assertIn(hyperlink.get(R + "id"), destination.part.rels)
        self.assertTrue(destination.part.rels[hyperlink.get(R + "id")].is_external)
        with zipfile.ZipFile(output) as package:
            self.assertIn("word/document.xml", package.namelist())
            self.assertIn("word/numbering.xml", package.namelist())
        validation = validate_package(output)
        self.assertTrue(validation["valid"], validation)

    def test_images_vml_recursive_part_and_dummy_ole_binary(self):
        image_path = os.path.join(self.tmp.name, "pixel.png")
        with open(image_path, "wb") as image:
            image.write(PNG)

        source = Document()
        picture = source.add_paragraph().add_run().add_picture(image_path)
        blip = picture._inline.graphic.graphicData.pic.blipFill.blip
        image_rid = blip.get(R + "embed")
        external_image_rid = source.part.relate_to(
            "https://example.test/remote-preview.png", RT.IMAGE, is_external=True)
        blip.set(R + "link", external_image_rid)
        # A VML preview reference exercises the old-Word shape route too.
        vml_p = source.add_paragraph("VML preview")
        pict = etree.SubElement(vml_p._p, W + "pict")
        shape = etree.SubElement(pict, "{%s}shape" % V_NS)
        image_data = etree.SubElement(shape, "{%s}imagedata" % V_NS)
        image_data.set(R + "id", image_rid)

        embedding_blob = b"not-a-real-mathtype-ole\x00payload"
        embedding = Part(PackURI("/word/embeddings/oleObject1.bin"),
                         "application/vnd.openxmlformats-officedocument.oleObject",
                         embedding_blob, source.part.package)
        nested = Part(PackURI("/customXml/importer-nested.xml"), "application/xml",
                      b"<fixture>nested relationship</fixture>", source.part.package)
        embedding.load_rel(CUSTOM_REL, nested, "rId1")
        ole_rid = source.part.relate_to(embedding, OLE_REL)
        ole_p = source.add_paragraph("OLE ")
        obj = etree.SubElement(ole_p._p, W + "object")
        ole = etree.SubElement(obj, "{urn:schemas-microsoft-com:office:office}OLEObject")
        ole.set(R + "id", ole_rid)

        # SmartArt-style source references exercise r:dm/lo/qs/cs remapping and
        # a relationship embedded in a copied XML child part.
        diagram_data = Part(PackURI("/word/diagrams/data1.xml"), "application/xml",
                            ("<fixture xmlns:r='%s'><ref r:id='rId51'/></fixture>" %
                             R_NS).encode("utf-8"), source.part.package)
        diagram_nested = Part(PackURI("/word/diagrams/nested.xml"), "application/xml",
                             b"<fixture>nested-diagram-part</fixture>", source.part.package)
        diagram_data.load_rel(CUSTOM_REL, diagram_nested, "rId51")
        layout = Part(PackURI("/word/diagrams/layout1.xml"), "application/xml",
                      b"<layout/>", source.part.package)
        quick_style = Part(PackURI("/word/diagrams/quickStyle1.xml"), "application/xml",
                           b"<quickStyle/>", source.part.package)
        colors = Part(PackURI("/word/diagrams/colors1.xml"), "application/xml",
                      b"<colors/>", source.part.package)
        diagram_rel_ids = etree.SubElement(
            source.add_paragraph("Diagram refs")._p,
            "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds")
        for attribute, target, reltype in (
                ("dm", diagram_data,
                 "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData"),
                ("lo", layout,
                 "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramLayout"),
                ("qs", quick_style,
                 "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramQuickStyle"),
                ("cs", colors,
                 "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramColors")):
            diagram_rel_ids.set(R + attribute,
                                source.part.relate_to(target, reltype))

        destination = Document()
        destination.add_paragraph().add_run().add_picture(image_path)  # docPr id collision
        importer = BlockImporter(source, destination)
        copied, report = importer.import_blocks(body_blocks(source))
        append_blocks(destination, copied)
        output = self.save_output(destination)

        self.assertEqual([], report["unsupported"])
        self.assertGreaterEqual(report["stats"].get("image_relationships_copied", 0), 2)
        self.assertGreaterEqual(report["stats"].get("external_relationships_copied", 0), 1)
        self.assertGreaterEqual(report["stats"].get("vml_image_relationships_copied", 0), 1)
        self.assertGreaterEqual(report["stats"].get("ole_references_copied", 0), 1)
        self.assertGreaterEqual(report["stats"].get("embedding_parts_copied", 0), 1)
        self.assertGreaterEqual(report["stats"].get("part_name_collisions_remapped", 0), 1)
        self.assertGreaterEqual(report["stats"].get("drawing_docPr_ids_remapped", 0), 1)
        self.assertGreaterEqual(report["stats"].get("picture_cNvPr_ids_remapped", 0), 1)
        with zipfile.ZipFile(output) as package:
            embedding_names = [n for n in package.namelist()
                               if n.startswith("word/embeddings/") and not n.endswith(".rels")]
            self.assertEqual(1, len(embedding_names))
            self.assertEqual(hashlib.sha256(embedding_blob).hexdigest(),
                             hashlib.sha256(package.read(embedding_names[0])).hexdigest())
            self.assertIn("customXml/importer-nested.xml", package.namelist())
            embedding_rels = os.path.dirname(embedding_names[0]) + "/_rels/" + \
                os.path.basename(embedding_names[0]) + ".rels"
            self.assertIn(embedding_rels, package.namelist())
            self.assertIn("word/diagrams/data1.xml", package.namelist())
            self.assertIn("word/diagrams/nested.xml", package.namelist())
            self.assertIn("word/diagrams/_rels/data1.xml.rels", package.namelist())
            self.assertTrue(any(n.startswith("word/media/") for n in package.namelist()))
        validation = validate_package(output)
        self.assertTrue(validation["valid"], validation)

    def test_conflicting_style_id_is_imported_under_a_safe_new_id(self):
        source = Document()
        source_style = source.styles.add_style("SameId", WD_STYLE_TYPE.PARAGRAPH)
        source_style.font.bold = True
        source.add_paragraph("source definition", style="SameId")
        destination = Document()
        destination_style = destination.styles.add_style("SameId", WD_STYLE_TYPE.PARAGRAPH)
        destination_style.font.size = Pt(18)

        copied, report = BlockImporter(source, destination).import_blocks(body_blocks(source))
        imported_id = copied[0].find(".//" + W + "pStyle").get(W + "val")
        self.assertEqual([], report["unsupported"])
        self.assertEqual("SameId_imported1", imported_id)
        self.assertTrue(any(s.get(W + "styleId") == imported_id
                            for s in destination.part._styles_part.element.findall(W + "style")))
        self.assertEqual(1, report["stats"].get("style_id_collisions_remapped"))

    def test_template_story_drawing_id_collisions_are_normalized(self):
        image_path = os.path.join(self.tmp.name, "pixel.png")
        with open(image_path, "wb") as image:
            image.write(PNG)
        source = Document()
        source.add_paragraph("Imported content")
        destination = Document()
        destination.add_paragraph().add_run().add_picture(image_path)
        destination.sections[0].header.paragraphs[0].add_run().add_picture(image_path)

        importer = BlockImporter(source, destination)
        copied, report = importer.import_blocks(body_blocks(source))
        append_blocks(destination, copied)
        output = self.save_output(destination)

        self.assertEqual([], report["unsupported"])
        self.assertGreater(report["stats"].get("destination_docPr_ids_normalized", 0), 0)
        self.assertGreater(report["stats"].get("destination_picture_ids_normalized", 0), 0)
        validation = validate_package(output)
        self.assertTrue(validation["valid"], validation)

    def test_custom_xml_data_binding_is_reported_as_unsupported(self):
        source = Document()
        paragraph = source.add_paragraph("Bound content")
        custom_xml = etree.SubElement(paragraph._p, W + "customXml")
        binding = etree.SubElement(custom_xml, W + "dataBinding")
        binding.set(W + "storeItemID", "{fixture-guid}")
        destination = Document()

        _, report = BlockImporter(source, destination).import_blocks(body_blocks(source))

        unsupported_codes = {item["code"] for item in report["unsupported"]}
        self.assertIn("unsupported_custom_xml_wrapper", unsupported_codes)
        self.assertIn("unsupported_custom_xml_data_binding", unsupported_codes)

    def test_bookmark_picture_id_and_section_properties_are_remapped_or_removed(self):
        image_path = os.path.join(self.tmp.name, "pixel.png")
        with open(image_path, "wb") as image:
            image.write(PNG)
        source = Document()
        p = source.add_paragraph("bookmarked")
        start = etree.Element(W + "bookmarkStart")
        start.set(W + "id", "1")
        start.set(W + "name", "same_name")
        end = etree.Element(W + "bookmarkEnd")
        end.set(W + "id", "1")
        p._p.insert(0, start)
        p._p.append(end)
        hyperlink = etree.SubElement(p._p, W + "hyperlink")
        hyperlink.set(W + "anchor", "same_name")
        p.add_run().add_picture(image_path)
        etree.SubElement(p._p.get_or_add_pPr(), W + "sectPr")

        destination = Document()
        existing = destination.add_paragraph("existing")
        old_start = etree.SubElement(existing._p, W + "bookmarkStart")
        old_start.set(W + "id", "1")
        old_start.set(W + "name", "same_name")
        existing.add_run().add_picture(image_path)
        importer = BlockImporter(source, destination)
        copied, report = importer.import_blocks(body_blocks(source))
        append_blocks(destination, copied)

        self.assertEqual([], report["unsupported"])
        self.assertEqual(1, report["stats"].get("source_sectPr_removed"))
        self.assertGreaterEqual(report["stats"].get("bookmark_ids_remapped", 0), 1)
        self.assertGreaterEqual(report["stats"].get("bookmark_names_remapped", 0), 1)
        self.assertGreaterEqual(report["stats"].get("drawing_docPr_ids_remapped", 0), 1)
        self.assertIsNone(copied[0].find(".//" + W + "sectPr"))
        starts = list(copied[0].iter(W + "bookmarkStart"))
        ends = list(copied[0].iter(W + "bookmarkEnd"))
        self.assertEqual(starts[0].get(W + "id"), ends[0].get(W + "id"))
        self.assertEqual(starts[0].get(W + "name"),
                         copied[0].find(".//" + W + "hyperlink").get(W + "anchor"))
        self.assertNotEqual("1", starts[0].get(W + "id"))


if __name__ == "__main__":
    unittest.main()
