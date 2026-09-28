# -*- coding: utf-8 -*-
"""Programmatic build-path checks for the v1.2 XML experiment."""
import os
import sys
import tempfile
import unittest
from unittest import mock

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import xml_engine


class XmlEngineBuildTests(unittest.TestCase):
    def _documents(self, directory):
        source_path = os.path.join(directory, "source.docx")
        template_path = os.path.join(directory, "template.docx")
        output_path = os.path.join(directory, "output.docx")

        source = Document()
        source.add_paragraph("SOURCE FIRST")
        table = source.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "SOURCE TABLE"
        source.add_paragraph("SOURCE LAST")
        source.save(source_path)

        template = Document()
        template.add_paragraph("知识精讲")
        template.save(template_path)
        return source_path, template_path, output_path

    def test_build_uses_block_importer_and_validates_saved_package(self):
        with tempfile.TemporaryDirectory() as directory:
            source, template, output = self._documents(directory)
            blocks = [("知识精讲", 1, 2)]
            with mock.patch.object(xml_engine, "split_ideal", return_value=blocks):
                used, stats = xml_engine.build(source, template, output)

            self.assertEqual(used, blocks)
            self.assertTrue(os.path.isfile(output))
            generated = Document(output)
            body_text = "".join(
                "".join(node.itertext())
                for node in generated.element.body.iter()
                if node.tag.endswith("}t")
            )
            self.assertIn("SOURCE FIRST", body_text)
            self.assertIn("SOURCE TABLE", body_text)
            self.assertIn("SOURCE LAST", body_text)
            self.assertEqual(stats["migration"]["unsupported"], [])
            self.assertTrue(stats["package_validation"]["valid"])

    def test_build_raises_and_removes_output_when_validation_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            source, template, output = self._documents(directory)
            blocks = [("知识精讲", 1, 2)]
            failed_report = {
                "valid": False, "errors": ["fixture dangling relationship"],
                "warnings": [], "stats": {},
            }
            with mock.patch.object(xml_engine, "split_ideal", return_value=blocks), \
                    mock.patch.object(xml_engine, "validate_package", return_value=failed_report):
                with self.assertRaises(xml_engine.PackageValidationError) as error:
                    xml_engine.build(source, template, output)

            self.assertFalse(os.path.exists(output))
            self.assertEqual(
                error.exception.report["package_validation"]["errors"],
                ["fixture dangling relationship"],
            )


if __name__ == "__main__":
    unittest.main()
