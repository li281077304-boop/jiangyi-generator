from pathlib import Path
from unittest.mock import Mock

from docx import Document
import pytest

from xml_degradation import build_degraded_plan, structural_snapshot
from template_slot_composer import render_slots
from product_integrity import validate_product_integrity
from slot_router import SlotRoutingError
import xml_degradation


def make_source(tmp_path, role):
    doc = Document()
    doc.add_paragraph('知识点', style='Heading 1')
    doc.add_paragraph('本专题使用实验测量的方法。')
    doc.add_paragraph('即时训练', style='Heading 1')
    doc.add_paragraph('1. ' + ('教师原稿独有题干' if role == 'teacher' else '学生原稿独有题干'))
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table._cells, ['电压', '1.9', '1.0', '序号', '2.9', '3.0']):
        cell.text = value
    if role == 'teacher':
        doc.add_paragraph('【答案】教师独有解析，不得进入另一份文件。')
    doc.add_paragraph('巩固练习', style='Heading 1')
    doc.add_paragraph('1. 独立完整练习组。')
    path = tmp_path / (role + '.docx')
    doc.save(path)
    return path


@pytest.mark.parametrize('template', ['1v1', 'class'])
def test_preservation_keeps_independent_text_table_order_and_answers(tmp_path, template):
    for role in ('teacher', 'student'):
        source = make_source(tmp_path, role)
        plan, evidence = build_degraded_plan(source, template, preserve_only=True)
        assert evidence['selected_tier'] == 'PRESERVATION'
        output = tmp_path / (role + '-' + template + '.docx')
        render_slots(str(source), plan, str(output))
        rendered = Document(output)
        text = '\n'.join(t.text or '' for t in rendered.element.body.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        own = '教师原稿独有题干' if role == 'teacher' else '学生原稿独有题干'
        peer = '学生原稿独有题干' if role == 'teacher' else '教师原稿独有题干'
        assert text.count(own) == 1 and peer not in text
        assert ('教师独有解析' in text) == (role == 'teacher')
        assert text.index(own) < text.index('电压') < text.index('独立完整练习组')
        assert '1.9' in text and '1.0' in text and '2.9' in text
        assert validate_product_integrity(source, output, plan=plan)['accepted']


def test_heading_rules_follow_semantic_refusal_and_preservation_follows_rules_refusal(tmp_path, monkeypatch):
    source = make_source(tmp_path, 'student')
    doc = Document(source)
    for paragraph in doc.paragraphs:
        paragraph.style = doc.styles['Normal']
    doc.save(source)
    actual = xml_degradation.build_slot_routing_plan

    def refuse_semantic(*args, **kwargs):
        if kwargs.get('canonical_projection_routes') is None and kwargs.get('split_mode') != 'full':
            raise SlotRoutingError('SLOT_ROUTING_AMBIGUOUS', 'cannot prove semantic ownership')
        return actual(*args, **kwargs)

    monkeypatch.setattr(xml_degradation, 'build_slot_routing_plan', refuse_semantic)
    plan, evidence = build_degraded_plan(source, '1v1', snapshot=structural_snapshot(source))
    assert evidence['selected_tier'] == 'RULES'
    assert [entry['tier'] for entry in evidence['attempts']] == ['NAVIGATION', 'SEMANTIC', 'RULES']
    monkeypatch.setattr(xml_degradation, 'heading_projection', Mock(side_effect=SlotRoutingError('SLOT_ROUTING_AMBIGUOUS', 'unknown boundaries')))
    plan, evidence = build_degraded_plan(source, '1v1', snapshot=structural_snapshot(source))
    assert evidence['selected_tier'] == 'PRESERVATION'
    assert [entry['tier'] for entry in evidence['attempts']] == ['NAVIGATION', 'SEMANTIC', 'RULES', 'PRESERVATION']
