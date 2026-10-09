"""Minimal, fail-closed OOXML projection, resource import, and insertion.

This is a B2 experiment only. It maps current StructDoc IDs to elements by
walking the same source document structure in lockstep. It never searches by
paragraph text or uses COM paragraph coordinates.

Supported payloads are whole body paragraphs, whole tables, and paragraph
spans that stay outside tables. The selected raw XML is passed to the adapted
BlockImporter for relationship, package-part, style, numbering, and XML-ID
migration. Unsupported import reports and failed saved-package validation
stop before replacing the requested output.
"""
from __future__ import annotations

from dataclasses import dataclass
import io
from pathlib import Path
import os
import re
import tempfile

from docx import Document

_APP = Path(__file__).resolve().parent
import sys
if str(_APP) not in sys.path:
    sys.path.insert(0, str(_APP))

from struct_doc import (  # noqa: E402
    W, W_P, W_TBL, W_TR, W_TC, W_SECTPR,
    read_struct_doc_bytes,
)
from struct_nodes import NodeIndex  # noqa: E402
from block_importer import BlockImporter  # noqa: E402
from package_validator import validate_package  # noqa: E402

class ProjectionError(ValueError):
    """The requested source content cannot be projected safely."""


@dataclass(frozen=True)
class BlockSpan:
    """A StructDoc leaf range or one atomic top-level table ID."""
    start: str
    end: str


@dataclass(frozen=True)
class TemplateTarget:
    """Insert before this zero-based direct child of template w:body."""
    body_child_index: int


@dataclass(frozen=True)
class RenderResult:
    output_path: str
    inserted_nodes: int
    resource_report: dict
    package_report: dict


def render_minimal(source_doc: str, template_doc: str,
                   blocks: list[BlockSpan], output_path: str,
                   target: TemplateTarget) -> RenderResult:
    """Clone selected source OOXML blocks into an explicit template position.

    Source and template are read without modification. The destination is
    written atomically after all selected nodes pass capability checks.
    """
    src_path, tpl_path, out_path = map(Path, (source_doc, template_doc, output_path))
    if not src_path.is_file() or not tpl_path.is_file():
        raise FileNotFoundError("source and template must be existing DOCX files")
    if out_path.resolve() in (src_path.resolve(), tpl_path.resolve()):
        raise ProjectionError("output path must not alias the source or template")
    if not blocks:
        raise ProjectionError("at least one StructDoc block span is required")
    source_bytes = _read_source_bytes(src_path)
    source_document = Document(io.BytesIO(source_bytes))
    destination_document = Document(str(tpl_path))
    src_body = source_document.element.body
    tpl_body = destination_document.element.body
    struct = read_struct_doc_bytes(source_bytes, name=str(src_path))
    index = NodeIndex(struct)
    node_map, table_ids, top_table_ids, table_leaf_ids = _map_structdoc_to_xml(src_body, struct)
    selected = _select_nodes(blocks, index, node_map, table_ids, top_table_ids,
                             table_leaf_ids, struct)
    bookmark_preflight = _validate_bookmark_scope(src_body, selected)
    permitted_bookmark_drops = bookmark_preflight["bookmark_markers_dropped"]
    for node in selected:
        _validate_payload(node)

    importer = BlockImporter(source_document, destination_document)
    imported, resource_report = importer.import_blocks(selected)
    if resource_report["unsupported"]:
        first = resource_report["unsupported"][0]
        raise ProjectionError("resource import unsupported: %s" % first)
    actual_bookmark_drops = sum(resource_report["stats"].get(key, 0) for key in (
        "unmatched_bookmark_starts_dropped", "unmatched_bookmark_ends_dropped"))
    if actual_bookmark_drops != permitted_bookmark_drops:
        raise ProjectionError("bookmark importer drop count did not match preflight")
    if actual_bookmark_drops:
        resource_report["stats"]["standalone_body_bookmark_markers_dropped"] = actual_bookmark_drops
    if bookmark_preflight["dangling_toc_anchors_dropped"]:
        resource_report["stats"]["dangling_toc_anchors_dropped"] = bookmark_preflight[
            "dangling_toc_anchors_dropped"]
    if bookmark_preflight["dangling_internal_anchors_dropped"]:
        resource_report["stats"]["dangling_internal_anchors_dropped"] = bookmark_preflight[
            "dangling_internal_anchors_dropped"]

    insert_at = target.body_child_index
    body_children = list(tpl_body)
    section_tail = next((i for i, child in enumerate(body_children)
                         if child.tag == W_SECTPR), len(body_children))
    if insert_at < 0 or insert_at > section_tail:
        raise ProjectionError("template target is outside direct w:body children")
    for offset, element in enumerate(imported):
        tpl_body.insert(insert_at + offset, element)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=".renderer-", suffix=".docx", dir=str(out_path.parent))
    os.close(fd)
    try:
        destination_document.save(temp_name)
        package_report = validate_package(temp_name)
        if not package_report["valid"]:
            errors = "; ".join(package_report["errors"][:5])
            raise ProjectionError("saved package failed validation: %s" % errors)
        os.replace(temp_name, out_path)
    except Exception:
        try:
            os.unlink(temp_name)
        except OSError:
            pass
        raise
    return RenderResult(str(out_path), len(imported), resource_report, package_report)


def _read_source_bytes(path):
    """Read once so XML and StructDoc parsing share one immutable snapshot."""
    return path.read_bytes()


def _map_structdoc_to_xml(body, doc):
    """Build ID -> exact element map through ordinal structural traversal."""
    mapping = {"body": body}
    table_ids = set()
    top_table_ids = set()
    table_leaf_ids = set()
    body_content = [(i, child) for i, child in enumerate(body)
                    if child.tag != W_SECTPR and child.tag not in (
                        "{%s}bookmarkStart" % W, "{%s}bookmarkEnd" % W)]
    if len(body_content) != len(doc.blocks):
        raise ProjectionError("StructDoc/body top-level block count mismatch")

    def add_table(block, table_el, table_id):
        if table_el.tag != W_TBL:
            raise ProjectionError("StructDoc table ordinal does not identify w:tbl")
        mapping[table_id] = table_el
        table_ids.add(table_id)
        rows = [child for child in table_el if child.tag == W_TR]
        model_rows = block.table.rows if block.table is not None else []
        if len(rows) != len(model_rows):
            raise ProjectionError("StructDoc/table row count mismatch at %s" % table_id)
        for r, (row_el, model_row) in enumerate(zip(rows, model_rows)):
            cells = [child for child in row_el if child.tag == W_TC]
            if len(cells) != len(model_row):
                raise ProjectionError("StructDoc/table cell count mismatch at %s.r%d" % (table_id, r))
            for c, (cell_el, model_cell) in enumerate(zip(cells, model_row)):
                cell_id = "%s.r%dc%d" % (table_id, r, c)
                mapping[cell_id] = cell_el
                content = [child for child in cell_el if child.tag in (W_P, W_TBL)]
                if len(content) != len(model_cell.blocks):
                    raise ProjectionError("StructDoc/cell block count mismatch at %s" % cell_id)
                for k, (child_el, child_block) in enumerate(zip(content, model_cell.blocks)):
                    child_id = "%s.n%d" % (cell_id, k)
                    add_block(child_block, child_el, child_id)
                    if child_block.kind == "paragraph":
                        table_leaf_ids.add(child_id)

    def add_block(block, element, node_id):
        mapping[node_id] = element
        if block.kind == "paragraph":
            if element.tag != W_P:
                raise ProjectionError("StructDoc paragraph ordinal does not identify w:p at %s" % node_id)
        elif block.kind == "table":
            add_table(block, element, node_id)
        else:
            raise ProjectionError("unknown StructDoc block kind at %s: %s" % (node_id, block.kind))

    for block, (body_idx, element) in zip(doc.blocks, body_content):
        if body_idx != block.body_idx:
            raise ProjectionError("StructDoc body index mismatch at b%d" % block.seq)
        add_block(block, element, "b%d" % block.seq)
        if block.kind == "table":
            top_table_ids.add("b%d" % block.seq)
    return mapping, table_ids, top_table_ids, table_leaf_ids


def _select_nodes(spans, index, mapping, table_ids, top_table_ids,
                  table_leaf_ids, struct_doc):
    selected = []
    top_level_kinds = [block.kind for block in struct_doc.blocks]
    for span in spans:
        start, end = index.resolve_ref(span.start), index.resolve_ref(span.end)
        if start is None or end is None:
            raise ProjectionError("span endpoints must be StructDoc IDs")
        if start == end and start in top_table_ids:
            selected.append(mapping[start])
            continue
        if start in table_ids or end in table_ids:
            raise ProjectionError("only a top-level bN table can be selected atomically")
        if start in table_leaf_ids or end in table_leaf_ids:
            raise ProjectionError("cell paragraph endpoints require the owning table as an atomic block")
        if start not in index.by_id or end not in index.by_id:
            raise ProjectionError("span endpoint is not a content node: %s..%s" % (start, end))
        a, b = index.order_of(start), index.order_of(end)
        if a is None or b is None or a > b:
            raise ProjectionError("span endpoints are reversed or unresolved")
        start_top = _top_level_body_id(start)
        end_top = _top_level_body_id(end)
        if start_top is None or end_top is None:
            raise ProjectionError("paragraph span endpoints must be top-level body paragraphs")
        if int(end_top[1:]) >= len(top_level_kinds) or int(start_top[1:]) >= len(top_level_kinds):
            raise ProjectionError("span endpoint is outside top-level StructDoc blocks")
        if any(kind == "table" for kind in top_level_kinds[int(start_top[1:]) + 1:int(end_top[1:])]):
            raise ProjectionError("paragraph span crosses a top-level table block")
        ids = index.order_ids[a:b + 1]
        if any(node_id in table_leaf_ids for node_id in ids):
            raise ProjectionError("paragraph span crosses table content")
        selected_tables = [node_id for node_id in ids if node_id in table_ids]
        if selected_tables:
            raise ProjectionError("paragraph span crosses table content: %s" % selected_tables[0])
        for node_id in ids:
            element = mapping.get(node_id)
            if element is None or element.tag != W_P:
                raise ProjectionError("span cannot map to a paragraph at %s" % node_id)
            selected.append(element)
    return selected


def _top_level_body_id(node_id):
    if not node_id.startswith("b"):
        return None
    base = node_id.split(".", 1)[0]
    return base if base[1:].isdigit() else None


def _validate_payload(element):
    if element.tag not in (W_P, W_TBL):
        raise ProjectionError("only w:p and w:tbl payloads are supported")
    package_scoped = {
        "commentRangeStart", "commentRangeEnd",
        "commentReference", "footnoteReference", "endnoteReference",
        "permStart", "permEnd", "sdt", "customXml", "ins", "del",
        "moveFrom", "moveTo", "moveFromRangeStart", "moveFromRangeEnd",
        "moveToRangeStart", "moveToRangeEnd",
        "customXmlInsRangeStart", "customXmlInsRangeEnd", "customXmlDelRangeStart",
        "customXmlDelRangeEnd", "customXmlMoveFromRangeStart",
        "customXmlMoveFromRangeEnd", "customXmlMoveToRangeStart",
        "customXmlMoveToRangeEnd", "rPrChange", "pPrChange", "tblPrChange",
        "trPrChange", "tcPrChange", "sectPrChange",
        "cellIns", "cellDel", "cellMerge", "conflictIns", "conflictDel",
    }
    revision_markers = {
        "ins", "del", "moveFrom", "moveTo", "moveFromRangeStart",
        "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd",
        "customXmlInsRangeStart", "customXmlInsRangeEnd", "customXmlDelRangeStart",
        "customXmlDelRangeEnd", "customXmlMoveFromRangeStart",
        "customXmlMoveFromRangeEnd", "customXmlMoveToRangeStart",
        "customXmlMoveToRangeEnd", "cellIns", "cellDel", "cellMerge",
        "conflictIns", "conflictDel",
    }
    for node in element.iter():
        local_name = node.tag.split("}", 1)[1] if node.tag.startswith("{%s}" % W) else ""
        is_revision_change = bool(local_name) and local_name.endswith("Change")
        if local_name in package_scoped or is_revision_change:
            category = "tracked-change" if local_name in revision_markers or is_revision_change else "package-scoped"
            raise ProjectionError("unsupported %s construct: %s" % (category, local_name))


def _validate_bookmark_scope(source_body, selected_elements):
    """Require selected bookmark ranges and internal hyperlinks to be complete.

    BlockImporter can remap complete bookmark ranges but deliberately drops
    orphan range markers. This preflight rejects partial or ambiguous source
    scopes before that importer can silently remove selected navigation data.
    """
    all_starts = list(source_body.iter("{%s}bookmarkStart" % W))
    all_ends = list(source_body.iter("{%s}bookmarkEnd" % W))
    selected_nodes = {node for root in selected_elements for node in root.iter()}
    selected_starts = [node for node in all_starts if node in selected_nodes]
    selected_ends = [node for node in all_ends if node in selected_nodes]
    selected_links = [node for root in selected_elements
                      for node in root.iter("{%s}hyperlink" % W)
                      if "{%s}anchor" % W in node.attrib]

    starts_by_id = {}
    ends_by_id = {}
    starts_by_name = {}
    for node in all_starts:
        bookmark_id = node.get("{%s}id" % W)
        name = node.get("{%s}name" % W)
        if bookmark_id:
            starts_by_id.setdefault(bookmark_id, []).append(node)
        if name:
            starts_by_name.setdefault(name, []).append(node)
    for node in all_ends:
        bookmark_id = node.get("{%s}id" % W)
        if bookmark_id:
            ends_by_id.setdefault(bookmark_id, []).append(node)

    permitted_drops = 0
    dangling_toc_anchors_dropped = 0
    dangling_internal_anchors_dropped = 0
    relevant_ids = {node.get("{%s}id" % W) for node in selected_starts + selected_ends}
    for bookmark_id in relevant_ids:
        if not bookmark_id:
            raise ProjectionError("selected bookmark has missing ID")
        starts = starts_by_id.get(bookmark_id, [])
        ends = ends_by_id.get(bookmark_id, [])
        if len(starts) != 1 or len(ends) != 1:
            raise ProjectionError("selected bookmark ID is missing or ambiguous: %s" % bookmark_id)
        start, end = starts[0], ends[0]
        name = start.get("{%s}name" % W)
        if not name:
            raise ProjectionError("selected bookmark has missing name: %s" % bookmark_id)
        if len(starts_by_name.get(name, [])) != 1:
            raise ProjectionError("selected bookmark name is ambiguous: %s" % name)
        start_selected, end_selected = start in selected_nodes, end in selected_nodes
        if start_selected != end_selected:
            omitted = end if start_selected else start
            omitted_tag = "{%s}%s" % (W, "bookmarkEnd" if start_selected else "bookmarkStart")
            if (omitted.tag != omitted_tag or omitted.getparent() is not source_body
                    or omitted not in list(source_body)):
                raise ProjectionError("selected span cuts bookmark pair: %s" % bookmark_id)
            if any(link.get("{%s}anchor" % W) == name for link in selected_links):
                raise ProjectionError("selected hyperlink depends on omitted bookmark: %s" % name)
            permitted_drops += 1

    for hyperlink in selected_links:
        anchor = hyperlink.get("{%s}anchor" % W)
        matches = starts_by_name.get(anchor, [])
        if not anchor or len(matches) != 1:
            # A target absent from the entire immutable source is already a
            # broken navigation link. Preserve all visible children/resources
            # and external relationship attributes; remove only that anchor.
            # Existing targets outside the projection and ambiguous targets
            # still fail closed below: this never hides a cut bookmark.
            if anchor and not matches:
                del hyperlink.attrib["{%s}anchor" % W]
                if re.fullmatch(r"_Toc\d+", anchor):
                    dangling_toc_anchors_dropped += 1
                else:
                    dangling_internal_anchors_dropped += 1
                continue
            raise ProjectionError("selected hyperlink anchor is missing or ambiguous: %s" % anchor)
        start = matches[0]
        bookmark_id = start.get("{%s}id" % W)
        ends = ends_by_id.get(bookmark_id, [])
        if (not bookmark_id or len(ends) != 1 or start not in selected_nodes
                or ends[0] not in selected_nodes):
            raise ProjectionError("selected hyperlink anchor is outside selected scope: %s" % anchor)
    return {"bookmark_markers_dropped": permitted_drops,
            "dangling_toc_anchors_dropped": dangling_toc_anchors_dropped,
            "dangling_internal_anchors_dropped": dangling_internal_anchors_dropped}
