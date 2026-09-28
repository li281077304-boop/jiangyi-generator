# -*- coding: utf-8 -*-
"""Tests for read-only UAT evidence collection.

The OLE fixture is intentionally a dummy binary. It proves package graph and
SHA256 collection only; it is not a substitute for real MathType/Word UAT.
"""
import hashlib
import os
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree


HERE = os.path.dirname(__file__)
APP = os.path.normpath(os.path.join(HERE, "..", "..", "res", "app"))
sys.path.insert(0, HERE)
sys.path.insert(0, APP)

from package_evidence import (A_NS, M_NS, O_NS, R_NS, V_NS, W_NS,
                              collect_package_evidence,
                              collect_selected_block_evidence)


W = "{%s}" % W_NS
R = "{%s}" % R_NS
V = "{%s}" % V_NS
OLE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
M = "{%s}" % M_NS
CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDAT\x08\xd7c\xf8\xcf"
    b"\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00IEND\xaeB`\x82"
)


def body_children(document):
    return [element for element in document.element.body.iterchildren()
            if element.tag in (W + "p", W + "tbl")]


def add_hyperlink(paragraph, rid, text):
    hyperlink = etree.SubElement(paragraph._p, W + "hyperlink")
    hyperlink.set(R + "id", rid)
    run = etree.SubElement(hyperlink, W + "r")
    etree.SubElement(run, W + "t").text = text


class PackageEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _fixture(self):
        image_path = os.path.join(self.tmp.name, "pixel.png")
        with open(image_path, "wb") as image:
            image.write(PNG)
        document = Document()
        document.styles.add_style("EvidenceParagraph", WD_STYLE_TYPE.PARAGRAPH)
        paragraph = document.add_paragraph("selected text", style="EvidenceParagraph")
        math_para = etree.SubElement(paragraph._p, M + "oMathPara")
        math = etree.SubElement(math_para, M + "oMath")
        math_run = etree.SubElement(math, M + "r")
        etree.SubElement(math_run, M + "t").text = "x"
        num_id = document.part.numbering_part.element.findall(W + "num")[0].get(W + "numId")
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = etree.SubElement(p_pr, W + "numPr")
        etree.SubElement(num_pr, W + "ilvl").set(W + "val", "0")
        etree.SubElement(num_pr, W + "numId").set(W + "val", num_id)
        picture = document.add_paragraph().add_run().add_picture(image_path)
        image_rid = picture._inline.graphic.graphicData.pic.blipFill.blip.get(R + "embed")
        vml = document.add_paragraph("VML preview")
        pict = etree.SubElement(vml._p, W + "pict")
        shape = etree.SubElement(pict, V + "shape")
        etree.SubElement(shape, V + "imagedata").set(R + "id", image_rid)
        link_rid = document.part.relate_to("https://example.test/lesson", RT.HYPERLINK,
                                           is_external=True)
        add_hyperlink(document.add_paragraph("link: "), link_rid, "external")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "table text"

        # Deliberately not a valid OLE/MathType object. This only exercises
        # evidence collection for an embedded binary and its recursive rel.
        blob = b"dummy-ole-is-not-real-mathtype\x00fixture"
        embedding = Part(PackURI("/word/embeddings/evidence1.bin"),
                         "application/vnd.openxmlformats-officedocument.oleObject",
                         blob, document.part.package)
        nested = Part(PackURI("/customXml/evidence-nested.xml"), "application/xml",
                      b"<fixture>nested</fixture>", document.part.package)
        embedding.load_rel(CUSTOM_REL, nested, "rIdNested")
        ole_rid = document.part.relate_to(embedding, OLE_REL)
        ole_p = document.add_paragraph("OLE: ")
        obj = etree.SubElement(ole_p._p, W + "object")
        preview_shape = etree.SubElement(obj, V + "shape")
        etree.SubElement(preview_shape, V + "imagedata").set(R + "id", image_rid)
        etree.SubElement(obj, "{%s}OLEObject" % O_NS).set(R + "id", ole_rid)
        path = os.path.join(self.tmp.name, "fixture with spaces 中文.docx")
        document.save(path)
        return document, path, blob, num_id

    def test_collect_package_evidence_reports_saved_package_structure(self):
        document, path, blob, num_id = self._fixture()
        with open(path, "rb") as package:
            before = package.read()
        evidence = collect_package_evidence(path)
        with open(path, "rb") as package:
            after = package.read()

        self.assertEqual(before, after, "collector must be read-only")
        self.assertEqual(os.path.getsize(path), evidence["bytes"])
        self.assertGreaterEqual(evidence["paragraphs"], 5)
        self.assertEqual(1, evidence["tables"])
        self.assertEqual(1, evidence["image_references"])
        self.assertEqual(2, evidence["vml_image_references"])
        self.assertEqual(1, evidence["ole_references"])
        self.assertEqual(1, evidence["ole_preview_image_references"])
        self.assertEqual(1, evidence["omml_math_elements"])
        self.assertEqual(1, evidence["omml_paragraphs"])
        self.assertGreater(evidence["relationship_count"], 0)
        self.assertIn("EvidenceParagraph", evidence["style_references"])
        self.assertIn(num_id, evidence["numbering_references"])
        self.assertEqual([{"partname": "/word/embeddings/evidence1.bin",
                           "sha256": hashlib.sha256(blob).hexdigest()}],
                         evidence["embedding_parts"])

    def test_selected_block_evidence_is_limited_and_walks_ole_graph(self):
        document, path, blob, num_id = self._fixture()
        before_xml = etree.tostring(document.element.body)
        children = body_children(document)
        evidence = collect_selected_block_evidence(
            document, children, [("selected", 1, 5)])

        self.assertEqual(before_xml, etree.tostring(document.element.body),
                         "collector must not modify source XML")
        self.assertEqual(6, evidence["body_blocks"])
        self.assertEqual(6, evidence["paragraphs"])
        self.assertEqual(1, evidence["tables"])
        self.assertEqual(1, evidence["image_references"])
        self.assertEqual(2, evidence["vml_image_references"])
        self.assertEqual(1, evidence["ole_references"])
        self.assertEqual(1, evidence["ole_preview_image_references"])
        self.assertEqual(1, evidence["omml_math_elements"])
        self.assertEqual(1, evidence["omml_paragraphs"])
        self.assertIn("EvidenceParagraph", evidence["style_references"])
        self.assertIn(num_id, evidence["numbering_references"])
        self.assertEqual([{"partname": "/word/embeddings/evidence1.bin",
                           "sha256": hashlib.sha256(blob).hexdigest()}],
                         evidence["embedding_parts"])
        self.assertGreaterEqual(evidence["relationship_count"], 4,
                                "includes direct OLE/image/link refs and nested OLE relation")
        rel_types = {item.get("reltype") for item in evidence["relationship_references"]}
        self.assertIn(OLE_REL, rel_types)
        self.assertIn(RT.HYPERLINK, rel_types)

    def test_selected_block_ranges_reject_invalid_input(self):
        document, _, _, _ = self._fixture()
        with self.assertRaises(ValueError):
            collect_selected_block_evidence(document, body_children(document),
                                            [("bad", 1, 999)])


if __name__ == "__main__":
    unittest.main()
