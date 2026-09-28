# -*- coding: utf-8 -*-
"""Safe, programmatic tests for the Windows real-corpus runner."""
import json
import os
import sys
import tempfile
import unittest
from collections import Counter
from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
from pathlib import Path
from unittest import mock

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "res" / "app"))

import corpus_uat


class CorpusUatTests(unittest.TestCase):
    def _manifest(self, root, repository):
        source = root / "source.docx"
        template = repository / "v1.1-stable" / "template.docx"
        source.write_bytes(b"fixture placeholder; not a real DOCX package")
        template.parent.mkdir(parents=True)
        template.write_bytes(b"fixture template placeholder")
        manifest = root / "manifest.json"
        manifest.write_text(json.dumps({
            "schema_version": 1,
            "samples": [{
                "id": "fixture-01",
                "source": "source.docx",
                "template": "v1.1-stable/template.docx",
                "meta": {key: "fixture" for key in corpus_uat._META_FIELDS},
            }],
        }), encoding="utf-8")
        return source, template, manifest

    def test_default_dry_run_never_opens_docx_or_creates_outputs(self):
        with tempfile.TemporaryDirectory() as tempdir:
            root = Path(tempdir) / "corpus"
            repository = Path(tempdir) / "repo"
            root.mkdir()
            repository.mkdir()
            source, template, manifest = self._manifest(root, repository)
            output = Path(tempdir) / "outputs"
            reports = Path(tempdir) / "reports"
            stdout = StringIO()
            with mock.patch.object(corpus_uat, "collect_package_evidence") as collect, \
                    redirect_stdout(stdout):
                result = corpus_uat.main([
                    "--manifest", str(manifest), "--corpus-root", str(root),
                    "--repo-root", str(repository), "--output-dir", str(output),
                    "--report-dir", str(reports),
                ])

            self.assertEqual(0, result)
            self.assertIn("DRY RUN", stdout.getvalue())
            collect.assert_not_called()
            self.assertTrue(source.exists())
            self.assertTrue(template.exists())
            self.assertFalse(output.exists())
            self.assertFalse(reports.exists())

    @unittest.skipUnless(os.name != "nt", "Mac-only guard regression")
    def test_mac_refuses_explicit_real_corpus_run(self):
        with tempfile.TemporaryDirectory() as tempdir:
            root = Path(tempdir) / "corpus"
            repository = Path(tempdir) / "repo"
            root.mkdir()
            repository.mkdir()
            self._manifest(root, repository)
            output = Path(tempdir) / "outputs"
            reports = Path(tempdir) / "reports"
            with mock.patch.object(corpus_uat, "collect_package_evidence") as collect, \
                    redirect_stdout(StringIO()), redirect_stderr(StringIO()):
                with self.assertRaises(SystemExit) as error:
                    corpus_uat.main([
                        "--manifest", str(root / "manifest.json"),
                        "--corpus-root", str(root), "--repo-root", str(repository),
                        "--output-dir", str(output), "--report-dir", str(reports), "--run",
                    ])

            self.assertEqual(2, error.exception.code)
            collect.assert_not_called()
            self.assertFalse(output.exists())
            self.assertFalse(reports.exists())

    def test_manifest_cannot_escape_corpus_root(self):
        with tempfile.TemporaryDirectory() as tempdir:
            root = Path(tempdir) / "corpus"
            repository = Path(tempdir) / "repo"
            root.mkdir()
            repository.mkdir()
            source, template, manifest = self._manifest(root, repository)
            entries = json.loads(manifest.read_text(encoding="utf-8"))
            entries["samples"][0]["source"] = "../outside.docx"
            manifest.write_text(json.dumps(entries), encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "escapes --corpus-root"):
                corpus_uat._load_manifest(manifest, root.resolve(), repository.resolve())

    def test_output_directory_cannot_target_repository_or_stable_branch(self):
        with tempfile.TemporaryDirectory() as tempdir:
            root = Path(tempdir).resolve()
            repository = root / "repo"
            corpus = root / "corpus"
            repository.mkdir()
            corpus.mkdir()
            with self.assertRaisesRegex(ValueError, "outside the repository"):
                corpus_uat._check_isolated_directory(
                    repository / "out", "output directory", corpus, repository)
            with self.assertRaisesRegex(ValueError, "outside the repository"):
                corpus_uat._check_isolated_directory(
                    repository / "v1.1-stable", "output directory", corpus, repository)

    def test_ole_hash_comparison_detects_missing_binary(self):
        sample_hash = "a" * 64
        source = [{"partname": "/word/embeddings/ole1.bin", "sha256": sample_hash}]
        template = [{"partname": "/word/embeddings/template.bin", "sha256": "b" * 64}]

        good = corpus_uat._ole_preservation(
            source, template,
            template + [{"partname": "/word/embeddings/ole1-imported.bin",
                         "sha256": sample_hash}])
        bad = corpus_uat._ole_preservation(source, template, template)

        self.assertTrue(good["binary_sha256_preserved"])
        self.assertFalse(bad["binary_sha256_preserved"])
        self.assertEqual(Counter([sample_hash]), Counter(bad["missing_sha256_counts"]))

    def test_ole_hash_comparison_preserves_duplicate_binary_parts_as_a_multiset(self):
        sample_hash = "c" * 64
        source = [
            {"partname": "/word/embeddings/ole1.bin", "sha256": sample_hash},
            {"partname": "/word/embeddings/ole2.bin", "sha256": sample_hash},
        ]

        result = corpus_uat._ole_preservation(
            source, [],
            [{"partname": "/word/embeddings/ole-import1.bin", "sha256": sample_hash},
             {"partname": "/word/embeddings/ole-import2.bin", "sha256": sample_hash}])

        self.assertTrue(result["binary_sha256_preserved"])
        self.assertEqual(2, result["expected_source_embedding_parts"])
        self.assertEqual(2, result["preserved_embedding_parts"])

    def test_generation_error_categories_are_structured(self):
        class _Error(Exception):
            report = {"code": "missing_template_anchor"}

        self.assertEqual("TEMPLATE_BUG", corpus_uat._failure_category(_Error()))

    def test_report_keeps_real_corpus_and_word_uat_pending(self):
        with tempfile.TemporaryDirectory() as tempdir:
            report_root = Path(tempdir) / "reports"
            records = [
                {"sample_id": "a", "generation_seconds": 2.0},
                {"sample_id": "b", "generation_seconds": 4.0},
                {"sample_id": "c", "generation_seconds": 6.0},
            ]
            json_path, csv_path = corpus_uat._write_reports(records, report_root)

            payload = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertEqual("NOT_CONFIRMED", payload["corpus_uat"])
            self.assertEqual("NOT_RUN", payload["word_uat"])
            self.assertEqual(4.0, payload["benchmark"]["p50_seconds"])
            self.assertEqual("c", payload["benchmark"]["slowest_sample_id"])
            self.assertTrue(csv_path.is_file())

    def test_fixture_run_records_output_validator_and_word_uat_pending(self):
        with tempfile.TemporaryDirectory() as tempdir:
            root = Path(tempdir)
            corpus_root = root / "corpus"
            repository = root / "repo"
            output_root = root / "outputs"
            corpus_root.mkdir()
            repository.mkdir()
            output_root.mkdir()

            source = corpus_root / "fixture.docx"
            source_doc = Document()
            source_doc.add_paragraph("知识内容")
            source_doc.add_paragraph("【题型1】练习")
            for index in range(1, 12):
                source_doc.add_paragraph("%s.（2024）题目%s" % (index, index))
            source_doc.save(source)

            template = repository / "template.docx"
            template_doc = Document()
            for marker in ("知识精讲", "即时训练", "六、巩固练习"):
                template_doc.add_paragraph(marker)
            template_doc.save(template)

            sample = {
                "id": "generated-fixture",
                "source": source,
                "template": template,
                "meta": {key: "fixture" for key in corpus_uat._META_FIELDS},
            }
            record = corpus_uat._run_sample(sample, corpus_root, output_root, repository)

            self.assertEqual("XML_PACKAGE_PASS_WORD_UAT_PENDING", record["status"])
            self.assertTrue(record["validator"]["valid"])
            self.assertTrue(record["output_package"]["bytes"] > 0)
            self.assertEqual("NOT_RUN", record["word_uat"])
            self.assertEqual("fixture.docx", record["source"])
            self.assertTrue((output_root / "generated-fixture" /
                             "generated-fixture_xml.docx").is_file())


if __name__ == "__main__":
    unittest.main()
